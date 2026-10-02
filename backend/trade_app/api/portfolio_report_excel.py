"""Streaming workbook from validated portfolio report data, never formulas."""
import io
import json

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from trade_app.api.excel_export import _cell
from trade_app.research.portfolio_report_domain import flatten, validate_portfolio_report


def portfolio_report_xlsx(payload, *, metadata=None):
    validate_portfolio_report(payload)
    result, context = flatten(payload), payload['input']
    workbook = Workbook(write_only=True)
    long_text = []

    def safe(value, identity):
        value = _cell(value)
        if isinstance(value, str):
            value = ''.join(f'\\u{ord(char):04x}' if ord(char) < 32 and char not in '\t\n\r' else char for char in value)
            if len(value) > 32000:
                reference = f'LONG-{len(long_text)+1}'
                for offset in range(0, len(value), 32000):
                    long_text.append({'reference': reference, 'field': identity, 'part': offset // 32000 + 1,
                                      'text': value[offset:offset+32000]})
                return reference + ' (complete text in LongText sheet)'
        return value

    def sheet(title, columns, rows):
        page = workbook.create_sheet(title)
        page.freeze_panes = 'A2'
        for index, key in enumerate(columns, 1):
            page.column_dimensions[get_column_letter(index)].width = min(40, max(16, len(key)+3))
        header = []
        for key in columns:
            cell = WriteOnlyCell(page, value=key)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='0A6B54')
            header.append(cell)
        page.append(header)
        for ordinal, row in enumerate(rows, 2):
            page.append([safe(row.get(key), f'{title}!{ordinal}:{key}') for key in columns])

    summary = {key: value for key, value in result.items() if key not in ('trades', 'equity', 'pool_history', 'decisions', 'open_positions')}
    summary.update(title=payload['title'], report_created_at=payload['created_at'], **payload['source'],
                   input_code_sha256=context['code_sha256'], **(metadata or {}))
    sheet('Summary', ['field', 'value'], ({'field': key, 'value': value} for key, value in summary.items()))
    sheet('Parameters', ['section', 'field', 'value'], ({'section': section, 'field': key, 'value': value}
        for section in ('params', 'config') for key, value in context[section].items()))
    sheet('EventProfile', ['field', 'value'], ({'field': key, 'value': value} for key, value in (context['event_profile'] or {}).items()))
    sheet('Trades', ['date', 'symbol', 'side', 'quantity', 'phase', 'reference_price', 'price', 'slippage_rate',
        'fees', 'commission', 'stamp', 'transfer', 'realized_pnl', 'signal_date', 'decision_at', 'known_at',
        'execution_at', 'execution_window', 'reason', 'reason_metrics', 'signal_sha256'], result['trades'])
    sheet('Equity', ['date', 'total_assets', 'cash', 'position_count', 'positions'], result['equity'])
    sheet('OpenPositions', ['symbol', 'quantity', 'cost', 'entry_date', 'entry_price', 'held_bars', 'peak', 'weak_days', 'last_observed_date',
        'entry_box_high', 'band_peak', 'band_observed_date', 'breakeven_active'],
          ({'symbol': symbol, **position} for symbol, position in result['open_positions'].items()))
    sheet('PoolHistory', ['date', 'source_date', 'decision_at', 'scope', 'trigger', 'members'], result['pool_history'])
    sheet('PoolSignals', ['date', 'symbol', 'in_pool', 'score', 'known_at', 'reason', 'components'],
          ({'date': pool['date'], **row} for pool in result['pool_history'] for row in pool['rows']))
    sheet('Decisions', ['date', 'symbol', 'status', 'reason'], result['decisions'])
    sheet('Datasets', ['dataset_id', 'symbol', 'bars_sha256', 'bar_count', 'first_date', 'last_date'],
          ({**{key: item[key] for key in ('dataset_id', 'symbol', 'bars_sha256')}, 'bar_count': len(item['bars']),
            'first_date': item['bars'][0]['event_date'], 'last_date': item['bars'][-1]['event_date']} for item in context['datasets']))
    sheet('Bars', ['symbol', 'event_date', 'open', 'high', 'low', 'close', 'volume', 'amount', 'available_at'],
          ({'symbol': item['symbol'], **bar} for item in context['datasets'] for bar in item['bars']))
    sheet('Checkpoints', ['ordinal', 'prior_sha256', 'result_sha256', 'cursor', 'cash', 'positions', 'pending', 'pool', 'quality_flags'],
          ({**{key: chunk[key] for key in ('ordinal', 'prior_sha256', 'result_sha256')}, **chunk['result']['checkpoint']}
           for chunk in payload['chunks']))
    # Remaining checkpoint fields are losslessly preserved independently of display columns.
    sheet('CheckpointState', ['ordinal', 'field', 'value'], ({'ordinal': chunk['ordinal'], 'field': key, 'value': value}
        for chunk in payload['chunks'] for key, value in chunk['result']['checkpoint'].items()))
    analysis = payload.get('analysis')
    if analysis:
        report = analysis['result']
        sheet('AnalysisMetadata', ['field', 'value'], ({'field': key, 'value': value} for key, value in analysis.items() if key != 'result'))
        sheet('RiskMetrics', ['field', 'value'], ({'field': key, 'value': value} for key, value in report['risk'].items()))
        sheet('AnalysisSections', ['section', 'field', 'value'], ({'section': section, 'field': key, 'value': value}
            for section in ('stability', 'regimes', 'monte_carlo', 'walk_forward', 'methodology', 'daily_detail', 'daily_plan')
            for key, value in report[section].items()))
        sheet('CompletedCycles', ['symbol', 'entry_date', 'exit_date', 'quantity', 'entry_cost', 'net_pnl', 'position_return', 'exit_legs',
            'entry_trade_index', 'exit_trade_indices', 'regime', 'known_20bar_return', 'regime_quality_flags'], report['completed_trades'])
        sheet('IncompleteCycles', ['symbol', 'entry_date', 'quantity', 'remaining_quantity', 'entry_cost', 'net_pnl', 'exit_legs',
            'entry_trade_index', 'exit_trade_indices'], report['open_cycles'])
        sheet('ConditionalPlans', ['symbol', 'side', 'status', 'reason', 'quantity_if_executable', 'price', 'reference_threshold',
            'candidate_order', 'score', 'known_at', 'source_date', 'remaining_observed_bars', 'conditions', 'components'], report['daily_plan'].get('plan_signals', []))
    # Chunk pieces are already <=32k, so this iteration cannot append itself.
    sheet('LongText', ['reference', 'field', 'part', 'text'], list(long_text))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
