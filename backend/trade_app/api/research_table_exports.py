"""Read-only table exports from persisted research; never invoke a strategy."""
import csv
import html
import io
import json
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from trade_app.api.excel_export import _cell
from trade_app.market.models import MarketDataset
from trade_app.platform.types import TradeError
from trade_app.research import backtest_service, legacy_report_service, plateau_service, portfolio_service, signal_workspace_service
from trade_app.research.report_domain import digest, encode
from trade_app.research.service import get_run
from trade_app.research.plateau_models import PlateauExperiment

FILL_COLUMNS = ['run_id', 'symbol', 'date', 'side', 'quantity', 'phase', 'reference_price', 'price',
    'slippage_rate', 'gross', 'fees', 'commission', 'stamp', 'transfer', 'realized_pnl', 'signal_date',
    'decision_at', 'known_at', 'execution_at', 'execution_window', 'reason', 'reason_metrics', 'signal_sha256']
LEGACY_TRADE_COLUMNS = ['序号', '股票代码', '股票名称', '信号日期', '入场日期', '离场日期', '入场事件', '入场阶段',
    '离场原因', '数量', '买入价', '卖出价', '持仓天数', '盈亏额', '盈亏比百分比', '质量分', '健康分', '事件分', '风险分', '事件等级', '确认状态']
POINT_COLUMNS = ['rank', 'status', 'error', 'plateau_score', 'local_score', 'point_score', 'score',
    'total_return', 'max_drawdown', 'win_rate', 'trade_count', 'annual_trades', 'candidate_count', 'skipped_count',
    'fill_rate', 'max_concurrent_positions', 'neighbor_pass_rate', 'neighbor_median_score', 'neighbor_p25_score',
    'sensitivity_penalty', 'passes_hard_filters', 'hard_filter_failures', 'region_id', 'region_rank', 'cache_hit',
    'window_days', 'min_score', 'stop_loss', 'take_profit', 'trailing_stop_pct', 'intraday_trailing_reduce_ratio',
    'max_positions', 'position_pct', 'max_symbols', 'priority_topk_per_day']


def csv_bytes(columns, rows):
    output = io.StringIO(newline='')
    writer = csv.writer(output, lineterminator='\r\n')
    writer.writerow([_cell(key) for key in columns])
    for row in rows: writer.writerow([_cell(row.get(key)) for key in columns])
    return ('\ufeff' + output.getvalue()).encode('utf-8')


def table_html(title, columns, rows, metadata):
    esc = lambda value: html.escape(str(_cell(value) if value is not None else ''), quote=True)
    meta = ''.join('<dt>' + esc(key) + '</dt><dd>' + esc(value) + '</dd>' for key, value in metadata.items())
    body = ''.join('<tr>' + ''.join('<td>' + esc(row.get(key)) + '</td>' for key in columns) + '</tr>' for row in rows)
    return ('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>' + esc(title) + '</title><body><h1>' + esc(title)
        + '</h1><dl>' + meta + '</dl><table><thead><tr>' + ''.join('<th>' + esc(key) + '</th>' for key in columns)
        + '</tr></thead><tbody>' + body + '</tbody></table></body></html>').encode('utf-8')


def workbook_bytes(sheets):
    """Keep all source text losslessly; Excel's 32767-char cells must not truncate."""
    workbook, overflow = Workbook(write_only=True), []
    def cell_value(value, identity):
        value = _cell(value)
        if isinstance(value, str):
            value = ''.join(f'\\u{ord(ch):04x}' if ord(ch) < 32 and ch not in '\t\r\n' else ch for ch in value)
            if len(value) > 32000:
                reference = f'LONG-{len(overflow) + 1}'
                overflow.extend({'reference': reference, 'field': identity, 'part': index // 32000 + 1, 'text': value[index:index + 32000]}
                    for index in range(0, len(value), 32000))
                return reference + ' (complete text in LongText)'
        return value
    for name, columns, rows in [*sheets, ('LongText', ['reference', 'field', 'part', 'text'], overflow)]:
        page = workbook.create_sheet(name); page.freeze_panes = 'A2'
        headers = []
        for key in columns:
            cell = WriteOnlyCell(page, value=cell_value(key, name + ':header')); cell.data_type = 's'; cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='0A6B54'); headers.append(cell)
        page.append(headers)
        for index, row in enumerate(rows, 2):
            values = []
            for key in columns:
                # Chunks already passed escaping; do not add an extra apostrophe
                # when an arbitrary split happens to begin with '=' or '+'.
                value = row.get(key) if name == 'LongText' and key == 'text' else cell_value(row.get(key), f'{name}!{index}:{key}')
                cell = WriteOnlyCell(page, value=value)
                if isinstance(value, str): cell.data_type = 's'
                values.append(cell)
            page.append(values)
    output = io.BytesIO(); workbook.save(output)
    return output.getvalue()


def native_trades(session, run_id, *, portfolio=False):
    run = portfolio_service.get_portfolio(session, run_id, full=True) if portfolio else backtest_service.get_backtest(session, run_id)
    if run['state'] != 'succeeded' or not run.get('result'):
        raise TradeError('RESEARCH_EXPORT_NOT_READY', '只导出已完成且有冻结结果的历史任务', 409)
    result = run['result']
    if not portfolio and digest(encode(result)) != run.get('result_sha256'):
        raise TradeError('RESEARCH_EXPORT_HASH_MISMATCH', '回测结果摘要不符，无法导出', 409)
    dataset = None if portfolio else session.get(MarketDataset, run['dataset_id'])
    trades = [{'run_id': run_id, 'symbol': dataset.symbol if dataset else None, **row} for row in result['trades']]
    # Include newly added native execution fields without silently dropping them.
    equity = result.get('equity', [])
    meta = {'run_id': run_id, 'date_from': equity[0]['date'] if equity else None, 'date_to': equity[-1]['date'] if equity else None,
        'row_count': len(trades), 'row_grain': 'executed_fill_in_source_order', 'result_sha256': run['result_sha256'],
        'code_sha256': run['code_sha256'], 'scope': 'frozen_portfolio' if portfolio else 'frozen_single_symbol',
        'note': '逐笔成交原记录；不将买卖腿伪装成旧版已完成交易回合，空值表示未记录。'}
    for row in trades:
        row.update(export_date_from=meta['date_from'], export_date_to=meta['date_to'], export_result_sha256=run['result_sha256'], export_scope=meta['scope'])
    columns = [*FILL_COLUMNS, *sorted(set().union(*(set(row) for row in trades)) - set(FILL_COLUMNS))]
    return columns, trades, meta


def legacy_trades(session, report_id, detail_key=None):
    report = legacy_report_service.get_report(session, report_id)
    if report['mode'] != 'legacy_readonly': raise TradeError('LEGACY_REPORT_READONLY_REQUIRED', '原件存档没有经过结构校验的成交表', 409)
    source = report['payload']
    if detail_key is not None:
        source = source['point_details'].get(detail_key)
        if source is None: raise TradeError('LEGACY_POINT_NOT_FOUND', '该原参数点明细不存在', 404)
    result, request = source.get('run_result'), source.get('run_request') or {}
    if result is None: raise TradeError('LEGACY_TRADES_UNAVAILABLE', '原文件没有完整回测成交结果', 409)
    columns = LEGACY_TRADE_COLUMNS
    keys = ['symbol', 'name', 'signal_date', 'entry_date', 'exit_date', 'entry_signal', 'entry_phase', 'exit_reason',
            'quantity', 'entry_price', 'exit_price', 'holding_days', 'pnl_amount', 'pnl_ratio', 'entry_quality_score',
            'health_score', 'event_score', 'risk_score', 'event_grade', 'confirmation_status']
    rows = []
    for index, trade in enumerate(result['trades'], 1):
        row = {'序号': index, **{column: trade.get(key) for column, key in zip(columns[1:], keys)}}
        row['盈亏比百分比'] = str(Decimal(str(trade['pnl_ratio'])) * 100) if trade.get('pnl_ratio') is not None else None
        rows.append(row)
    return columns, rows, {'report_id': report_id, 'detail_key': detail_key, 'date_from': request.get('date_from'),
        'date_to': request.get('date_to'), 'row_count': len(rows), 'source_sha256': report['source_sha256'],
        'note': '旧来源逐回合记录，未重新计算。原值保留，缺少评分/价格不补0；列序沿用旧 buildTradesExportRows。'}


def legacy_plateau_xlsx(session, report_id):
    report = legacy_report_service.get_report(session, report_id)
    plateau = report['payload']['plateau_result']
    if report['mode'] != 'legacy_readonly' or plateau is None: raise TradeError('LEGACY_PLATEAU_UNAVAILABLE', '旧档案没有已校验的平原结果', 409)
    points = []
    for index, point in enumerate(plateau['points'], 1):
        points.append({**point, **point['params'], **point['stats'], 'rank': index, 'status': 'failed' if point.get('error') else 'succeeded'})
    meta = {key: value for key, value in plateau.items() if key not in ('points', 'base_payload', 'regions', 'correlations', 'notes')}
    meta.update(report_id=report_id, source_sha256=report['source_sha256'], mode='legacy_readonly', points_count=len(points),
                note='原顺序、全部点含失败；所有百分比数值是比例值，0.1=10%。原得分未重算。')
    def records(value): return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []
    regions = []
    for region in records(plateau.get('regions')):
        center = region.get('center_point') if isinstance(region.get('center_point'), dict) else {}
        stats = center.get('stats') if isinstance(center.get('stats'), dict) else {}
        regions.append({**region, 'center_plateau_score': center.get('plateau_score'), 'center_total_return': stats.get('total_return'),
            'center_max_drawdown': stats.get('max_drawdown'), 'center_params': center.get('params'), 'source_region_json': region})
    region_columns = ['region_id', 'region_rank', 'point_count', 'region_score', 'median_local_score', 'p25_local_score',
        'median_point_score', 'median_total_return', 'best_total_return', 'center_margin_score', 'size_score', 'oos_pass_rate',
        'center_plateau_score', 'center_total_return', 'center_max_drawdown', 'center_params', 'parameter_ranges', 'source_region_json']
    correlations = records(plateau.get('correlations'))
    corr_columns = ['parameter', 'parameter_label', 'score_corr', 'score_direction', 'total_return_corr', 'total_return_direction', 'win_rate_corr', 'win_rate_direction', 'source_correlation_json']
    def direction(value): return None if not isinstance(value, (int, float)) else '正相关' if value > .05 else '负相关' if value < -.05 else '弱相关'
    correlation_rows = [{**row, 'score_direction': direction(row.get('score_corr')), 'total_return_direction': direction(row.get('total_return_corr')),
        'win_rate_direction': direction(row.get('win_rate_corr')), 'source_correlation_json': row} for row in correlations]
    return workbook_bytes([
        ('PlateauSummary', ['field', 'value'], [{'field': key, 'value': value} for key, value in meta.items()]),
        ('BasePayload', list(plateau['base_payload']), [plateau['base_payload']]),
        ('PlateauPoints', [*POINT_COLUMNS, 'detail_key', 'source_point_json'], [{**row, 'source_point_json': point} for row, point in zip(points, plateau['points'])]),
        ('PlateauRegions', region_columns, regions),
        ('PlateauCorr', corr_columns, correlation_rows),
        ('PlateauNotes', ['index', 'note'], [{'index': index + 1, 'note': value} for index, value in enumerate(plateau.get('notes') or [])]),
    ])


def native_plateau_xlsx(session, experiment_id):
    experiment = plateau_service.get_plateau(session, experiment_id)
    stored = session.get(PlateauExperiment, experiment_id)
    if digest(encode(json.loads(stored.input_json))) != stored.input_sha256:
        raise TradeError('PLATEAU_EXPORT_HASH_MISMATCH', '实验冻结输入摘要不符', 409)
    if experiment['state'] not in ('succeeded', 'failed', 'cancelled'):
        raise TradeError('PLATEAU_EXPORT_NOT_FINAL', '请先完成或取消实验后导出固定状态的全部参数点', 409)
    analysis = experiment.get('analysis') or {}
    scored = {row['id']: row for row in analysis.get('points', [])}
    points = []
    for point in experiment['points']:
        # Payload and result digests are checked, independently of current code version.
        if digest(encode({key: point[key] for key in ('params', 'config', 'axis_values')})) != point['point_sha256']:
            raise TradeError('PLATEAU_EXPORT_HASH_MISMATCH', '参数点输入摘要不符', 409)
        plateau_service.get_point(session, experiment_id, point['id'])
        if point['id'] in scored and scored[point['id']].get('metrics') != point['metrics']:
            raise TradeError('PLATEAU_EXPORT_HASH_MISMATCH', '参数点指标与冻结排名快照不符', 409)
        row = {**point, **(point['metrics'] or {}), **scored.get(point['id'], {}), 'status': point['state']}
        row.update({'axis.' + key: value for key, value in point['axis_values'].items()})
        row.update({'params.' + key: value for key, value in point['params'].items()})
        row.update({'config.' + key: value for key, value in point['config'].items()})
        points.append(row)
    columns = ['id', 'ordinal', 'rank', 'status', 'error', 'attempt_number', 'total_return', 'max_drawdown', 'win_rate',
        'trade_count', 'point_score', 'plateau_score', 'local_score', 'neighbor_pass_rate', 'neighbor_median_score',
        'neighbor_p25_score', 'sensitivity_penalty', 'passes_hard_filters', 'hard_filter_failures', 'point_sha256', 'result_sha256']
    columns += sorted(set().union(*(set(row) for row in points)) - set(columns))
    meta = {key: value for key, value in experiment.items() if key not in ('points', 'analysis', 'frozen_context')}
    meta['note'] = '全部采样点，保留失败/取消/无效项；ordinal是采样顺序，rank为已保存排名；比例0.1=10%，缺分不补0。'
    return workbook_bytes([
        ('PlateauSummary', ['field', 'value'], [{'field': key, 'value': value} for key, value in meta.items()]),
        ('BasePayload', ['field', 'value'], [{'field': key, 'value': value} for key, value in experiment['frozen_context'].items()]),
        ('PlateauPoints', columns, points),
        ('PlateauRegions', ['source_region_json'], [{'source_region_json': row} for row in analysis.get('regions', [])]),
        ('PlateauCorr', ['source_correlation_json'], [{'source_correlation_json': row} for row in analysis.get('correlations', [])]),
        ('PlateauNotes', ['note'], [{'note': note} for note in analysis.get('notes', [])]),
    ])


def cross_validation_csv(session, report_id):
    report = signal_workspace_service.get_report(session, report_id)
    result = report['result']; strategies = result['strategies']
    detail_columns = ['评分', '健康分', '事件分', '阶段', '触发日', '出现天数', '出现日期']
    headers = ['代码', '名称', '重叠策略数', '共振强度', '最高分', '平均分', '区间涨幅%', '最大回撤%']
    labels = {item['id']: f"{item['name']} [{item['id']}]" for item in strategies}
    headers += [labels[item['id']] + '_' + field for item in strategies for field in detail_columns]
    meta = {'报告ID': report_id, '扫描ID': report['scan_id'], '报告截至日': report['as_of_date'],
        '扫描开始日': result['scan_request'].get('date_from') or result['scan_request'].get('as_of_date'),
        '扫描结束日': result['scan_request'].get('date_to') or result['scan_request'].get('as_of_date'),
        '冻结过滤条件': report['request'], '算法版本': result['version'],
        '口径': '代表点=各策略最高质量分，平分取最早观察日；共振=策略数×出现日期去重数。跨策略最高/均分仅旧格式字段，不是统一排名；价格变化不含费用；缺失留空。'}
    headers += [*meta, '实际价格起日', '实际价格止日', '价格根数']
    rows = []
    for item in result['per_symbol']:
        perf = item['range_performance']; selected = [row for row in result['rows'] if row['symbol'] == item['symbol'] and not row['excluded_reasons']]
        row = {'代码': item['symbol'], '名称': item.get('name'), '重叠策略数': item['overlap_count'],
            '共振强度': item['overlap_count'] * max(1, len(item['signal_dates'])), '区间涨幅%': perf['return_pct'],
            '最大回撤%': perf['max_drawdown_pct'], **meta, '实际价格起日': perf['start_date'],
            '实际价格止日': perf['end_date'], '价格根数': perf['bar_count']}
        scores = []
        for strategy in item['strategies']:
            evidence = sorted([row for row in selected if row['strategy_id'] == strategy['strategy_id']], key=lambda row: (row['decision_date'], row['run_id']))
            if not evidence: continue
            best = max(evidence, key=lambda row: row.get('entry_quality_score') if row.get('entry_quality_score') is not None else -1)
            quality = best.get('entry_quality_score')
            if quality is not None: scores.append(quality)
            source = get_run(session, best['run_id'])
            if signal_workspace_service.digest(source) != result['research_run_hashes'].get(best['run_id']):
                raise TradeError('SIGNAL_EVIDENCE_CHANGED', '导出所需的原研究摘要与冻结报告不符', 409)
            indicator = source['result'].get('indicator') or {}
            values = [quality, best.get('health_score', indicator.get('health_score')), best.get('event_score', indicator.get('event_score')),
                best.get('phase'), best.get('trigger_date'), strategy['appearance_count'], '; '.join(strategy['dates'])]
            row.update({labels[strategy['strategy_id']] + '_' + field: value for field, value in zip(detail_columns, values)})
        row.update({'最高分': max(scores) if scores else None, '平均分': sum(scores) / len(scores) if scores else None})
        rows.append(row)
    return csv_bytes(headers, rows)
