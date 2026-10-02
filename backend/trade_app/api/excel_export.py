"""Account-scoped workbook export for the rebuilt ledger."""
from __future__ import annotations

import io
import json
import re
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from trade_app.analytics.service import projection_status
from trade_app.platform.types import TradeError
from trade_app.research.backtest_service import get_backtest
from trade_app.trading.service import account_or_error, list_flows, list_snapshots, list_trades
from trade_app.trading.simulation import list_fills, list_orders, portfolio, sim_account


REAL_COLUMNS = {
    'Trades': ['id', 'trade_date', 'symbol', 'name', 'side', 'quantity', 'price',
               'fee', 'calculated_fee', 'fee_source', 'fee_rule_version', 'note', 'revision'],
    'Flows': ['id', 'flow_date', 'kind', 'amount', 'note', 'revision'],
    'Snapshots': ['id', 'snap_date', 'total_assets', 'available_cash', 'position_value',
                  'positions', 'revision'],
    'Rounds': ['id', 'symbol', 'name', 'start_date', 'end_date', 'status', 'pnl', 'trade_ids'],
    'NAV': ['date', 'nav', 'assets', 'drawdown_pct', 'quality'],
}
SIM_COLUMNS = {
    'Orders': ['id', 'symbol', 'side', 'quantity', 'limit_price', 'signal_date',
               'submit_date', 'status', 'revision'],
    'Fills': ['id', 'order_id', 'fill_date', 'symbol', 'side', 'quantity',
              'fill_price', 'gross', 'commission', 'stamp', 'transfer',
              'realized_pnl', 'price_source'],
    'Positions': ['symbol', 'quantity', 'sellable_quantity', 'cost_basis'],
}


def _cell(value: object) -> str | int | float | None:
    if value is None:
        return None
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, (date, datetime, Decimal)):
        value = str(value)
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        if not re.fullmatch(r'-?\d+(?:\.\d+)?', value):
            return "'" + value
    return value


def _sheet(workbook: Workbook, title: str, columns: list[str], rows: list[dict]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(columns)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{max(2, len(rows) + 1)}'
    for index, header in enumerate(columns, 1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0A6B54')
        cell.alignment = Alignment(wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 3, 15), 35)
    for row in rows:
        sheet.append([_cell(row.get(column)) for column in columns])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and len(cell.value) > 60:
                cell.alignment = Alignment(wrap_text=True, vertical='top')


def account_xlsx(session: Session, account_id: str) -> bytes:
    account = account_or_error(session, account_id)
    workbook = Workbook()
    workbook.remove(workbook.active)
    if account.kind == 'real':
        projection = projection_status(session, account_id)
        result = projection['result'] or {}
        meta = [{'field': 'account_id', 'value': account_id},
                {'field': 'account_name', 'value': account.name},
                {'field': 'account_kind', 'value': account.kind},
                {'field': 'currency', 'value': account.currency},
                {'field': 'input_revision', 'value': projection['account_input_revision']},
                {'field': 'projection_status', 'value': projection['status']},
                {'field': 'projection_version', 'value': projection['projection_version']},
                {'field': 'scope', 'value': '全部已确认记录；作废记录不含在交易 / 资金表'}]
        _sheet(workbook, 'Meta', ['field', 'value'], meta)
        tables = {'Trades': list_trades(session, account_id),
                  'Flows': list_flows(session, account_id),
                  'Snapshots': list_snapshots(session, account_id),
                  'Rounds': result.get('rounds', []),
                  'NAV': result.get('nav', {}).get('points', [])}
        for name, columns in REAL_COLUMNS.items():
            _sheet(workbook, name, columns, tables[name])
    else:
        sim_account(session, account_id)
        state = portfolio(session, account_id)
        orders = list_orders(session, account_id)
        order_map = {row['id']: row for row in orders}
        fills = [{**row, 'symbol': order_map[row['order_id']]['symbol'],
                  'side': order_map[row['order_id']]['side'],
                  'quantity': order_map[row['order_id']]['quantity']}
                 for row in list_fills(session, account_id)]
        meta = [{'field': 'account_id', 'value': account_id},
                {'field': 'account_name', 'value': account.name},
                {'field': 'account_kind', 'value': account.kind},
                {'field': 'currency', 'value': account.currency},
                {'field': 'as_of_date', 'value': state['as_of_date']},
                {'field': 'wallet_revision', 'value': state['wallet_revision']},
                {'field': 'config_version', 'value': state['config_version']},
                {'field': 'valuation_quality', 'value': state['valuation_quality']},
                {'field': 'scope', 'value': '模拟委托、成交与当前持仓；不含外部行情估值'}]
        _sheet(workbook, 'Meta', ['field', 'value'], meta)
        for name, columns, rows in (
            ('Orders', SIM_COLUMNS['Orders'], orders),
            ('Fills', SIM_COLUMNS['Fills'], fills),
            ('Positions', SIM_COLUMNS['Positions'], state['positions']),
        ):
            _sheet(workbook, name, columns, rows)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def backtest_xlsx(session: Session, run_id: str) -> bytes:
    run = get_backtest(session, run_id)
    return backtest_report_xlsx(run)


def backtest_report_xlsx(run: dict, *, report_metadata: dict | None = None) -> bytes:
    """Render a frozen successful run, including imported report snapshots."""
    if run['state'] != 'succeeded' or not run.get('result'):
        raise TradeError('BACKTEST_NOT_READY', '回测尚未完成，无法导出报告', 409)
    result = run['result']
    workbook = Workbook()
    workbook.remove(workbook.active)
    meta = [{'field': key, 'value': run.get(key)} for key in (
        'id', 'dataset_id', 'strategy_id', 'strategy_version', 'execution_version',
        'calculation_version', 'code_sha256', 'result_sha256', 'attempt_number', 'state', 'created_at', 'updated_at')]
    meta.extend({'field': key, 'value': result.get(key)} for key in (
        'initial_capital', 'ending_assets', 'total_return', 'max_drawdown',
        'trade_count', 'win_rate', 'signal_count', 'realized_pnl'))
    meta.extend([{'field': 'quality_flags', 'value': result.get('quality_flags', [])},
                 {'field': 'limitations', 'value': result.get('limitations', [])},
                 {'field': 'scope', 'value': '单股新执行配置；不等同于旧矩阵或传统回测'}])
    meta.extend({'field': key, 'value': value} for key, value in (report_metadata or {}).items())
    _sheet(workbook, 'Summary', ['field', 'value'], meta)
    _sheet(workbook, 'Params', ['field', 'value'], [
        {'field': key, 'value': value} for key, value in run['params'].items()])
    _sheet(workbook, 'Config', ['field', 'value'], [
        {'field': key, 'value': value} for key, value in run['config'].items()])
    _sheet(workbook, 'Trades', ['date', 'side', 'quantity', 'price', 'reference_price', 'slippage_rate', 'fees',
                                'realized_pnl', 'signal_date', 'known_at', 'decision_at',
                                'execution_at', 'reason', 'reason_metrics', 'signal_sha256'], result.get('trades', []))
    _sheet(workbook, 'Equity', ['date', 'total_assets', 'cash', 'quantity'],
           result.get('equity', []))
    _sheet(workbook, 'Decisions', ['date', 'decision_at', 'known_at', 'source_date', 'status',
                                   'signal', 'draft_eligible', 'primary_event', 'trigger_date',
                                   'reasons', 'draft_block_reason', 'quality_flags'], result.get('decisions', []))
    analysis = result.get('advanced_analysis') or {'status': 'not_generated', 'reason': '旧结果未生成高级分析'}
    _sheet(workbook, 'Analysis', ['field', 'value'], [{'field': key, 'value': value} for key, value in analysis.items()
                                                  if key not in {'completed_trades', 'stability', 'regimes'}])
    if analysis.get('status') == 'generated':
        _sheet(workbook, 'Monthly', ['month', 'first_observed_date', 'last_observed_date', 'sessions', 'start_assets', 'end_assets', 'return'], analysis['stability']['months'])
        _sheet(workbook, 'Regimes', ['regime', 'trade_count', 'win_rate', 'net_pnl', 'pnl_contribution', 'mean_position_return', 'closed_trade_sequence_drawdown'], analysis['regimes']['buckets'])
        _sheet(workbook, 'RoundTrips', ['entry_date', 'exit_date', 'quantity', 'net_pnl', 'position_return', 'account_return', 'holding_bars'], analysis['completed_trades'])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
