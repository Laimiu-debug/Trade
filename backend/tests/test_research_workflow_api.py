import io
import csv

from openpyxl import load_workbook

from trade_app.platform.backup import create_backup, restore_to_new_directory
from trade_app.research.scan_job_service import process_one_scan_chunk
from test_wyckoff_domain import event_rich_bars
from test_wyckoff_research_api import client_for, custom_payload, data, dataset, write


def test_scan_basket_reports_and_templates_survive_full_backup(tmp_path):
    original = tmp_path / 'original'
    bars = event_rich_bars()
    with client_for(original) as client:
        ds = dataset(client)
        profile = data(write(client, '/research/event-profiles', {'profile': custom_payload('固定模板')}))
        body = {'dataset_ids': [ds['id']], 'as_of_date': bars[84]['event_date'],
                'strategies': [{'strategy_id': strategy, 'params': {'min_score': '0', 'min_event_count': '0'}}
                               for strategy in ('wyckoff_trend_v1', 'score_only_rank_v1')],
                'event_profile_id': profile['profile_id'], 'event_profile_revision': 1}
        job = data(write(client, '/research/scan-jobs', body))
        assert job['state'] == 'queued'
        for _ in range(100):
            assert process_one_scan_chunk(client.app.state.db_factory, original)
            job = data(client.get('/api/v1/research/scan-jobs/' + job['id']))
            if job['state'] == 'succeeded':
                break
            assert job['state'] in ('queued', 'running'), job
        assert job['state'] == 'succeeded'
        scan = data(client.get('/api/v1/research/scans/' + job['scan_id']))
        assert scan['intersection_count'] == 1
        name = '=HYPERLINK("https://invalid.example","text")'
        basket = data(write(client, '/research/baskets', {'name': name, 'scan_id': scan['id'],
                         'selection': 'intersection', 'holding_bars': 3}))
        assert basket['total_constituents'] == 1
        assert len(basket['constituents'][0]['run_ids']) == 2
        assert basket['source']['kind'] == 'scan'
        path = '/research/baskets/' + basket['id']
        report = data(write(client, path + '/evaluations', {'expected_revision': 1,
                           'as_of_date': bars[-1]['event_date']}))
        assert report['summary']['t1']['completed_count'] == 1
        csv_response = client.get('/api/v1' + path + '/export.csv')
        assert csv_response.status_code == 200
        rows = list(csv.DictReader(io.StringIO(csv_response.content.decode('utf-8-sig'))))
        assert rows[0]['basket_name'] == "'" + name
        assert rows[0]['basket_revision'] == '1' and rows[0]['evaluation_id'] == report['id']
        xlsx = client.get('/api/v1' + path + '/export.xlsx')
        assert xlsx.status_code == 200
        workbook = load_workbook(io.BytesIO(xlsx.content))
        assert 'Returns' in workbook.sheetnames
        assert {row[0].value: row[1].value for row in workbook['Basket'].iter_rows(min_row=2)}['name'] == "'" + name
        for worksheet in workbook:
            assert not any(cell.data_type == 'f' for row in worksheet for cell in row)
        invalid = write(client, path, {'expected_revision': 1, 'quantity': 101}, 'PUT')
        assert invalid.status_code == 422
        renamed = data(write(client, path, {'expected_revision': 1, 'name': '修订后的篮子'}, 'PUT'))
        assert renamed['revision'] == 2
        # A revision with no evaluation must not inherit old revision's returns.
        current_csv = list(csv.DictReader(io.StringIO(client.get('/api/v1' + path + '/export.csv').content.decode('utf-8-sig'))))
        assert current_csv[0]['status'] == 'not_evaluated' and current_csv[0]['raw_return'] == ''
        assert write(client, '/research/baskets/delete-batch', {'items': [{'id': basket['id'], 'expected_revision': 1}]}).status_code == 409
        data(write(client, '/research/baskets/delete-batch', {'items': [{'id': basket['id'], 'expected_revision': 2}]}))
        assert client.get('/api/v1' + path).status_code == 404
        assert data(client.get('/api/v1' + path + '?include_deleted=true'))['deleted'] is True
        report_path = path + '/evaluations/' + report['id']
        assert data(client.get('/api/v1' + report_path)) == report
        assert client.get('/api/v1' + report_path + '/export.xlsx').status_code == 200
        data(write(client, '/research/scans/' + scan['id'], {}, 'DELETE'))
        assert data(client.get('/api/v1/research/runs/' + scan['result']['run_ids'][0]))
    restored = tmp_path / 'restored'
    restore_to_new_directory(create_backup(original), restored)
    with client_for(restored) as client:
        assert data(client.get('/api/v1' + report_path)) == report
        assert data(client.get('/api/v1' + path + '?include_deleted=true'))['revision'] == 3
        assert len(data(client.get('/api/v1' + path + '/audit'))) == 4
        assert any(item['profile_id'] == profile['profile_id'] for item in data(client.get('/api/v1/research/event-profiles'))['profiles'])
        assert client.get('/api/v1' + report_path + '/export.xlsx').status_code == 200


def test_scan_write_requires_csrf_and_idempotency(tmp_path):
    with client_for(tmp_path) as client:
        ds = dataset(client)
        body = {'dataset_ids': [ds['id']], 'as_of_date': ds['last_date'],
                'strategies': [{'strategy_id': 'relative_strength_breakout_v1'}]}
        assert client.post('/api/v1/research/scan-jobs', json=body).status_code == 400
        token = client.headers.pop('X-CSRF-Token')
        assert write(client, '/research/scan-jobs', body).status_code == 403
        client.headers['X-CSRF-Token'] = token
        first = data(write(client, '/research/scan-jobs', body, key='fixed-key'))
        assert data(write(client, '/research/scan-jobs', body, key='fixed-key')) == first
        assert write(client, '/research/scan-jobs', {**body, 'strict': False}, key='fixed-key').status_code == 409
        disabled = write(client, '/research/scans', body)
        assert disabled.status_code == 409
        assert disabled.json()['error']['code'] == 'ASYNC_SCAN_REQUIRED'
