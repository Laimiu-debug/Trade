from copy import deepcopy
import io
import json
import stat
import zipfile
from uuid import uuid4

from openpyxl import load_workbook
import pytest
from sqlalchemy import event, select, text

from trade_app.api.portfolio_report_excel import portfolio_report_xlsx
from trade_app.market.domain import normalize_bars
from trade_app.platform.backup import create_backup, restore_to_new_directory
from trade_app.platform.db import open_database
from trade_app.platform.types import TradeError, utc_now
from trade_app.research import portfolio_report_domain as domain, portfolio_report_service as service, portfolio_service
from trade_app.research.portfolio_domain import VERSION, MATRIX_VERSION, initial_checkpoint, run_chunk, summary
from trade_app.research.portfolio_report_models import PortfolioReport
from test_portfolio import context, create_run
from test_wyckoff_research_api import client_for, data, write


@pytest.fixture
def payload():
    ctx = context(count=75, start=60)
    ctx.update(version=VERSION, mode='matrix_raw_s1_s9', strategy_id='matrix_raw_s1_s9', strategy_version=MATRIX_VERSION,
               universe_scope='fixed_research_sample', event_profile=None, params={}, code_sha256='c'*64, budget={})
    bars = ctx['datasets'][0]['bars']
    bars[62].update(close='10.4', high='10.5', volume=5000)
    for item in ctx['datasets']:
        item['bars'] = normalize_bars(item['bars'])
        item['bars_sha256'] = domain.digest(domain.encode(item['bars']))
    checkpoint, chunks = initial_checkpoint(ctx), []
    initial = deepcopy(checkpoint)
    while checkpoint['cursor'] < len(ctx['calendar']):
        result = run_chunk(ctx, checkpoint)
        chunks.append({'ordinal': len(chunks), 'prior_sha256': domain.digest(domain.encode(checkpoint)),
                       'result_sha256': domain.digest(domain.encode(result)), 'result': result})
        checkpoint = result['checkpoint']
    result = summary(ctx, checkpoint)
    input_sha = domain.digest(domain.encode(ctx))
    now = utc_now()
    value = {'format': domain.FORMAT, 'version': 1, 'scope': domain.SCOPE, 'title': '组合冻结报告', 'created_at': now,
        'source': {'run_id': 'run-portfolio', 'state': 'succeeded', 'created_at': now, 'updated_at': now,
            'input_sha256': input_sha, 'result_sha256': domain.digest(domain.encode({'summary': result,
            'chunks': [chunk['result_sha256'] for chunk in chunks], 'input_sha256': input_sha}))},
        'input': ctx, 'initial_checkpoint': initial, 'summary': result, 'chunks': chunks}
    return domain.validate_portfolio_report(domain.strict_json(domain.encode(value)))


def pack(value):
    return domain.package_portfolio_report(value, portfolio_report_xlsx(value))


def files_for(value):
    with zipfile.ZipFile(io.BytesIO(pack(value))) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def archive(files, *, update=True):
    if update:
        manifest = json.loads(files['manifest.json'])
        manifest['files'] = {name: {'size': len(raw), 'sha256': domain.digest(raw)} for name, raw in files.items() if name != 'manifest.json'}
        files['manifest.json'] = domain.encode(manifest)
    target = io.BytesIO()
    with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as output:
        for name, raw in files.items():
            output.writestr(name, raw)
    return target.getvalue()


def test_full_checkpoint_report_roundtrip_html_escape_and_workbook_complete(payload):
    payload['title'] = '<script>alert("untrusted")</script>'
    contents = pack(payload)
    restored = domain.unpack_portfolio_report(contents)
    assert restored == payload
    rendered = domain.render_portfolio_html(restored, origin='import')
    assert '<script>' not in rendered and '&lt;script&gt;' in rendered
    assert '导入快照，未重新计算' in rendered and '固定研究样本不等同于历史全市场' in rendered
    result = domain.flatten(restored)
    book = load_workbook(io.BytesIO(portfolio_report_xlsx(restored)), read_only=True, data_only=False)
    assert {'Summary', 'Parameters', 'EventProfile', 'Trades', 'Equity', 'PoolHistory', 'PoolSignals', 'Bars', 'Checkpoints', 'CheckpointState', 'OpenPositions', 'LongText'} <= set(book.sheetnames)
    assert len(list(book['Bars'].rows)) == 76
    assert len(list(book['Trades'].rows)) == len(result['trades']) + 1
    assert len(list(book['Equity'].rows)) == 16 and len(list(book['Checkpoints'].rows)) == 4
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
    book.close()


def test_embedded_html_xlsx_never_used_to_render(payload):
    files = files_for(payload)
    files['report.html'] = b'<script>fetch("https://attacker.invalid")</script>'
    files['report.xlsx'] = b'not a workbook'
    restored = domain.unpack_portfolio_report(archive(files))
    assert restored == payload and 'attacker.invalid' not in domain.render_portfolio_html(restored)
    assert portfolio_report_xlsx(restored).startswith(b'PK')


@pytest.mark.parametrize('field', ['bar', 'checkpoint', 'chain', 'result', 'input', 'scope', 'cursor', 'completed'])
def test_report_rejects_tampering_and_incomplete_output(payload, field):
    changed = deepcopy(payload)
    if field == 'bar': changed['input']['datasets'][0]['bars'][0]['close'] = '999'
    elif field == 'checkpoint': changed['chunks'][0]['result']['checkpoint']['cash'] = '999999'
    elif field == 'chain': changed['chunks'][1]['prior_sha256'] = 'e'*64
    elif field == 'result': changed['source']['result_sha256'] = 'e'*64
    elif field == 'input': changed['source']['input_sha256'] = 'e'*64
    elif field == 'scope': changed['summary']['universe_scope'] = 'survivorship_bias_free'
    elif field == 'cursor': changed['chunks'][0]['result']['checkpoint']['cursor'] = 9
    else: changed['chunks'].pop()
    with pytest.raises(TradeError):
        domain.validate_portfolio_report(changed)


@pytest.mark.parametrize('kind', ['traversal', 'duplicate_key', 'noncanonical', 'wrong_scope', 'checksum', 'deep_json', 'symlink', 'duplicate_member'])
def test_safe_shared_transport_rejects_hostile_packages(payload, kind):
    files = files_for(payload)
    if kind == 'traversal':
        files['../report.html'] = files.pop('report.html')
    elif kind == 'duplicate_key':
        files['report.json'] = b'{"version":1,"version":1}'
    elif kind == 'noncanonical':
        files['report.json'] = json.dumps(payload, ensure_ascii=False, indent=2).encode()
    elif kind == 'wrong_scope':
        manifest = json.loads(files['manifest.json']); manifest['scope'] = 'single_symbol_backtest'; files['manifest.json'] = domain.encode(manifest)
    elif kind == 'checksum':
        files['report.html'] += b'tamper'
    elif kind == 'deep_json':
        files['report.json'] = b'{"x":' * 40 + b'0' + b'}' * 40
    if kind in ('symlink', 'duplicate_member'):
        out = io.BytesIO()
        with zipfile.ZipFile(out, 'w') as target:
            for name, raw in files.items():
                if name == 'report.html' and kind == 'symlink':
                    item = zipfile.ZipInfo(name); item.create_system = 3; item.external_attr = (stat.S_IFLNK | 0o777) << 16
                    target.writestr(item, raw)
                else: target.writestr(name, raw)
            if kind == 'duplicate_member': target.writestr('report.html', b'duplicate')
        contents = out.getvalue()
    else:
        contents = archive(files, update=kind != 'checksum')
    with pytest.raises(TradeError):
        domain.unpack_portfolio_report(contents)


def test_shared_package_limit_and_formula_long_text_safety(payload, monkeypatch):
    payload['title'] = '=HYPERLINK("https://invalid","label")'
    workbook = portfolio_report_xlsx(payload, metadata={'long_field': '=formula' + '字' * 40000})
    book = load_workbook(io.BytesIO(workbook), read_only=True, data_only=False)
    assert len(list(book['LongText'].rows)) == 3
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
    book.close()
    import trade_app.research.report_domain as shared
    monkeypatch.setattr(shared, 'MAX_PACKAGE_BYTES', 1024)
    with pytest.raises(TradeError) as error:
        domain.package_portfolio_report(payload, workbook)
    assert error.value.code == 'REPORT_PACKAGE_TOO_LARGE'


def test_report_preserves_native_decimal_ratio_precision(payload):
    ratio = '0.01234567890123456789012345678'
    payload['summary']['max_drawdown'] = ratio
    for chunk in payload['chunks']:
        chunk['result']['checkpoint']['max_drawdown'] = ratio
        chunk['result_sha256'] = domain.digest(domain.encode(chunk['result']))
    for index, chunk in enumerate(payload['chunks'][1:], 1):
        chunk['prior_sha256'] = domain.digest(domain.encode(payload['chunks'][index - 1]['result']['checkpoint']))
    payload['source']['result_sha256'] = domain.digest(domain.encode({'summary': payload['summary'],
        'chunks': [chunk['result_sha256'] for chunk in payload['chunks']], 'input_sha256': payload['source']['input_sha256']}))
    assert domain.unpack_portfolio_report(pack(payload))['summary']['max_drawdown'] == ratio


def test_service_metadata_light_list_dedup_soft_delete_and_independent_import(payload, tmp_path):
    with client_for(tmp_path) as client:
        factory = client.app.state.db_factory
        with factory.begin() as session:
            saved = service.import_report(session, pack(payload))
            assert service.import_report(session, pack(payload))['id'] == saved['id']
        statements = []
        def record(_connection, _cursor, statement, *_): statements.append(statement)
        event.listen(factory.kw['bind'], 'before_cursor_execute', record)
        try:
            with factory() as session:
                assert service.list_reports(session)[0]['origin'] == 'import'
        finally:
            event.remove(factory.kw['bind'], 'before_cursor_execute', record)
        assert not any('payload_json' in query for query in statements)
        with factory.begin() as session:
            service.delete_report(session, saved['id'])
            assert service.list_reports(session) == []
            assert session.get(PortfolioReport, saved['id']).payload_json
            again = service.import_report(session, pack(payload))
            assert again['id'] != saved['id'] and again['payload'] == payload


def test_actual_portfolio_to_report_api_and_backup_restore(tmp_path):
    with client_for(tmp_path) as client:
        run, _ = create_run(client, tmp_path, mode='matrix_raw_s1_s9')
        before = write(client, '/research/portfolio-reports', {'portfolio_run_id': run['id'], 'title': '尚未完成'})
        assert before.status_code == 409
        factory = client.app.state.db_factory
        while portfolio_service.process_one_portfolio(factory, tmp_path):
            pass
        body = {'portfolio_run_id': run['id'], 'title': '完整真实组合报告'}
        saved = data(write(client, '/research/portfolio-reports', body, key='report-portfolio'))
        assert data(write(client, '/research/portfolio-reports', body, key='report-portfolio')) == saved
        path = '/api/v1/research/portfolio-reports/' + saved['id']
        zipped = client.get(path + '/export.zip')
        assert zipped.status_code == 200
        assert domain.unpack_portfolio_report(zipped.content) == saved['payload']
        assert client.get(path + '/report.html').status_code == 200
        assert client.get(path + '/export.xlsx').content.startswith(b'PK')
        data(write(client, '/research/portfolios/' + run['id'], {}, method='DELETE'))
        assert data(client.get(path))['payload']['source']['run_id'] == run['id']
        data(write(client, '/research/portfolio-reports/' + saved['id'], {}, method='DELETE'))
        imported = client.post('/api/v1/research/portfolio-reports/import', files={'file': ('report.zip', zipped.content, 'application/zip')},
                              headers={'Idempotency-Key': str(uuid4())})
        imported = data(imported)
        assert imported['origin'] == 'import' and imported['content_sha256'] == saved['content_sha256']
    backup = create_backup(tmp_path)
    restored_dir = tmp_path / 'restored'
    restore_to_new_directory(backup, restored_dir)
    with client_for(restored_dir) as client:
        assert data(client.get('/api/v1/research/portfolio-reports/' + imported['id'])) == imported
