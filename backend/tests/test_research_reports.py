"""Complete report packages, untrusted archives, and database-only portability."""
from copy import deepcopy
import io
import json
import stat
import zipfile
from uuid import uuid4
from pathlib import Path

from openpyxl import load_workbook
import pytest
from sqlalchemy import event, text

from trade_app.api.excel_export import backtest_report_xlsx
from trade_app.market.domain import normalize_bars
from trade_app.platform.backup import create_backup, restore_to_new_directory
from trade_app.platform.db import open_database
from trade_app.platform.types import TradeError, utc_now
from trade_app.research import backtest_domain, report_domain as domain, report_service
from trade_app.research.backtest_service import process_one_backtest
from trade_app.research.domain import DEFAULT_PARAMS
from trade_app.research.report_models import ResearchReport
from test_strategy_backtests import flat_bars, execute, one_signal
from test_wyckoff_research_api import client_for, data, dataset, write


@pytest.fixture
def payload(monkeypatch):
    monkeypatch.setattr(backtest_domain, 'evaluate_strategy', one_signal)
    bars = normalize_bars(flat_bars())
    result = execute(bars)
    now = utc_now()
    run = {'id': 'r' * 64, 'dataset_id': 'd' * 64, 'strategy_id': 'relative_strength_breakout_v1',
           'strategy_version': '1.0.0-alpha', 'execution_version': result['execution_profile_version'],
           'calculation_version': result['calculation_version'], 'code_sha256': 'c' * 64,
           'result_sha256': domain.digest(domain.encode(result)), 'attempt_number': 1, 'params': DEFAULT_PARAMS,
           'config': {'fee_config': {'slippage_rate': '0'}, 'strict': True}, 'state': 'succeeded',
           'created_at': now, 'updated_at': now, 'result': result}
    ds = {'id': run['dataset_id'], 'symbol': result['symbol'], 'provider': 'manual_import', 'adjustment': 'none',
          'first_date': bars[0]['event_date'], 'last_date': bars[-1]['event_date'], 'bar_count': len(bars),
          'availability_quality': 'provided_availability', 'bars': bars}
    return domain.freeze_report(run, ds, title='完整测试报告', created_at=now)


def pack(payload):
    return domain.package_report(payload, backtest_report_xlsx(payload['run']))


def files_for(payload):
    with zipfile.ZipFile(io.BytesIO(pack(payload))) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def rebuild(files, update_manifest=True, entries=None):
    if update_manifest:
        manifest = json.loads(files['manifest.json'])
        for name in set(files) & (domain.PACKAGE_MEMBERS - {'manifest.json'}):
            manifest['files'][name] = {'size': len(files[name]), 'sha256': domain.digest(files[name])}
        files['manifest.json'] = domain.encode(manifest)
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, contents in (entries if entries is not None else files.items()):
            archive.writestr(name, contents)
    return output.getvalue()


def test_package_roundtrip_contains_complete_frozen_bars_excel_slippage_and_offline_html(payload):
    contents = pack(payload)
    imported = domain.unpack_report(contents)
    assert imported == payload
    assert len(imported['dataset']['bars']) == 40
    with zipfile.ZipFile(io.BytesIO(contents)) as archive:
        assert set(archive.namelist()) == domain.PACKAGE_MEMBERS
        manifest = json.loads(archive.read('manifest.json'))
        assert manifest['scope'] == 'single_symbol_backtest'
        for name, spec in manifest['files'].items():
            assert domain.digest(archive.read(name)) == spec['sha256']
        workbook = load_workbook(io.BytesIO(archive.read('report.xlsx')))
        columns = [cell.value for cell in workbook['Trades'][1]]
        assert 'reference_price' in columns and 'slippage_rate' in columns
        assert 'Decisions' in workbook.sheetnames
        document = archive.read('report.html').decode()
        assert 'Content-Security-Policy' in document and '<svg' in document
        assert '<script' not in document and '<link' not in document


def test_embedded_html_is_never_trusted_and_json_content_is_escaped(payload):
    malicious = deepcopy(payload)
    malicious['title'] = '<img src=x onerror=alert(1)>'
    files = files_for(malicious)
    files['report.html'] = b'<script>window.stolen=true</script>'
    files['report.xlsx'] = b'untrusted workbook bytes are never parsed'
    imported = domain.unpack_report(rebuild(files))
    regenerated = domain.render_html(imported, origin='import')
    assert '<script>' not in regenerated and 'window.stolen' not in regenerated
    assert '<img src=x' not in regenerated and '&lt;img src=x onerror=alert(1)&gt;' in regenerated
    assert '导入报告，未重新计算' in regenerated
    workbook = load_workbook(io.BytesIO(backtest_report_xlsx(imported['run'])))
    assert workbook.sheetnames[0] == 'Summary'


@pytest.mark.parametrize('case', ['extra', 'traversal', 'absolute', 'duplicate', 'symlink', 'large', 'checksum'])
def test_archive_whitelist_size_and_integrity_rejections(payload, case):
    files = files_for(payload)
    entries = None
    if case in ('extra', 'traversal', 'absolute'):
        files[{'extra': 'readme.txt', 'traversal': '../report.json', 'absolute': '/report.json'}[case]] = b'x'
    elif case == 'duplicate':
        entries = list(files.items()) + [('report.json', files['report.json'])]
    elif case == 'symlink':
        link = zipfile.ZipInfo('report.html')
        link.create_system = 3
        link.external_attr = (stat.S_IFLNK | 0o777) << 16
        entries = [(link if name == 'report.html' else name, content) for name, content in files.items()]
    elif case == 'large':
        files['report.html'] = b'x' * (domain.MAX_PACKAGE_BYTES + 1)
    else:
        files['report.json'] += b' '
    with pytest.raises(TradeError) as failure:
        domain.unpack_report(rebuild(files, update_manifest=case != 'checksum', entries=entries))
    assert failure.value.code in ('INVALID_REPORT_PACKAGE', 'REPORT_PACKAGE_TOO_LARGE')


@pytest.mark.parametrize('raw', [b'{"version":1,"version":1}', b'{"x":NaN}', b'{"x":1e9999}',
    b'{"x":' + b'[' * 40 + b'0' + b']' * 40 + b'}', b'PK\x03\x04not-json', b'{"x":Infinity}'])
def test_strict_json_rejects_duplicates_nonfinite_nested_and_binary(raw):
    with pytest.raises(TradeError) as failure:
        domain.strict_json(raw)
    assert failure.value.code == 'INVALID_REPORT_JSON'


@pytest.mark.parametrize('change', ['legacy_scope', 'wrong_version', 'bars', 'result_digest', 'huge_number'])
def test_report_rejects_unknown_formats_mismatched_frozen_data_and_unbounded_numbers(payload, change):
    value = deepcopy(payload)
    if change == 'legacy_scope':
        value['scope'] = 'legacy_matrix'
    elif change == 'wrong_version':
        value['version'] = 2
    elif change == 'bars':
        value['dataset']['bars'][0]['volume'] += 1
    elif change == 'result_digest':
        value['run']['result']['ending_assets'] = '999999.00'
    else:
        value['run']['result']['total_return'] = '1e9999999'
        value['run']['result_sha256'] = domain.digest(domain.encode(value['run']['result']))
    with pytest.raises(TradeError) as failure:
        domain.validate_report(value)
    assert failure.value.code == 'INVALID_REPORT_DATA'


def test_manifest_and_report_duplicate_keys_are_rejected_even_with_updated_outer_checksums(payload):
    files = files_for(payload)
    files['report.json'] = b'{"format":"x","format":"x"}'
    with pytest.raises(TradeError) as failure:
        domain.unpack_report(rebuild(files))
    assert failure.value.code == 'INVALID_REPORT_JSON'
    files = files_for(payload)
    files['manifest.json'] = b'{"version":1,"version":1}'
    with pytest.raises(TradeError) as failure:
        domain.unpack_report(rebuild(files, update_manifest=False))
    assert failure.value.code == 'INVALID_REPORT_JSON'


def test_freeze_checks_existing_result_digest_and_allows_older_unhashed_completed_job(payload):
    run, ds = deepcopy(payload['run']), payload['dataset']
    run['result_sha256'] = None
    frozen = domain.freeze_report(run, ds, title='早期完成结果', created_at=utc_now())
    assert frozen['run']['result_sha256'] == domain.digest(domain.encode(run['result']))
    run['result_sha256'] = '0' * 64
    with pytest.raises(TradeError) as failure:
        domain.freeze_report(run, ds, title='破损结果', created_at=utc_now())
    assert failure.value.code == 'BACKTEST_RESULT_CORRUPT'


def test_report_directory_reads_small_metadata_only_and_legacy_metadata_backfills_once(tmp_path, payload, monkeypatch):
    engine, factory = open_database(tmp_path)
    try:
        with factory.begin() as session:
            report = report_service._store(session, payload, 'import')
            row = session.get(ResearchReport, report['id'])
            row.metadata_json = None
            session.flush()
            minimal = report_service.list_reports(session)[0]
            assert minimal['metadata_missing'] is True
            migration = Path(__file__).parents[1] / 'trade_app/migrations/0041_report_metadata.sql'
            session.execute(text(migration.read_text().split(';', 1)[1].strip().rstrip(';')))
            session.expire(row, ['metadata_json'])
            assert json.loads(row.metadata_json)['summary']['quality_flags'] == payload['run']['result']['quality_flags']
        statements = []
        event.listen(engine, 'before_cursor_execute', lambda _conn, _cursor, statement, *_: statements.append(statement))
        monkeypatch.setattr(report_service, '_payload', lambda *_: pytest.fail('List must not decode full report bodies'))
        with factory() as session:
            listing = report_service.list_reports(session)
        assert listing[0]['summary'] == report['summary']
        assert all('payload_json' not in statement for statement in statements)
    finally:
        engine.dispose()


def upload(client, contents, key=None):
    return client.post('/api/v1/research/reports/import', files={'file': ('report.zip', contents, 'application/zip')},
                       headers={'Idempotency-Key': key or str(uuid4())})


def test_api_freeze_export_import_without_source_then_backup_restore_and_delete(tmp_path, monkeypatch):
    original_dir = tmp_path / 'original'
    with client_for(original_dir) as client:
        ds = dataset(client, flat_bars())
        run = data(write(client, '/backtests', {'dataset_id': ds['id'], 'strategy_id': 'relative_strength_breakout_v1'}))
        before = write(client, '/research/reports', {'backtest_run_id': run['id'], 'title': '未完成报告'})
        assert before.status_code == 409 and before.json()['error']['code'] == 'BACKTEST_NOT_READY'
        assert process_one_backtest(client.app.state.db_factory, original_dir)
        report = data(write(client, '/research/reports', {'backtest_run_id': run['id'], 'title': '实际冻结报告'}, key='report-once'))
        assert data(write(client, '/research/reports', {'backtest_run_id': run['id'], 'title': '实际冻结报告'}, key='report-once')) == report
        contents = client.get(f'/api/v1/research/reports/{report["id"]}/export.zip')
        assert contents.status_code == 200
        package = contents.content
        assert domain.unpack_report(package) == report['payload']
        assert client.get(f'/api/v1/research/reports/{report["id"]}/export.xlsx').status_code == 200
        html = client.get(f'/api/v1/research/reports/{report["id"]}/report.html')
        assert html.status_code == 200 and "default-src 'none'" in html.headers['content-security-policy']
        assert data(client.get('/api/v1/research/reports'))[0]['id'] == report['id']

    imported_dir = tmp_path / 'imported'
    def never_compute(*_, **__):
        pytest.fail('Import must not read or run source backtests')
    monkeypatch.setattr(report_service, 'get_backtest', never_compute)
    monkeypatch.setattr(report_service, 'get_dataset', never_compute)
    with client_for(imported_dir) as client:
        imported = data(upload(client, package))
        assert imported['origin'] == 'import' and imported['payload'] == report['payload']
        assert data(upload(client, package))['id'] == imported['id']
        assert data(client.get('/api/v1/backtests')) == []
        assert data(client.get('/api/v1/market/datasets')) == []
        assert '导入报告，未重新计算' in client.get(f'/api/v1/research/reports/{imported["id"]}/report.html').text
        regenerated = client.get(f'/api/v1/research/reports/{imported["id"]}/export.zip')
        with zipfile.ZipFile(io.BytesIO(regenerated.content)) as archive:
            assert '导入报告，未重新计算' in archive.read('report.html').decode()
            workbook = load_workbook(io.BytesIO(archive.read('report.xlsx')))
            summary = {row[0].value: row[1].value for row in workbook['Summary'].iter_rows(min_row=2)}
            assert summary['report_origin'] == '导入，未重新计算'
    backup = create_backup(imported_dir)
    restored_dir = tmp_path / 'restored'
    restore_to_new_directory(backup, restored_dir)
    with client_for(restored_dir) as client:
        assert data(client.get('/api/v1/research/reports/' + imported['id'])) == imported
        deleted = data(write(client, '/research/reports/' + imported['id'], {}, 'DELETE'))
        assert deleted['deleted'] is True
        assert client.get('/api/v1/research/reports/' + imported['id']).status_code == 404
        assert data(client.get('/api/v1/research/reports')) == []
        assert data(upload(client, package))['id'] != imported['id']
