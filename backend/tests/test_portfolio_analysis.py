from copy import deepcopy
from datetime import date, timedelta
from decimal import Decimal
import math

import pytest

from trade_app.platform.compute_process import ComputeProcessError
from trade_app.platform.types import TradeError
from trade_app.research import portfolio_analysis_domain as domain, portfolio_analysis_service as service
from trade_app.research import portfolio_service as portfolio
from trade_app.research.portfolio_analysis_models import PortfolioAnalysis
from trade_app.research import portfolio_domain, portfolio_plan_domain as plans
from test_portfolio import context, create_run, signals
from test_wyckoff_research_api import client_for, data, write


def sample():
    ctx = context(symbols=('sh600000', 'sh600001'), count=100)
    result = {'initial_capital': '10000', 'buy_count': 3, 'quality_flags': ['selection_membership_unverified'],
        'equity': [{'date': (date(2025, 1, 1) + timedelta(days=i)).isoformat(), 'total_assets': str(10000 + i * 10)} for i in range(65)],
        'decisions': [{'reason': 'MAX_POSITIONS'}, {'reason': 'DAILY_TOP_K'}, {'reason': 'POOL_OR_RISK_INVALIDATION'}],
        'trades': [
            {'symbol': 'sh600000', 'side': 'buy', 'quantity': 100, 'price': '10', 'fees': '5', 'date': '2025-02-01'},
            {'symbol': 'sh600001', 'side': 'buy', 'quantity': 100, 'price': '10', 'fees': '5', 'date': '2025-02-02'},
            {'symbol': 'sh600000', 'side': 'sell', 'quantity': 50, 'realized_pnl': '100', 'date': '2025-02-03'},
            {'symbol': 'sh600001', 'side': 'sell', 'quantity': 100, 'realized_pnl': '-50', 'date': '2025-02-04'},
            {'symbol': 'sh600000', 'side': 'sell', 'quantity': 50, 'realized_pnl': '-20', 'date': '2025-02-05'},
            {'symbol': 'sh600001', 'side': 'buy', 'quantity': 100, 'price': '10', 'fees': '5', 'date': '2025-02-06'},
            {'symbol': 'sh600001', 'side': 'sell', 'quantity': 50, 'realized_pnl': '500', 'date': '2025-02-07'}]}
    return ctx, result


def test_partial_exits_count_one_completed_cycle_and_keep_open_cycle_separate():
    ctx, result = sample()
    actual = domain.build_analysis(ctx, result, iterations=100)
    assert actual['risk']['completed_trade_count'] == 2
    assert actual['risk']['profit_factor'] == 1.6
    assert actual['risk']['expectancy'] == pytest.approx((80 / 1005 - 50 / 1005) / 2)
    assert actual['risk']['max_concurrent_positions'] == 2
    assert actual['risk']['fill_rate'] == .6
    assert actual['open_cycles'][0]['remaining_quantity'] == 50
    assert actual['open_cycles'][0]['net_pnl'] == '500'
    assert [row['net_pnl'] for row in actual['completed_trades']] == ['-50', '80']
    assert actual['stability']['stability_score'] is None
    assert actual['stability']['neighborhood']['status'] == 'not_generated'


@pytest.mark.parametrize('price, fees, expected_cost', [
    ('10.00004', '0', '1000.00'),
    ('10.00005', '0', '1000.01'),
    ('10.00006', '5.01', '1005.02'),
])
def test_cycle_cost_matches_ledger_cents_before_fees_across_partial_exits(price, fees, expected_cost):
    from trade_app.trading.domain import FeeBreakdown, FillState, apply_fill
    bought = apply_fill(FillState(Decimal('10000'), 0, 0, Decimal(0)), side='buy',
        price=Decimal(price), quantity=100, fees=FeeBreakdown(Decimal(fees), Decimal(0), Decimal(0)))
    assert str(bought.cost_basis) == expected_cost
    trades = [
        {'symbol': 'sh600000', 'side': 'buy', 'date': '2025-01-01', 'quantity': 100, 'price': price, 'fees': fees},
        {'symbol': 'sh600000', 'side': 'sell', 'date': '2025-01-02', 'quantity': 50, 'realized_pnl': '100.01'},
    ]
    completed, open_cycles, peak = domain.complete_cycles(trades)
    assert not completed and peak == 1
    assert open_cycles[0]['entry_cost'] == expected_cost
    assert open_cycles[0]['remaining_quantity'] == 50
    trades.append({'symbol': 'sh600000', 'side': 'sell', 'date': '2025-01-03', 'quantity': 50, 'realized_pnl': '179.99'})
    completed, open_cycles, _ = domain.complete_cycles(trades)
    assert not open_cycles and len(completed) == 1
    assert completed[0]['entry_cost'] == expected_cost
    assert completed[0]['net_pnl'] == '280.00' and completed[0]['exit_legs'] == 2
    assert completed[0]['position_return'] == round(float(Decimal('280') / bought.cost_basis), 10)


def test_risk_includes_initial_capital_loss_and_recovery_is_explicit():
    _, result = sample()
    result.update(buy_count=0, trades=[], decisions=[], equity=[
        {'date': '2025-01-01', 'total_assets': '9000'}, {'date': '2025-01-03', 'total_assets': '8000'},
        {'date': '2025-01-04', 'total_assets': '9500'}, {'date': '2025-01-06', 'total_assets': '10000'}])
    daily = domain.daily_returns(result)
    assert daily[0] == pytest.approx(-.1)
    metrics = domain.risk_metrics(result, [], daily, 0)
    assert metrics['max_drawdown'] == .2 and metrics['annualized_return'] == 0
    assert metrics['recovery']['calendar_days'] == 3 and metrics['recovery']['trading_sessions'] == 2
    assert metrics['profit_factor'] is None and metrics['expectancy'] is None
    result['equity'].pop()
    metrics = domain.risk_metrics(result, [], domain.daily_returns(result), 0)
    assert metrics['recovery']['status'] == 'unrecovered' and metrics['recovery']['calendar_days'] is None


def test_seeded_bootstrap_uses_daily_aggregate_and_reports_insufficient_and_extreme_samples():
    daily = [.02, -.012, .01, -.015, .003] * 12
    kwargs = {'seed': 8, 'iterations': 100, 'block_size': 4}
    result = domain.monte_carlo(daily, **kwargs)
    assert result == domain.monte_carlo(daily, **kwargs)
    assert result != domain.monte_carlo(daily, **{**kwargs, 'seed': 9})
    assert result['total_return']['p5'] <= result['total_return']['p50'] <= result['total_return']['p95']
    assert result['max_drawdown']['p95'] >= 0 and result['ruin_probability'] == 0
    insufficient = domain.monte_carlo(daily[:29], **kwargs)
    assert insufficient['status'] == 'insufficient_sample' and insufficient['total_return'] is None and insufficient['iterations'] == 0
    ruin = domain.monte_carlo([-1] * 30, seed=2, iterations=100, block_size=1)
    assert ruin['ruin_probability'] == 1 and ruin['total_return']['p95'] == -1
    extreme = domain.monte_carlo([1e20] * 30, seed=2, iterations=100, block_size=1)
    assert extreme['total_return'] is None and extreme['numerical_overflow_paths'] == 100


def test_regime_uses_entry_known_prefix_and_never_later_confirmation():
    ctx, result = sample()
    cycles, _, _ = domain.complete_cycles(result['trades'])
    baseline, annotated = domain.regimes(ctx, cycles, result['initial_capital'])
    changed = deepcopy(ctx)
    for dataset in changed['datasets']:
        for bar in dataset['bars']:
            if bar['event_date'] >= '2025-02-01': bar['close'] = '999999'
    assert domain.regimes(changed, cycles, result['initial_capital'])[1][1] == annotated[1]
    # A missing old observation cannot silently compress strict classification.
    ctx['datasets'][0]['bars'][2]['available_at'] = '2030-01-01T00:00:00+00:00'
    assert domain.regimes(ctx, cycles, result['initial_capital'])[1][1]['regime'] == 'unknown'


@pytest.mark.parametrize('options', [{'seed': True}, {'iterations': 99}, {'iterations': 2001}, {'block_size': 0}, {'seed': 2**32}])
def test_options_are_explicit_bounded_and_not_coerced(options):
    ctx, result = sample()
    with pytest.raises(ValueError): domain.build_analysis(ctx, result, **options)


def prepare(client, tmp_path):
    source, _ = create_run(client, tmp_path, count=95)
    factory = client.app.state.db_factory
    while portfolio.process_one_portfolio(factory, tmp_path): pass
    with factory.begin() as session:
        created = service.create_analysis(session, {'source_run_id': source['id'], 'iterations': 100})
    return factory, source, created


def test_actual_process_persists_digest_survives_source_deletion_and_detects_corruption(tmp_path):
    with client_for(tmp_path) as client:
        factory, source, created = prepare(client, tmp_path)
        with factory.begin() as session: portfolio.control_portfolio(session, source['id'], 'delete')
        assert service.process_one_analysis(factory, tmp_path)
        with factory.begin() as session:
            result = service.get_analysis(session, created['id'])
            assert result['state'] == 'succeeded', result
            assert result['result']['monte_carlo']['status'] == 'generated'
            assert result['result']['daily_plan']['status'] == 'generated'
            assert result['compute_metrics']['pid'] > 0
            assert len(result['result_sha256']) == 64
            assert 'result' not in service.list_analyses(session)[0]
            row = session.get(PortfolioAnalysis, created['id'])
            row.result_json = row.result_json.replace('generated', 'tampered', 1)
        with factory() as session, pytest.raises(TradeError) as error:
            service.get_analysis(session, created['id'])
        assert error.value.code == 'PORTFOLIO_ANALYSIS_CORRUPT'


def test_shutdown_recovery_and_retry_reject_changed_calculator(tmp_path, monkeypatch):
    with client_for(tmp_path) as client:
        factory, _, created = prepare(client, tmp_path)
        service.recover_interrupted_analyses(factory)
        with factory.begin() as session:
            assert service.get_analysis(session, created['id'])['state'] == 'paused'
            service.control_analysis(session, created['id'], 'resume')
        def stopped(*args, **kwargs): raise ComputeProcessError('COMPUTE_SHUTDOWN', 'stopped')
        monkeypatch.setattr(service, 'run_json_process', stopped)
        assert service.process_one_analysis(factory, tmp_path)
        with factory.begin() as session:
            actual = service.get_analysis(session, created['id'])
            assert actual['state'] == 'paused' and actual['result_sha256'] is None
            monkeypatch.setattr(service, 'code_sha256', lambda: '0' * 64)
            with pytest.raises(TradeError) as error: service.control_analysis(session, created['id'], 'resume')
            assert error.value.code == 'CODE_VERSION_CHANGED'


def test_conditional_plan_is_causal_no_next_quotes_and_matches_partial_exit_rule(monkeypatch):
    ctx = context(count=8, trailing_stop_pct='.05', intraday_trailing=True, trailing_reduce_ratio='.5')
    signals(monkeypatch)
    monkeypatch.setattr(plans, 'evaluate_strategy', portfolio_domain.evaluate_strategy)
    # Buy at 10 on index3, then observe its low/close after that entry.
    ctx['datasets'][0]['bars'][3].update(low='8.9', close='9')
    state = portfolio_domain.run_chunk(ctx, days=1)['checkpoint']
    day = ctx['calendar'][0]
    plan = plans.build_plan(ctx, state, day)
    holding = plan['holdings'][0]
    assert holding['sellable_at_cutoff'] == 0 and holding['sellable_next_session_if_open'] == 1000
    exit_plan = plan['open_signals'][0]
    assert exit_plan['reason'] == 'OBSERVED_TRAILING_NEXT_OPEN' and exit_plan['quantity_if_executable'] == 500
    assert exit_plan['price'] is None and plan['target_session_date'] is None
    original = deepcopy(ctx)
    for bar in original['datasets'][0]['bars'][4:]:
        bar.update(open='1234', close='1999', high='3000', low='1', volume=1)
    assert plans.build_plan(original, state, day) == plan
    actual_exit = portfolio_domain.run_chunk(ctx, state, days=1)['trades'][0]
    assert actual_exit['reason'] == exit_plan['reason'] and actual_exit['quantity'] == exit_plan['quantity_if_executable']
    ctx['datasets'][0]['bars'][3]['available_at'] = day + 'T20:00:00+00:00'
    late = plans.build_plan(ctx, state, day)
    assert not late['open_signals'] and 'unavailable_or_stale_prior_history' in late['quality_flags']


def test_weekly_candidate_and_delayed_intent_do_not_assume_a_future_session(monkeypatch):
    ctx = context(count=8, entry_delay_bars=3, pool_roll='weekly')
    signals(monkeypatch, lambda symbol, bars: True)
    monkeypatch.setattr(plans, 'evaluate_strategy', portfolio_domain.evaluate_strategy)
    state = portfolio_domain.run_chunk(ctx, days=1)['checkpoint']
    plan = plans.build_plan(ctx, state, ctx['calendar'][0])
    assert plan['weekly_refresh_date_unknown']
    assert plan['plan_signals'][0]['status'] == 'waiting_observed_bars'
    assert plan['plan_signals'][0]['quantity_if_executable'] is None
    assert not plan['open_signals']
    assert plan['plan_signals'][0]['remaining_observed_bars'] == 1


def test_v2_report_freezes_analysis_and_keeps_v1_readable(tmp_path):
    from trade_app.research import portfolio_report_service as reports
    from trade_app.research.portfolio_report_domain import validate_portfolio_report, render_portfolio_html, package_portfolio_report, unpack_portfolio_report
    from trade_app.api.portfolio_report_excel import portfolio_report_xlsx
    from openpyxl import load_workbook
    from io import BytesIO
    with client_for(tmp_path) as client:
        factory, source, created = prepare(client, tmp_path)
        assert service.process_one_analysis(factory, tmp_path)
        with factory.begin() as session:
            v1 = reports.create_report(session, {'portfolio_run_id': source['id'], 'title': '旧协议保留'})
            v2 = reports.create_report(session, {'portfolio_run_id': source['id'], 'title': '附加分析', 'analysis_run_id': created['id']})
            assert v1['payload']['version'] == 1 and v2['payload']['version'] == 2
            assert validate_portfolio_report(v1['payload']) == v1['payload']
            assert v2['payload']['analysis']['result']['daily_plan']['status'] == 'generated'
            service.control_analysis(session, created['id'], 'delete')
        workbook = portfolio_report_xlsx(v2['payload'])
        assert 'ConditionalPlans' in load_workbook(BytesIO(workbook), read_only=True).sheetnames
        assert 'Bootstrap' in render_portfolio_html(v2['payload'])
        package = package_portfolio_report(v2['payload'], workbook)
        assert unpack_portfolio_report(package) == v2['payload']
        with factory() as session: assert reports.get_report(session, v2['id']) == v2
        # Stored pre-rounding analysis reports stay readable; never recalculate
        # their frozen numeric evidence using the new analysis revision.
        legacy = deepcopy(v2['payload'])
        legacy['analysis']['version'] = 'portfolio-complete-cycles-daily-block-bootstrap-v1'
        legacy['analysis']['result']['version'] = legacy['analysis']['version']
        legacy['analysis']['result_sha256'] = portfolio_domain.digest(legacy['analysis']['result'])
        assert validate_portfolio_report(legacy) == legacy
        assert unpack_portfolio_report(package_portfolio_report(legacy, portfolio_report_xlsx(legacy))) == legacy
        corrupt = deepcopy(v2['payload'])
        corrupt['analysis']['source_result_sha256'] = '0' * 64
        with pytest.raises(TradeError): validate_portfolio_report(corrupt)
        corrupt = deepcopy(v2['payload'])
        corrupt['analysis']['result']['monte_carlo']['seed'] = 7
        with pytest.raises(TradeError): validate_portfolio_report(corrupt)


def test_api_idempotence_date_validation_cancel_and_deep_result(tmp_path):
    with client_for(tmp_path) as client:
        factory, source, first = prepare(client, tmp_path)
        with factory.begin() as session: service.control_analysis(session, first['id'], 'cancel')
        body = {'source_run_id': source['id'], 'name': '真实分析API', 'seed': 17, 'iterations': 100, 'block_size': 3}
        created = data(write(client, '/research/portfolio-analyses', body, key='analysis-api-idempotency'))
        assert data(write(client, '/research/portfolio-analyses', body, key='analysis-api-idempotency')) == created
        invalid = write(client, '/research/portfolio-analyses', {**body, 'plan_date': '1999-01-01'})
        assert invalid.status_code == 400 and invalid.json()['error']['code'] == 'INVALID_ANALYSIS_DATE'
        assert write(client, '/research/portfolio-analyses', {**body, 'iterations': True}).status_code == 422
        assert service.process_one_analysis(factory)
        actual = data(client.get('/api/v1/research/portfolio-analyses/' + created['id']))
        assert actual['state'] == 'succeeded' and actual['result']['monte_carlo']['seed'] == 17
        exported = client.get('/api/v1/research/portfolio-analyses/' + created['id'] + '/export.json')
        assert exported.status_code == 200 and exported.json()['result_sha256'] == actual['result_sha256']
        assert write(client, '/research/portfolio-analyses/' + created['id'] + '/cancel', {}).status_code == 409
        assert len(data(client.get('/api/v1/research/portfolio-analyses?source_run_id=' + source['id']))) == 2
        data(write(client, '/research/portfolio-analyses/' + created['id'], {}, method='DELETE'))
        assert client.get('/api/v1/research/portfolio-analyses/' + created['id']).status_code == 404


def test_old_execution_version_allows_statistics_but_cannot_regenerate_plan(tmp_path, monkeypatch):
    with client_for(tmp_path) as client:
        factory, source, first = prepare(client, tmp_path)
        with factory.begin() as session: service.control_analysis(session, first['id'], 'cancel')
        monkeypatch.setattr(service, 'execution_code_sha256', lambda: 'f' * 64)
        with factory.begin() as session:
            created = service.create_analysis(session, {'source_run_id': source['id'], 'iterations': 100})
        assert service.process_one_analysis(factory)
        with factory() as session:
            actual = service.get_analysis(session, created['id'])
        assert actual['state'] == 'succeeded'
        assert actual['result']['risk']['daily_sample_count'] >= 30
        assert actual['result']['daily_plan']['status'] == 'not_generated'


def test_selected_day_checkpoint_replay_stops_at_exact_day(tmp_path):
    import json
    with client_for(tmp_path) as client:
        factory, source, first = prepare(client, tmp_path)
        with factory.begin() as session:
            service.control_analysis(session, first['id'], 'cancel')
            portfolio_result = portfolio.get_portfolio(session, source['id'], full=True)
            day = portfolio_result['result']['equity'][6]['date']
            created = service.create_analysis(session, {'source_run_id': source['id'], 'iterations': 100, 'plan_date': day})
            frozen = json.loads(session.get(PortfolioAnalysis, created['id']).input_json)
            assert frozen['plan']['checkpoint']['cursor'] == 5 and frozen['plan']['replay_days'] == 2
        assert service.process_one_analysis(factory)
        with factory() as session: actual = service.get_analysis(session, created['id'])
        assert actual['state'] == 'succeeded'
        assert actual['result']['daily_plan']['as_of_date'] == day
        assert actual['result']['daily_plan']['cash'] == portfolio_result['result']['equity'][6]['cash']
        assert actual['result']['daily_detail']['equity'] == portfolio_result['result']['equity'][6]
        assert all(row['source_date'] <= day for row in actual['result']['daily_plan']['plan_signals'])
