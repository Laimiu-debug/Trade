from pathlib import Path
import asyncio
import csv
import hashlib
import io
import json
import struct
import zipfile
from datetime import date, datetime, timedelta, timezone
from math import sin
from types import SimpleNamespace
from contextlib import contextmanager
from decimal import Decimal
from dataclasses import asdict
from PIL import Image
from openpyxl import load_workbook

from fastapi.testclient import TestClient
from sqlalchemy import text

from trade_app.analytics.service import process_one
from trade_app.analytics.domain import project_account
from trade_app.research.backtest_service import process_one_backtest
from trade_app.research.backtest_models import BacktestRun
from trade_app.main import create_app
from trade_app.platform.backup import create_backup, restore_to_new_directory, verify_backup
from trade_app.platform.instance_lock import InstanceLock
from trade_app.reviews.periods import bounds
from trade_app.trading.nav import FlowFact, SnapshotFact, calculate_nav
from trade_app.trading.domain import calculate_fees
from trade_app.trading.simulation import DEFAULT_CONFIG, fee_rule
from trade_app.trading.service import create_trade
from trade_app.platform.types import TradeError, price_units
from trade_app.market.domain import eligible_bars
from trade_app.market.tdx import normalize_tdx_symbol
from trade_app.market import akshare_online
from trade_app.market import baostock_online
from trade_app.market.sync_jobs import process_one_sync_symbol
from trade_app.market import sync_jobs
from trade_app.market.models import MarketSyncJob
from trade_app.research.screener_domain import FunnelConfig, ScreenerCandidate, run_funnel
from trade_app.research.screener_metrics import build_candidate
from trade_app.research.tdx_universe import process_one_universe_symbol
from trade_app.research.tdx_universe_models import TdxUniverseJob
from trade_app.research.b1_domain import B1Params as NewB1Params, check_b1 as new_check_b1
from app.core.b1_strategy import B1Params as OldB1Params, check_b1 as old_check_b1
from app.core.market_momentum import _compute_daily_top_symbols as legacy_trend_top
from app.core.market_momentum import count_consecutive_limit_up as legacy_limit_height
from app.core import sector_capital_flow as legacy_sector_flow
from app.core.abnormal_movement import analyze_symbol_abnormal_events as legacy_analyze_abnormal
from trade_app.research.abnormal_domain import analyze_symbol_abnormal_events as new_analyze_abnormal
from app.core.sentiment_valuation import ValuationInputs as OldValuationInputs, calc_theoretical_cap as old_valuation_calc
from trade_app.research.valuation_domain import ValuationInputs as NewValuationInputs, calc_theoretical_cap as new_valuation_calc
from trade_app.research import valuation_service
from app.core.sector_analyzer import SECTOR_CODES as legacy_sector_codes
from trade_app.research.sector_codes import SECTOR_CODES as new_sector_codes
from trade_app.research.sector_service import compute_sector_flow
from app.models import ScreenerResult, ScreenerStepConfigs
from app.store import InMemoryStore
from app.tdx_loader import _build_row as legacy_build_screener_row
from app.core.strategy_plugins import RelativeStrengthBreakoutPlugin
from app.core.trend_king_strategy import calculate_trend_king_signal as legacy_trend_king, evaluate_trend_king_signal as legacy_trend_king_evaluate
from app.models import CandlePoint as LegacyCandlePoint
from trade_app.research.trend_king_domain import CandlePoint as NewTrendKingCandlePoint, calculate_trend_king_signal as new_trend_king, evaluate_trend_king_signal as new_trend_king_evaluate
from app.core.limit_up_arb_strategy import calculate_limit_up_arb_signal as legacy_limit_up_arb, evaluate_limit_up_arb_signal as legacy_limit_up_arb_evaluate
from trade_app.research.limit_up_arb_domain import calculate_limit_up_arb_signal as new_limit_up_arb, evaluate_limit_up_arb_signal as new_limit_up_arb_evaluate
from app.core.ths_volume_signal import calculate_ths_main_retail_signal as legacy_ths_signal
from trade_app.research.ths_volume_domain import calculate_ths_main_retail_signal as new_ths_signal
from app.core.emotion_limit_up_strategy import calculate_emotion_limit_up_signal as legacy_emotion_signal, evaluate_emotion_limit_up_signal as legacy_emotion_evaluate
from trade_app.research.emotion_limit_up_domain import calculate_emotion_limit_up_signal as new_emotion_signal, evaluate_emotion_limit_up_signal as new_emotion_evaluate
from app.core.strategy_plugins import calculate_wulong_cluster_signal as legacy_wulong_signal, evaluate_wulong_cluster_signal as legacy_wulong_evaluate
from trade_app.research.wulong_domain import calculate_wulong_cluster_signal as new_wulong_signal, evaluate_wulong_cluster_signal as new_wulong_evaluate
from app.core.force_rhythm_strategy import calculate_force_rhythm_signal as legacy_rhythm_signal, evaluate_force_rhythm_signal as legacy_rhythm_evaluate
from trade_app.research.force_rhythm_domain import calculate_force_rhythm_signal as new_rhythm_signal, evaluate_force_rhythm_signal as new_rhythm_evaluate
from trade_app.research.domain import normalize_params as normalize_relative_strength_params


def test_relative_strength_params_follow_legacy_catalog_bounds() -> None:
    assert normalize_relative_strength_params({'min_vol_slope20': '-0.25'})['min_vol_slope20'] == '-0.25'
    for key, value in [('min_ret40', '1.51'), ('max_retrace20', '0.61'),
                       ('min_up_down_volume_ratio', '3.01'), ('min_vol_slope20', '-0.51'),
                       ('min_ai_confidence', '1.01')]:
        try:
            normalize_relative_strength_params({key: value})
        except TradeError as exc:
            assert exc.code == 'INVALID_STRATEGY_PARAM'
        else:
            raise AssertionError(f'{key} accepted an out-of-range value')


def test_trend_king_four_modes_match_legacy_on_frozen_bars(tmp_path: Path) -> None:
    bars = []
    for index in range(140):
        day = date(2025, 1, 1) + timedelta(days=index)
        close = 10 + index * 0.08
        if index == 139:
            close *= 1.10
        bars.append({'event_date': day.isoformat(), 'open': f'{close * 0.99:.4f}',
                     'high': f'{close * 1.02:.4f}', 'low': f'{close * 0.98:.4f}',
                     'close': f'{close:.4f}', 'volume': 1000 + index * 2,
                     'available_at': datetime.combine(day + timedelta(days=1),
                                                       datetime.min.time(), timezone.utc).isoformat()})
    old_points = [LegacyCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                    high=float(bar['high']), low=float(bar['low']),
                                    close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
    new_points = [NewTrendKingCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                          high=float(bar['high']), low=float(bar['low']),
                                          close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
    expected_indicator = legacy_trend_king(old_points)
    assert new_trend_king(new_points) == expected_indicator
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        catalog = client.get('/api/v1/research/strategies').json()['data']
        assert sum(row['status'] == 'trend_king_signal' for row in catalog) == 4
        assert sum(row['status'] == 'not_migrated' for row in catalog) == 0
        dataset_response = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'bars': bars, 'adjustment': 'none'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'tk-market'})
        assert dataset_response.status_code == 200, dataset_response.text
        dataset_id = dataset_response.json()['data']['id']
        for mode, strategy_id in [('all', 'trend_king_v1'), ('a', 'trend_king_limitup_v1'),
                                  ('b', 'trend_king_rally_v1'), ('c', 'trend_king_pullback_v1')]:
            body = {'dataset_id': dataset_id, 'strategy_id': strategy_id,
                    'decision_at': '2025-05-22T12:00:00+00:00', 'strict': True,
                    'params': {'min_hist': '15'}}
            response = client.post('/api/v1/research/runs', json=body,
                                   headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'tk-{mode}'})
            assert response.status_code == 200, response.text
            run = response.json()['data']
            assert run['result']['indicator'] == expected_indicator
            assert run['result']['evaluation'] == legacy_trend_king_evaluate(
                expected_indicator, {'mode': mode, 'min_hist': '15'})
            assert run['result']['signal'] == run['result']['evaluation']['signal']
            assert 'sector_context_unavailable' in run['result']['quality_flags']
            assert len(run['result']['code_sha256']) == 64
            assert run['result']['calculation_version'] == 'legacy-trend-king-frozen-v1'
            assert client.get('/api/v1/research/runs/' + run['id']).json()['data'] == run
        earlier = client.post('/api/v1/research/runs', json={
            **body, 'strategy_id': 'trend_king_v1', 'decision_at': '2025-05-20T12:00:00+00:00'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'tk-before-last-available'})
        assert earlier.status_code == 200, earlier.text
        assert earlier.json()['data']['result']['indicator'] == legacy_trend_king(old_points[:-1])
        bad = client.post('/api/v1/research/runs', json={**body, 'params': {'mode': 'all'}},
                          headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'tk-bad'})
        assert bad.status_code == 400
        invalid = client.post('/api/v1/research/runs', json={**body, 'params': {'min_hist': '1000'}},
                              headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'tk-invalid'})
        assert invalid.status_code == 400


def test_limit_up_arb_signal_matches_legacy_on_frozen_bars(tmp_path: Path) -> None:
    closes = [10, 10, 10, 10, 11, 11.44]
    volumes = [1000, 1000, 1000, 1000, 1000, 1600]
    bars = []
    for index, (close, volume) in enumerate(zip(closes, volumes)):
        day = date(2025, 1, 1) + timedelta(days=index)
        bars.append({'event_date': day.isoformat(), 'open': str(close),
                     'high': str(round(close * 1.01, 4)), 'low': str(round(close * 0.99, 4)),
                     'close': str(close), 'volume': volume,
                     'available_at': datetime.combine(day + timedelta(days=1),
                                                       datetime.min.time(), timezone.utc).isoformat()})
    old_points = [LegacyCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                    high=float(bar['high']), low=float(bar['low']),
                                    close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
    new_points = [NewTrendKingCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                          high=float(bar['high']), low=float(bar['low']),
                                          close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
    expected = legacy_limit_up_arb(old_points, symbol='sh600000')
    assert new_limit_up_arb(new_points, symbol='600000') == expected
    for prefix, plain, threshold in [('sz', '300001', 1.20), ('bj', '430001', 1.30)]:
        board_old = old_points[:-2] + [old_points[-2].model_copy(update={'close': 10 * threshold}),
                                        old_points[-1].model_copy(update={'close': 10 * threshold * 1.04})]
        board_new = new_points[:-2] + [NewTrendKingCandlePoint(**{**vars(new_points[-2]), 'close': 10 * threshold}),
                                        NewTrendKingCandlePoint(**{**vars(new_points[-1]), 'close': 10 * threshold * 1.04})]
        assert new_limit_up_arb(board_new, symbol=plain)['prev_limit_up'] == legacy_limit_up_arb(
            board_old, symbol=prefix + plain)['prev_limit_up']
    params = {'min_day_gain': '3', 'min_volume_ratio_prev': '1.2',
              'min_volume_ratio_ma5': '0', 'sector_rank_weight': '0.08'}
    assert new_limit_up_arb_evaluate(expected, params) == legacy_limit_up_arb_evaluate(expected, params)
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        catalog = client.get('/api/v1/research/strategies').json()['data']
        assert next(item for item in catalog if item['id'] == 'limit_up_arb_v1')['status'] == 'limit_up_arb_signal'
        assert sum(item['status'] == 'not_migrated' for item in catalog) == 0
        imported = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'bars': bars, 'adjustment': 'none'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'arb-market'})
        assert imported.status_code == 200, imported.text
        body = {'dataset_id': imported.json()['data']['id'], 'strategy_id': 'limit_up_arb_v1',
                'decision_at': '2025-01-08T12:00:00+00:00', 'strict': True,
                'params': params}
        response = client.post('/api/v1/research/runs', json=body,
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'arb-run'})
        assert response.status_code == 200, response.text
        run = response.json()['data']
        assert run['result']['indicator'] == expected
        assert run['result']['evaluation'] == legacy_limit_up_arb_evaluate(expected, params)
        assert run['result']['signal'] is True
        assert len(run['result']['code_sha256']) == 64
        assert run['result']['evaluation']['entry_timing'] == 'same_day_close'
        assert 'st_status_unknown' in run['result']['quality_flags']
        assert client.get('/api/v1/research/runs/' + run['id']).json()['data'] == run
        short = client.post('/api/v1/market/datasets', json={
            'symbol': '600001', 'bars': bars[:4], 'adjustment': 'none'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'arb-short-market'})
        assert short.status_code == 200, short.text
        short_run = client.post('/api/v1/research/runs', json={**body,
            'dataset_id': short.json()['data']['id'],
            'params': {**params, 'min_volume_ratio_ma5': '1.0'}},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'arb-short-run'})
        assert short_run.status_code == 200, short_run.text
        assert short_run.json()['data']['result']['status'] == 'insufficient_data'
        assert 'ma5_volume_unavailable' in short_run.json()['data']['result']['quality_flags']


def test_ths_purple_and_golden_signals_match_legacy(tmp_path: Path) -> None:
    cases = [
        ('ths_main_force_flip_v1', 'purple_to_yellow', [10, 10.2, 10.4, 10.6, 10.4, 10.2, 10, 9.8, 10]),
        ('ths_main_force_golden_cross_v1', 'golden_cross', [10, 9.8, 9.6, 9.4, 9.2, 9.4, 9.6, 9.8]),
    ]
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        catalog = client.get('/api/v1/research/strategies').json()['data']
        assert sum(row['status'] == 'ths_signal' for row in catalog) == 2
        assert sum(row['status'] == 'not_migrated' for row in catalog) == 0
        for index, (strategy_id, trigger, closes) in enumerate(cases):
            bars = []
            for offset, close in enumerate(closes):
                day = date(2025, 1, 1) + timedelta(days=offset)
                bars.append({'event_date': day.isoformat(), 'open': f'{close:.4f}',
                             'high': f'{close * 1.01:.4f}', 'low': f'{close * 0.99:.4f}',
                             'close': f'{close:.4f}', 'volume': 1000,
                             'available_at': datetime.combine(day + timedelta(days=1),
                                                               datetime.min.time(), timezone.utc).isoformat()})
            old_points = [LegacyCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                            high=float(bar['high']), low=float(bar['low']),
                                            close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
            new_points = [NewTrendKingCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                                  high=float(bar['high']), low=float(bar['low']),
                                                  close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
            expected = legacy_ths_signal(old_points)
            assert expected[trigger] is True
            assert new_ths_signal(new_points) == expected
            imported = client.post('/api/v1/market/datasets', json={
                'symbol': f'60000{index}', 'bars': bars, 'adjustment': 'none'},
                headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'ths-market-{index}'})
            assert imported.status_code == 200, imported.text
            body = {'dataset_id': imported.json()['data']['id'], 'strategy_id': strategy_id,
                    'decision_at': '2025-01-12T12:00:00+00:00', 'strict': True, 'params': {}}
            response = client.post('/api/v1/research/runs', json=body,
                                   headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'ths-run-{index}'})
            assert response.status_code == 200, response.text
            run = response.json()['data']
            assert run['result']['indicator'] == expected
            assert run['result']['signal'] is True
            assert run['result']['evaluation']['trigger_key'] == trigger
            assert run['result']['code_sha256']
            assert 'sector_context_unavailable' not in run['result']['quality_flags']
            assert 'indicator_warmup_short' in run['result']['quality_flags']
            assert client.get('/api/v1/research/runs/' + run['id']).json()['data'] == run
            bad = client.post('/api/v1/research/runs', json={**body, 'params': {'min_score': '50'}},
                              headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'ths-bad-{index}'})
            assert bad.status_code == 400


def test_emotion_limit_up_signal_matches_legacy_on_frozen_bars(tmp_path: Path) -> None:
    closes = [10] * 12 + [10, 10.2, 10.4, 10.6, 10.4, 10.2, 10, 9.8, 11]
    bars = []
    for index, close in enumerate(closes):
        day = date(2025, 1, 1) + timedelta(days=index)
        bars.append({'event_date': day.isoformat(), 'open': f'{close:.4f}',
                     'high': f'{close * 1.01:.4f}', 'low': f'{close * 0.99:.4f}',
                     'close': f'{close:.4f}', 'volume': 1000,
                     'available_at': datetime.combine(day + timedelta(days=1),
                                                       datetime.min.time(), timezone.utc).isoformat()})
    old_points = [LegacyCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                    high=float(bar['high']), low=float(bar['low']),
                                    close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
    new_points = [NewTrendKingCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                          high=float(bar['high']), low=float(bar['low']),
                                          close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
    expected = legacy_emotion_signal(old_points)
    assert new_emotion_signal(new_points) == expected
    assert legacy_emotion_evaluate(expected)['signal'] is True
    params = {'min_day_gain': '9.5', 'lookback_days': '5', 'min_score': '0'}
    assert new_emotion_evaluate(expected, params) == legacy_emotion_evaluate(expected, params)
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        catalog = client.get('/api/v1/research/strategies').json()['data']
        assert next(item for item in catalog if item['id'] == 'emotion_limit_up_v1')['status'] == 'emotion_signal'
        assert sum(item['status'] == 'not_migrated' for item in catalog) == 0
        imported = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'bars': bars, 'adjustment': 'none'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'emotion-market'})
        assert imported.status_code == 200, imported.text
        body = {'dataset_id': imported.json()['data']['id'], 'strategy_id': 'emotion_limit_up_v1',
                'decision_at': '2025-01-24T12:00:00+00:00', 'strict': True, 'params': params}
        response = client.post('/api/v1/research/runs', json=body,
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'emotion-run'})
        assert response.status_code == 200, response.text
        run = response.json()['data']
        assert run['result']['indicator'] == expected
        assert run['result']['evaluation'] == legacy_emotion_evaluate(expected, params)
        assert run['result']['signal'] is True
        assert 'sector_context_unavailable' in run['result']['quality_flags']
        assert client.get('/api/v1/research/runs/' + run['id']).json()['data'] == run
        rejected = client.post('/api/v1/research/runs', json={**body, 'params': {'require_next_day_confirm': 'true'}},
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'emotion-reject'})
        assert rejected.status_code == 400


def test_wulong_single_symbol_indicator_matches_legacy(tmp_path: Path, monkeypatch) -> None:
    bars = []
    for index in range(90):
        day = date(2025, 1, 1) + timedelta(days=index)
        close = 10 + index * 0.05
        bars.append({'event_date': day.isoformat(), 'open': f'{close:.4f}',
                     'high': f'{close + 0.1:.4f}', 'low': f'{close - 0.1:.4f}',
                     'close': f'{close:.4f}', 'volume': 1000 + index * 2,
                     'available_at': datetime.combine(day + timedelta(days=1),
                                                       datetime.min.time(), timezone.utc).isoformat()})
    old_points = [LegacyCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                    high=float(bar['high']), low=float(bar['low']),
                                    close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
    new_points = [NewTrendKingCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                          high=float(bar['high']), low=float(bar['low']),
                                          close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
    expected = legacy_wulong_signal(old_points)
    assert new_wulong_signal(new_points) == expected
    params = {'convergence_threshold_pct': '0.04', 'pre_convergence_min_spread_pct': '0.02',
              'min_current_spread_pct': '0.01', 'min_spread_expansion_multiple': '1.6',
              'min_volume_ratio20': '1.2', 'min_breakout_return_pct': '0.006',
              'max_convergence_age_days': '10', 'min_rising_ma_count': '3'}
    assert new_wulong_evaluate(expected, params) == legacy_wulong_evaluate(expected, params)
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        catalog = client.get('/api/v1/research/strategies').json()['data']
        assert next(item for item in catalog if item['id'] == 'wulong_cluster_v1')['status'] == 'wulong_signal'
        assert sum(item['status'] == 'not_migrated' for item in catalog) == 0
        imported = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'bars': bars, 'adjustment': 'none'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'wulong-market'})
        assert imported.status_code == 200, imported.text
        body = {'dataset_id': imported.json()['data']['id'], 'strategy_id': 'wulong_cluster_v1',
                'decision_at': '2025-04-04T12:00:00+00:00', 'strict': True, 'params': params}
        response = client.post('/api/v1/research/runs', json=body,
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'wulong-run'})
        assert response.status_code == 200, response.text
        run = response.json()['data']
        assert run['result']['indicator'] == expected
        assert run['result']['evaluation'] == legacy_wulong_evaluate(expected, params)
        assert run['result']['quality_flags'] == ['candidate_universe_rejected']
        assert run['result']['universe']['failed_conditions'] == ['allow_upper_shadow_risk']
        assert client.get('/api/v1/research/runs/' + run['id']).json()['data'] == run
        bad = client.post('/api/v1/research/runs', json={**body, 'params': {'max_convergence_age_days': '1.5'}},
                          headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'wulong-bad'})
        assert bad.status_code == 400
        sparse = client.post('/api/v1/market/datasets', json={
            'symbol': '600001', 'bars': bars[:2], 'adjustment': 'none'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'wulong-sparse'})
        assert sparse.status_code == 200, sparse.text
        sparse_run = client.post('/api/v1/research/runs', json={**body, 'dataset_id': sparse.json()['data']['id']},
                                 headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'wulong-sparse-run'})
        assert sparse_run.status_code == 200, sparse_run.text
        assert sparse_run.json()['data']['result']['status'] == 'insufficient_data'
        assert sparse_run.json()['data']['result']['indicator']['convergence_spread_pct'] is None
        from trade_app.research import runtime as research_service
        monkeypatch.setattr(research_service, 'evaluate_wulong_cluster_signal',
                            lambda _indicator, _params: {'signal': True, 'signal_score': 90})
        incomplete = client.post('/api/v1/research/runs', json={
            **body, 'params': {**params, 'min_rising_ma_count': '4', 'min_ma10_above_ma20_days': '30'}},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'wulong-incomplete'})
        assert incomplete.status_code == 200, incomplete.text
        incomplete_run = incomplete.json()['data']
        assert incomplete_run['result']['shape_signal'] is True
        assert incomplete_run['result']['signal'] is False
        assert incomplete_run['result']['draft_eligible'] is False
        sim = client.post('/api/v1/sim-accounts', json={
            'name': '五龙草稿隔离', 'initial_capital': '10000', 'start_date': '2025-04-04'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'wulong-sim'})
        assert sim.status_code == 200, sim.text
        draft = client.post(f"/api/v1/sim-accounts/{sim.json()['data']['id']}/drafts", json={
            'source_run_id': incomplete_run['id'], 'quantity': 100, 'limit_price': '15'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'wulong-draft'})
        assert draft.status_code == 409


def test_force_rhythm_signal_and_boolean_params_match_legacy(tmp_path: Path) -> None:
    bars = []
    for index in range(100):
        day = date(2025, 1, 1) + timedelta(days=index)
        close = 10 + sin(index * 0.7) * 0.5
        bars.append({'event_date': day.isoformat(), 'open': f'{close:.4f}',
                     'high': f'{close + 0.1:.4f}', 'low': f'{close - 0.1:.4f}',
                     'close': f'{close:.4f}', 'volume': 1000 + int(sin((index + 1) * 0.7) * 600),
                     'available_at': datetime.combine(day + timedelta(days=1),
                                                       datetime.min.time(), timezone.utc).isoformat()})
    old_points = [LegacyCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                    high=float(bar['high']), low=float(bar['low']),
                                    close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
    new_points = [NewTrendKingCandlePoint(time=bar['event_date'], open=float(bar['open']),
                                          high=float(bar['high']), low=float(bar['low']),
                                          close=float(bar['close']), volume=bar['volume'], amount=0) for bar in bars]
    expected = legacy_rhythm_signal(old_points)
    assert new_rhythm_signal(new_points) == expected
    assert legacy_rhythm_evaluate(expected)['signal'] is True
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        catalog = client.get('/api/v1/research/strategies').json()['data']
        descriptor = next(item for item in catalog if item['id'] == 'ths_force_rhythm_v1')
        assert descriptor['status'] == 'rhythm_signal'
        assert descriptor['signal_params']['require_trough_turn'] == 'true'
        assert sum(item['status'] == 'not_migrated' for item in catalog) == 0
        imported = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'bars': bars, 'adjustment': 'none'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'rhythm-market'})
        assert imported.status_code == 200, imported.text
        params = {'require_trough_turn': 'false', 'block_buy_on_retail_sell': 'false'}
        body = {'dataset_id': imported.json()['data']['id'], 'strategy_id': 'ths_force_rhythm_v1',
                'decision_at': '2025-04-12T12:00:00+00:00', 'strict': True, 'params': params}
        response = client.post('/api/v1/research/runs', json=body,
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'rhythm-run'})
        assert response.status_code == 200, response.text
        run = response.json()['data']
        assert run['params']['require_trough_turn'] is False
        assert run['params']['block_buy_on_retail_sell'] is False
        assert run['result']['indicator'] == expected
        assert run['result']['evaluation'] == legacy_rhythm_evaluate(expected, run['params'])
        assert run['result']['signal'] is True
        assert run['result']['calculation_version'] == 'legacy-force-rhythm-frozen-v1'
        assert client.get('/api/v1/research/runs/' + run['id']).json()['data'] == run
        invalid = client.post('/api/v1/research/runs', json={**body,
            'params': {'require_trough_turn': 'maybe'}},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'rhythm-invalid'})
        assert invalid.status_code == 400


@contextmanager
def started_client(app):
    lifespan = app.router.lifespan_context(app)
    with asyncio.Runner() as runner:
        runner.run(lifespan.__aenter__())
        try:
            with TestClient(app) as client:
                yield client
        finally:
            runner.run(lifespan.__aexit__(None, None, None))


def test_account_ledger_nav_and_revision(tmp_path: Path) -> None:
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(method: str, path: str, body: dict | None = None):
            nonlocal serial
            serial += 1
            response = client.request(method, path, json=body,
                                      headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'test-{serial}'})
            assert response.status_code == 200, response.text
            return response.json()['data']

        account = write('POST', '/api/v1/accounts', {'name': '主账户'})
        root = f"/api/v1/accounts/{account['id']}"
        write('POST', root + '/cash-flows', {'flow_date': '2025-01-02', 'kind': 'initial', 'amount': '10000'})
        write('PUT', root + '/snapshots/2025-01-02', {'snap_date': '2025-01-02', 'total_assets': '11000',
                                                        'expected_revision': 0})
        write('POST', root + '/cash-flows', {'flow_date': '2025-01-03', 'kind': 'deposit', 'amount': '1100'})
        write('PUT', root + '/snapshots/2025-01-03', {'snap_date': '2025-01-03', 'total_assets': '12100',
                                                        'expected_revision': 0})
        trade = write('POST', root + '/trades', {'trade_date': '2025-01-03', 'symbol': '600000',
                                                 'side': 'buy', 'quantity': 100, 'price': '10.00'})
        repeat = client.post(root + '/trades', json={'trade_date': '2025-01-03', 'symbol': '600000',
                              'side': 'buy', 'quantity': 100, 'price': '10.00'},
                              headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'test-6'})
        assert repeat.status_code == 200 and repeat.json()['data']['id'] == trade['id']
        review = write('PUT', root + '/daily-reviews/2025-01-03',
                       {'expected_revision': 0, 'title': '复盘', 'decision_review': '遵守计划'})
        assert review['revision'] == 1
        assert client.get(root + '/analytics').json()['data']['status'] == 'recalculating'
        assert process_one(app.state.db_factory)
        analytic = client.get(root + '/analytics').json()['data']
        assert analytic['status'] == 'fresh'
        assert Decimal(analytic['result']['nav']['current']['nav']) == Decimal('1.1')
        assert analytic['result']['positions'][0]['quantity'] == 100

        changed = write('PUT', root + '/trades/' + trade['id'], {**{
            'trade_date': '2025-01-03', 'symbol': '600000', 'side': 'buy',
            'quantity': 200, 'price': '10.00'}, 'expected_revision': trade['revision']})
        assert changed['revision'] == 2
        assert client.get(root + '/analytics').json()['data']['status'] == 'recalculating'
        assert process_one(app.state.db_factory)
        assert client.get(root + '/analytics').json()['data']['result']['positions'][0]['quantity'] == 200
        assert client.get(root + '/daily-reviews/2025-01-03').json()['data']['decision_review'] == '遵守计划'
        assert client.get(root + '/snapshots').json()['data'][-1]['total_assets'] == '12100.00'
        conflict = client.put(root + '/trades/' + trade['id'], json={
            'trade_date': '2025-01-03', 'symbol': '600000', 'side': 'buy',
            'quantity': 300, 'price': '10.00', 'expected_revision': 1},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'conflict'})
        assert conflict.status_code == 409
        assert conflict.json()['error']['code'] == 'REVISION_CONFLICT'

        second = write('POST', '/api/v1/accounts', {'name': '第二账户'})
        wrong_scope = client.put(f"/api/v1/accounts/{second['id']}/trades/{trade['id']}", json={
            'trade_date': '2025-01-03', 'symbol': '600000', 'side': 'buy',
            'quantity': 300, 'price': '10.00', 'expected_revision': 2},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'wrong-scope'})
        assert wrong_scope.status_code == 404


def test_confirmed_snapshot_period_performance_and_missing_baseline(tmp_path: Path) -> None:
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(method: str, path: str, body: dict):
            nonlocal serial
            serial += 1
            response = client.request(method, path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': f'perf-{serial}'})
            assert response.status_code == 200, response.text
            return response.json()['data']

        account = write('POST', '/api/v1/accounts', {'name': '统计账户'})
        root = f"/api/v1/accounts/{account['id']}"
        write('POST', root + '/cash-flows', {'flow_date': '2025-01-01',
            'kind': 'initial', 'amount': '10000'})
        for day, assets in [('2025-01-01', '10000'), ('2025-01-02', '11000')]:
            write('PUT', root + '/snapshots/' + day, {
                'snap_date': day, 'total_assets': assets, 'expected_revision': 0})
        write('POST', root + '/cash-flows', {'flow_date': '2025-01-03',
            'kind': 'deposit', 'amount': '1100'})
        write('PUT', root + '/snapshots/2025-01-03', {
            'snap_date': '2025-01-03', 'total_assets': '12100', 'expected_revision': 0})
        write('PUT', root + '/snapshots/2025-01-06', {
            'snap_date': '2025-01-06', 'total_assets': '13310', 'expected_revision': 0})
        assert process_one(app.state.db_factory)
        daily = client.get(root + '/performance?kind=daily').json()['data']
        assert daily['projection_status'] == 'fresh'
        rows = {item['key']: item for item in daily['items']}
        assert rows['2025-01-01']['return_pct'] is None
        assert rows['2025-01-02']['return_pct'] == '10.0000'
        assert rows['2025-01-03']['return_pct'] == '0.0000'
        assert rows['2025-01-06']['return_pct'] == '10.0000'
        assert rows['2025-01-06']['return_quality'] == 'gap'
        weekly = client.get(root + '/performance?kind=weekly').json()['data']
        assert weekly['items'][0]['key'] == '2025-W02'
        assert weekly['items'][0]['return_pct'] == '10.0000'
        assert weekly['items'][1]['key'] == '2025-W01'
        assert weekly['items'][1]['return_pct'] is None
        assert client.get(root + '/performance?kind=yearly').status_code == 400
        assert client.get(root + '/performance?limit=367').status_code == 400
        sim = write('POST', '/api/v1/sim-accounts', {'name': '模拟隔离',
            'initial_capital': '1000', 'start_date': '2025-01-01'})
        assert client.get(f"/api/v1/accounts/{sim['id']}/performance").status_code == 404


def test_daily_review_full_manual_sections_and_history(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        def write(path: str, body: dict, key: str):
            return client.put(path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': key})
        account = client.post('/api/v1/accounts', json={'name': '完整日复盘'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'review-account'}).json()['data']
        root = f"/api/v1/accounts/{account['id']}"
        path = root + '/daily-reviews/2025-01-06'
        body = {'expected_revision': 0, 'title': '周一复盘',
                'overall_summary': '总述', 'market_observation': '指数',
                'decision_review': '遵守计划', 'reflection': '反思', 'mistakes': '冲动',
                'tomorrow_plan': '耐心等待', 'tags': ['突破', '纪律'],
                'next_market_forecast': '震荡',
                'next_watchlist': [{'code': '600000', 'name': '浦发', 'condition': '突破十日高点', 'action': '观察'}],
                'next_position_plan': '四成仓', 'next_risk_plan': '跌破止损',
                'next_position_rehearsal': [{'code': '600000', 'name': '浦发', 'qty': 100, 'note': '已有持仓'}]}
        saved = write(path, body, 'review-full')
        assert saved.status_code == 200, saved.text
        data = saved.json()['data']
        assert data['overall_summary'] == '总述'
        assert data['reflection'] == '反思'
        assert data['tags'] == ['突破', '纪律']
        assert data['next_watchlist'][0]['condition'] == '突破十日高点'
        assert data['next_position_rehearsal'][0]['qty'] == 100
        history = client.get(root + '/daily-reviews').json()['data']
        assert history[0]['date'] == '2025-01-06' and history[0]['tags'] == ['突破', '纪律']
        assert write(path, {**body, 'expected_revision': 1,
                            'next_watchlist': [{'code': '', 'name': '', 'condition': '', 'action': ''}]},
                     'review-invalid').status_code == 400
        assert client.get(path).json()['data']['revision'] == 1
        assert write(path, {**body, 'expected_revision': 0}, 'review-stale').status_code == 409
        other = client.post('/api/v1/accounts', json={'name': '隔离日复盘'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'review-other'}).json()['data']
        assert client.get(f"/api/v1/accounts/{other['id']}/daily-reviews").json()['data'] == []


def test_review_attachment_upload_scope_and_backup_round_trip(tmp_path: Path) -> None:
    source = tmp_path / 'review-image-source'
    with started_client(create_app(source, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        headers = {'X-CSRF-Token': csrf, 'Idempotency-Key': 'attachment-account'}
        account = client.post('/api/v1/accounts', json={'name': '图片复盘'},
                              headers=headers).json()['data']
        other = client.post('/api/v1/accounts', json={'name': '其他账户'},
                            headers={**headers, 'Idempotency-Key': 'attachment-other'}).json()['data']
        root = f"/api/v1/accounts/{account['id']}"
        day = '2025-01-06'
        output = io.BytesIO()
        Image.new('RGB', (16, 12), color=(30, 80, 120)).save(output, format='PNG')
        content = output.getvalue()
        endpoint = root + f'/daily-reviews/{day}/attachments'
        added = client.post(endpoint, files={'file': ('截图.png', content, 'image/png')},
                            headers={**headers, 'Idempotency-Key': 'attachment-upload'})
        assert added.status_code == 200, added.text
        item = added.json()['data']
        assert item['width'] == 16 and item['height'] == 12
        assert client.get(endpoint).json()['data'][0]['id'] == item['id']
        image = client.get(item['url'])
        assert image.status_code == 200 and image.content == content
        assert image.headers['content-type'].startswith('image/png')
        assert client.get(f"/api/v1/accounts/{other['id']}/review-attachments/{item['id']}/content").status_code == 404
        assert client.get(f"/api/v1/accounts/{other['id']}/daily-reviews/{day}/attachments").json()['data'] == []
        duplicate = client.post(endpoint, files={'file': ('截图.png', content, 'image/png')},
                                headers={**headers, 'Idempotency-Key': 'attachment-upload'})
        assert duplicate.json()['data']['id'] == item['id']
        invalid = client.post(endpoint, files={'file': ('fake.png', b'not an image', 'image/png')},
                              headers={**headers, 'Idempotency-Key': 'attachment-invalid'})
        assert invalid.status_code == 400
        backup = create_backup(source)
        manifest = verify_backup(backup)
        assert f"attachments/{hashlib.sha256(content).hexdigest()}.bin" in manifest['files']
        assert manifest['format'] == 'trade-rebuild-backup-v3'
        deleted = client.delete(f"{root}/review-attachments/{item['id']}?expected_revision=1",
                                headers={**headers, 'Idempotency-Key': 'attachment-delete'})
        assert deleted.status_code == 200
        assert client.get(item['url']).status_code == 404
        assert client.get(endpoint).json()['data'] == []
    restored = tmp_path / 'review-image-restored'
    restore_to_new_directory(backup, restored)
    with started_client(create_app(restored, auto_rebuild=False)) as client:
        client.get('/api/v1/session')
        assert client.get(item['url']).content == content
    (source / 'attachments' / f"{hashlib.sha256(content).hexdigest()}.bin").write_bytes(b'corrupt')
    try:
        create_backup(source)
    except ValueError:
        pass
    else:
        raise AssertionError('损坏的附件不应进入备份')


def test_manual_daily_trade_and_t_group_scores_are_separate(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0
        def write(path: str, body: dict, method: str = 'POST'):
            nonlocal serial
            serial += 1
            return client.request(method, path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': f'scores-{serial}'})
        account = write('/api/v1/accounts', {'name': '人工评分'}).json()['data']
        root = f"/api/v1/accounts/{account['id']}"
        day = '2025-01-06'
        buy = write(root + '/trades', {'trade_date': day, 'symbol': '600000',
            'side': 'buy', 'quantity': 100, 'price': '10'}).json()['data']
        sell = write(root + '/trades', {'trade_date': day, 'symbol': '600000',
            'side': 'sell', 'quantity': 100, 'price': '11'}).json()['data']
        path = root + f'/daily-reviews/{day}/scores'
        daily = write(path, {'scope': 'daily', 'trade_ids': [],
            'scores': {'position': {'final': 8, 'comment': '仓位适中'},
                       'emotion': {'final': 7, 'comment': ''}},
            'comment': '总体稳定', 'expected_revision': 0}, 'PUT')
        assert daily.status_code == 200, daily.text
        assert daily.json()['data']['scores']['position']['final_source'] == 'manual'
        assert daily.json()['data']['scores']['position']['ai'] is None
        trade = write(path, {'scope': 'trade', 'trade_ids': [buy['id']],
            'scores': {'timing': {'final': 9, 'comment': '买点好'}},
            'expected_revision': 0}, 'PUT')
        assert trade.status_code == 200, trade.text
        group = write(path, {'scope': 't_group', 'trade_ids': [sell['id'], buy['id']],
            'scores': {'timing': {'final': 6, 'comment': '价差偏小'}},
            'expected_revision': 0}, 'PUT')
        assert group.status_code == 200, group.text
        assert group.json()['data']['trade_ids'] == sorted([buy['id'], sell['id']])
        assert len(client.get(path).json()['data']) == 3
        assert write(path, {'scope': 'trade', 'trade_ids': [buy['id']],
            'scores': {'position': {'final': 5}}, 'expected_revision': 1}, 'PUT').status_code == 400
        assert write(path, {'scope': 't_group', 'trade_ids': [buy['id']],
            'scores': {}, 'expected_revision': 0}, 'PUT').status_code == 400
        assert write(path, {'scope': 'daily', 'trade_ids': [],
            'scores': {'position': {'final': 11}}, 'expected_revision': 1}, 'PUT').status_code == 422
        assert write(path, {'scope': 'daily', 'trade_ids': [],
            'scores': {'position': {'final': 9}}, 'expected_revision': 0}, 'PUT').status_code == 409
        other = write('/api/v1/accounts', {'name': '评分隔离'}).json()['data']
        assert client.get(f"/api/v1/accounts/{other['id']}/daily-reviews/{day}/scores").json()['data'] == []
        assert write(root + '/trades/' + buy['id'] + '?expected_revision=1', {}, 'DELETE').status_code == 200
        saved_sheets = client.get(path).json()['data']
        assert next(row for row in saved_sheets if row['scope'] == 'trade')['association_changed']
        assert next(row for row in saved_sheets if row['scope'] == 't_group')['association_changed']


def test_next_plan_target_date_and_confirmed_actual_comparison(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0
        def write(method: str, path: str, body: dict):
            nonlocal serial
            serial += 1
            return client.request(method, path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': f'plan-{serial}'})
        account = write('POST', '/api/v1/accounts', {'name': '计划对照'}).json()['data']
        root = f"/api/v1/accounts/{account['id']}"
        plan = write('PUT', root + '/daily-reviews/2025-01-03', {
            'expected_revision': 0, 'next_market_forecast': '震荡',
            'next_target_date': '2025-01-06',
            'next_watchlist': [{'code': '600000', 'name': '浦发', 'condition': '十日新高', 'action': '小仓试买'}],
            'next_position_rehearsal': [{'code': '600000', 'name': '浦发', 'qty': 100, 'note': ''}]})
        assert plan.status_code == 200, plan.text
        assert plan.json()['data']['next_target_date'] == '2025-01-06'
        comparison = client.get(root + '/plan-comparison/2025-01-06').json()['data']
        assert comparison['plans'][0]['rehearsal'][0]['actual_qty'] is None
        assert comparison['plans'][0]['rehearsal'][0]['quality'] == 'missing_snapshot'
        assert write('POST', root + '/cash-flows', {
            'flow_date': '2025-01-03', 'kind': 'initial', 'amount': '10000'}).status_code == 200
        trade = write('POST', root + '/trades', {'trade_date': '2025-01-06',
            'symbol': '600000', 'side': 'buy', 'quantity': 100, 'price': '10'})
        assert trade.status_code == 200
        snapshot = write('PUT', root + '/snapshots/2025-01-06', {
            'snap_date': '2025-01-06', 'total_assets': '10000',
            'positions': [{'symbol': '600000', 'name': '浦发', 'quantity': 80,
                           'market_value': '800'}], 'expected_revision': 0})
        assert snapshot.status_code == 200, snapshot.text
        compared = client.get(root + '/plan-comparison/2025-01-06').json()['data']
        assert compared['plans'][0]['watchlist'][0]['matching_trade_ids'] == [trade.json()['data']['id']]
        assert compared['plans'][0]['rehearsal'][0]['actual_qty'] == 80
        assert compared['plans'][0]['rehearsal'][0]['quantity_delta'] == -20
        assert compared['snapshot']['total_assets'] == '10000.00'
        assert write('PUT', root + '/daily-reviews/2025-01-03', {
            'expected_revision': 1, 'next_target_date': '2025-01-03'}).status_code == 400
        other = write('POST', '/api/v1/accounts', {'name': '计划隔离'}).json()['data']
        assert client.get(f"/api/v1/accounts/{other['id']}/plan-comparison/2025-01-06").json()['data']['plans'] == []


def test_pending_trades_batch_duplicate_and_history(tmp_path: Path) -> None:
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def request(method: str, path: str, body: dict | None = None, expected: int = 200):
            nonlocal serial
            serial += 1
            response = client.request(method, path, json=body,
                                      headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'pending-{serial}'})
            assert response.status_code == expected, response.text
            return response.json()

        account = request('POST', '/api/v1/accounts', {'name': '待确认测试'})['data']
        root = f"/api/v1/accounts/{account['id']}"
        trade = {'trade_date': '2025-01-02', 'symbol': '600000', 'side': 'buy',
                 'quantity': 100, 'price': '10.00', 'fee': '5'}
        one = request('POST', root + '/pending-trades', trade)['data']
        assert request('GET', root + '/trades')['data'] == []
        assert client.get(root + '/analytics').json()['data']['account_input_revision'] == 0
        changed = request('PUT', root + '/pending-trades/' + one['id'],
                          {**trade, 'note': '已核对', 'expected_revision': 1})['data']
        assert changed['revision'] == 2
        assert request('POST', root + '/pending-trades/' + one['id'] + '/confirm',
                       {'expected_revision': 1}, 409)['error']['code'] == 'REVISION_CONFLICT'
        confirmed = request('POST', root + '/pending-trades/' + one['id'] + '/confirm',
                            {'expected_revision': 2})['data']
        assert request('POST', root + '/pending-trades/' + one['id'] + '/confirm',
                       {'expected_revision': 2})['data']['trade_id'] == confirmed['trade_id']
        assert len(request('GET', root + '/trades')['data']) == 1
        duplicate = request('POST', root + '/pending-trades', trade)['data']
        assert duplicate['duplicate_trade_ids'] == [confirmed['trade_id']]
        other = request('POST', root + '/pending-trades', {**trade, 'symbol': '000001'})['data']
        second = request('POST', '/api/v1/accounts', {'name': '隔离账户'})['data']
        assert request('GET', f"/api/v1/accounts/{second['id']}/pending-trades")['data'] == []
        assert request('POST', f"/api/v1/accounts/{second['id']}/pending-trades/{other['id']}/confirm",
                       {'expected_revision': 1}, 404)['error']['code'] == 'PENDING_TRADE_NOT_FOUND'
        batch = request('POST', root + '/pending-trades/confirm-batch', {'items': [
            {'id': duplicate['id'], 'expected_revision': 1},
            {'id': other['id'], 'expected_revision': 1},
        ]})['data']
        assert batch['confirmed_count'] == 1
        assert batch['results'][0]['code'] == 'POSSIBLE_DUPLICATE'
        assert batch['results'][1]['status'] == 'confirmed'
        assert len(request('GET', root + '/trades')['data']) == 2
        assert len(request('GET', root + '/pending-trades')['data']) == 1
        assert request('POST', root + '/pending-trades/' + duplicate['id'] + '/confirm',
                       {'expected_revision': 1}, 409)['error']['code'] == 'POSSIBLE_DUPLICATE'
        request('POST', root + '/pending-trades/' + duplicate['id'] + '/confirm',
                {'expected_revision': 1, 'acknowledge_duplicate': True})
        to_delete = request('POST', root + '/pending-trades', {**trade, 'symbol': '000002'})['data']
        request('DELETE', root + '/pending-trades/' + to_delete['id'] + '?expected_revision=1')
        assert request('GET', root + '/pending-trades')['data'] == []
        assert client.get(root + '/analytics').json()['data']['account_input_revision'] == 3


def test_real_fee_settings_auto_manual_and_pending_snapshot(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(method: str, path: str, body: dict):
            nonlocal serial
            serial += 1
            response = client.request(method, path, json=body,
                                      headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'fee-{serial}'})
            assert response.status_code == 200, response.text
            return response.json()['data']

        account = write('POST', '/api/v1/accounts', {'name': '费用测试'})
        base = f"/api/v1/accounts/{account['id']}"
        sample = {'trade_date': '2025-01-03', 'symbol': '600000', 'side': 'sell',
                  'quantity': 100, 'price': '10', 'fee_mode': 'auto'}
        preview = write('POST', base + '/fee-preview', sample)
        assert preview['fee'] == '6.01'
        auto = write('POST', base + '/trades', sample)
        assert auto['fee'] == auto['calculated_fee'] == '6.01'
        assert auto['fee_source'] == 'auto' and auto['fee_rule_version'] == 1
        assert auto['fee_breakdown']['commission'] == '5.00'
        overridden = write('POST', base + '/trades', {**sample, 'symbol': '000001',
                            'fee_mode': 'manual', 'fee': '3.25'})
        assert overridden['fee'] == '3.25' and overridden['calculated_fee'] == '6.01'
        assert overridden['fee_source'] == 'manual'
        staged = write('POST', base + '/pending-trades', {**sample, 'symbol': '000002'})
        assert staged['fee'] == '6.01' and staged['fee_source'] == 'auto'
        current = client.get(base + '/fee-settings').json()['data']
        assert current['version'] == 1
        changed = write('PUT', base + '/fee-settings', {
            'expected_version': 1,
            'config': {**current['config'], 'minimum_commission': '1.00',
                       'sell_stamp_rate': '0.0005'}})
        assert changed['version'] == 2
        assert write('POST', base + '/fee-preview', sample)['fee'] == '1.51'
        confirmed = write('POST', base + '/pending-trades/' + staged['id'] + '/confirm',
                          {'expected_revision': staged['revision']})
        historical = next(row for row in client.get(base + '/trades').json()['data']
                          if row['id'] == confirmed['trade_id'])
        assert historical['fee'] == '6.01' and historical['fee_rule_version'] == 1
        assert historical['fee_source'] == 'auto'


def test_versioned_target_nodes_and_projection_upgrade(tmp_path: Path) -> None:
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(method: str, path: str, body: dict):
            nonlocal serial
            serial += 1
            response = client.request(method, path, json=body,
                                      headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'node-{serial}'})
            assert response.status_code == 200, response.text
            return response.json()['data']

        account = write('POST', '/api/v1/accounts', {'name': '节点测试'})
        base = f"/api/v1/accounts/{account['id']}"
        write('POST', base + '/cash-flows',
              {'flow_date': '2025-01-01', 'kind': 'initial', 'amount': '100'})
        write('PUT', base + '/snapshots/2025-01-02',
              {'snap_date': '2025-01-02', 'total_assets': '130', 'expected_revision': 0})
        assert process_one(app.state.db_factory)
        before = client.get(base + '/analytics').json()['data']
        assert before['result']['nav']['current']['lit_count'] == 1
        assert before['result']['nav']['first_achievements'][0]['days_from_start'] == 1
        assert before['result']['nav']['target_config']['version'] == 1
        changed = write('PUT', base + '/target-config', {'expected_version': 1,
                                                        'multiplier': '2.0', 'node_count': 2})
        assert changed['version'] == 2
        assert process_one(app.state.db_factory)
        after = client.get(base + '/analytics').json()['data']
        assert after['result']['nav']['current']['lit_count'] == 0
        assert after['result']['nav']['current']['next_target_assets'] == '200.00'
        assert after['result']['nav']['current']['next_gap_amount'] == '70.00'
        assert after['result']['nav']['target_config']['version'] == 2
        with app.state.db_factory() as session:
            rows = session.execute(text(
                'SELECT state, payload_json FROM projection_versions WHERE account_id=:a ORDER BY created_at'),
                {'a': account['id']}).all()
            assert len(rows) == 2 and rows[0][0] == 'historical' and rows[1][0] == 'current'
            assert json.loads(rows[0][1])['nav']['current']['lit_count'] == 1
        write('PUT', base + '/snapshots/2025-01-03',
              {'snap_date': '2025-01-03', 'total_assets': '220', 'expected_revision': 0})
        assert process_one(app.state.db_factory)
        lit = client.get(base + '/analytics').json()['data']['result']['nav']
        assert lit['current']['lit_count'] == 1
        assert lit['first_achievements'][0]['days_from_start'] == 2
        write('PUT', base + '/snapshots/2025-01-04',
              {'snap_date': '2025-01-04', 'total_assets': '190', 'expected_revision': 0})
        assert process_one(app.state.db_factory)
        extinguished = client.get(base + '/analytics').json()['data']['result']['nav']
        assert extinguished['current']['lit_count'] == 0
        assert extinguished['node_events'][-1]['kind'] == 'extinguished'


def test_outdated_projection_rebuilds_without_overwriting_history(tmp_path: Path) -> None:
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        account = client.post('/api/v1/accounts', json={'name': '升级测试'}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'create'}).json()['data']
        assert process_one(app.state.db_factory) is False
        with app.state.db_factory.begin() as session:
            session.execute(text("INSERT INTO projection_versions(id, account_id, input_revision, calculation_version, payload_json, state, created_at) VALUES ('old', :account, 0, 'nav-rounds-v4', '{}', 'current', '2025-01-01T00:00:00Z')"),
                            {'account': account['id']})
    second = create_app(tmp_path, auto_rebuild=False)
    with started_client(second) as client:
        client.get('/api/v1/session')
        base = f"/api/v1/accounts/{account['id']}"
        assert client.get(base + '/analytics').json()['data']['status'] == 'recalculating'
        assert process_one(second.state.db_factory)
        latest = client.get(base + '/analytics').json()['data']
        assert latest['status'] == 'fresh' and latest['calculation_version'] == 'nav-rounds-v5'
        with second.state.db_factory() as session:
            rows = session.execute(text('SELECT state, calculation_version FROM projection_versions WHERE account_id=:account ORDER BY created_at'),
                                   {'account': account['id']}).all()
            assert rows == [('historical', 'nav-rounds-v4'), ('current', 'nav-rounds-v5')]


def test_round_ids_streaks_and_anomalies_are_separate() -> None:
    trades = []
    prices = [(10, 11), (10, 12), (10, 9), (10, 8)]
    for index, (buy, sell) in enumerate(prices):
        trades.extend([
            {'id': f'b{index}', 'trade_date': f'2025-01-{index * 2 + 1:02}', 'sequence': 1,
             'symbol': '600000', 'name': 'A', 'side': 'buy', 'quantity': 100,
             'price': str(buy), 'fee': '0'},
            {'id': f's{index}', 'trade_date': f'2025-01-{index * 2 + 2:02}', 'sequence': 1,
             'symbol': '600000', 'name': 'A', 'side': 'sell', 'quantity': 100,
             'price': str(sell), 'fee': '0'},
        ])
    trades.append({'id': 'bad', 'trade_date': '2025-01-09', 'sequence': 1,
                   'symbol': '600000', 'name': 'A', 'side': 'sell', 'quantity': 1,
                   'price': '10', 'fee': '0'})
    result = project_account([], [], trades)
    closed = [row for row in result['rounds'] if row['status'] == 'closed']
    assert [row['id'] for row in closed] == ['round:b0', 'round:b1', 'round:b2', 'round:b3']
    assert [row['pnl'] for row in closed] == ['100.00', '200.00', '-100.00', '-200.00']
    assert result['trade_stats']['closed_rounds'] == 4
    assert result['trade_stats']['max_consecutive_wins'] == 2
    assert result['trade_stats']['max_consecutive_losses'] == 2
    assert result['trade_stats']['profit_factor'] == '1.0000'
    assert result['anomalies'][0]['available_quantity'] == 0
    assert next(row for row in result['rounds'] if row['status'] == 'anomaly')['trade_ids'] == ['bad']


def test_review_gap_dates_are_account_scoped(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(method: str, path: str, body: dict):
            nonlocal serial
            serial += 1
            response = client.request(method, path, json=body,
                                      headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'gap-{serial}'})
            assert response.status_code == 200, response.text
            return response.json()['data']

        account = write('POST', '/api/v1/accounts', {'name': '复盘提醒'})
        other = write('POST', '/api/v1/accounts', {'name': '其他账户'})
        base = f"/api/v1/accounts/{account['id']}"
        write('POST', base + '/trades', {'trade_date': '2025-01-02', 'symbol': '600000',
                                        'side': 'buy', 'quantity': 100, 'price': '10'})
        write('PUT', base + '/snapshots/2025-01-03', {'snap_date': '2025-01-03',
                                                      'total_assets': '0', 'expected_revision': 0})
        assert client.get(base + '/review-gaps').json()['data'] == [
            {'date': '2025-01-03', 'has_trades': False, 'has_snapshot': True},
            {'date': '2025-01-02', 'has_trades': True, 'has_snapshot': False}]
        assert client.get(f"/api/v1/accounts/{other['id']}/review-gaps").json()['data'] == []
        write('PUT', base + '/daily-reviews/2025-01-02',
              {'expected_revision': 0, 'title': '已完成'})
        assert [row['date'] for row in client.get(base + '/review-gaps').json()['data']] == ['2025-01-03']
        serial += 1
        stale = client.put(base + '/daily-reviews/2025-01-02',
                           json={'expected_revision': 0, 'title': '过期草稿'},
                           headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'gap-{serial}'})
        assert stale.status_code == 409 and stale.json()['error']['code'] == 'REVISION_CONFLICT'
        assert client.get(base + '/daily-reviews/2025-01-02').json()['data']['title'] == '已完成'


def test_snapshot_position_details_revisions_and_reconciliation(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(path: str, body: dict, expected: int = 200):
            nonlocal serial
            serial += 1
            response = client.put(path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': f'positions-{serial}'})
            assert response.status_code == expected, response.text
            return response.json()

        account = client.post('/api/v1/accounts', json={'name': '快照持仓'}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'create'}).json()['data']
        base = f"/api/v1/accounts/{account['id']}"
        client.post(base + '/cash-flows', json={'flow_date': '2025-01-01',
                    'kind': 'initial', 'amount': '10000'}, headers={
                        'X-CSRF-Token': csrf, 'Idempotency-Key': 'cash'})
        path = base + '/snapshots/2025-01-02'
        positions = [{'symbol': '600000', 'name': '浦发', 'quantity': 100,
                      'market_value': '1200.00'},
                     {'symbol': '000001', 'name': '平安', 'quantity': 200,
                      'market_value': '1800.00'}]
        bad = write(path, {'snap_date': '2025-01-02', 'total_assets': '10000',
                           'available_cash': '7000', 'position_value': '2999',
                           'positions': positions, 'expected_revision': 0}, 400)
        assert bad['error']['code'] == 'POSITION_DETAIL_MISMATCH'
        saved = write(path, {'snap_date': '2025-01-02', 'total_assets': '10000',
                             'available_cash': '7000', 'positions': positions,
                             'expected_revision': 0})['data']
        assert saved['position_value'] == '3000.00'
        assert [row['symbol'] for row in saved['positions']] == ['600000', '000001']
        assert len(client.get(base + '/snapshots').json()['data'][0]['positions']) == 2
        revised = write(path, {'snap_date': '2025-01-02', 'total_assets': '10100',
                               'available_cash': '7000', 'positions': [
                                   {**positions[0], 'market_value': '3100.00'}],
                               'expected_revision': 1})['data']
        assert revised['revision'] == 2 and len(revised['positions']) == 1
        assert revised['positions'][0]['market_value'] == '3100.00'
        with client.app.state.db_factory() as session:
            audit = session.execute(text("SELECT before_json, after_json FROM audit_events WHERE entity_type='asset_snapshot' AND operation='update'")).one()
            assert len(json.loads(audit[0])['positions']) == 2
            assert len(json.loads(audit[1])['positions']) == 1


def test_real_asset_estimate_is_read_only_and_reports_quote_quality(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(method: str, path: str, body: dict):
            nonlocal serial
            serial += 1
            response = client.request(method, path, json=body,
                                      headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'est-{serial}'})
            assert response.status_code == 200, response.text
            return response.json()['data']

        account = write('POST', '/api/v1/accounts', {'name': '估算账户'})
        base = f"/api/v1/accounts/{account['id']}"
        write('POST', base + '/cash-flows', {'flow_date': '2025-01-01',
                                             'kind': 'initial', 'amount': '10000'})
        write('POST', base + '/trades', {'trade_date': '2025-01-02', 'symbol': '600000',
                                        'side': 'buy', 'quantity': 100, 'price': '10',
                                        'fee': '0', 'fee_mode': 'manual'})
        write('PUT', base + '/snapshots/2025-01-03', {'snap_date': '2025-01-03',
                                                      'total_assets': '10100', 'expected_revision': 0})
        dataset = write('POST', '/api/v1/market/datasets', {'symbol': '600000', 'bars': [
            {'event_date': '2025-01-02', 'open': '10', 'high': '10', 'low': '10',
             'close': '10', 'volume': 1000, 'available_at': '2025-01-02T15:00:00+00:00'},
            {'event_date': '2025-01-03', 'open': '11', 'high': '11', 'low': '11',
             'close': '11', 'volume': 1000, 'available_at': '2025-01-03T15:00:00+00:00'},
        ]})
        route = base + '/asset-estimate/2025-01-03'
        unknown = client.get(route, params={'decision_at': '2025-01-03T16:00:00Z'})
        assert unknown.json()['data']['total_assets'] is None
        assert 'dataset_missing' in unknown.json()['data']['quality_flags']
        known = client.get(route, params={'decision_at': '2025-01-03T16:00:00Z',
                                          'dataset_id': dataset['id']}).json()['data']
        assert known['cash'] == '9000.00' and known['known_position_value'] == '1100.00'
        assert known['total_assets'] == '10100.00' and known['difference_from_manual'] == '0.00'
        assert known['positions'][0]['quote_date'] == '2025-01-03'
        early = client.get(route, params={'decision_at': '2025-01-03T00:00:00Z',
                                          'dataset_id': dataset['id']}).json()['data']
        assert early['positions'][0]['quote_date'] == '2025-01-02'
        assert 'stale_quote' in early['quality_flags']
        assert client.get(base + '/snapshots').json()['data'][0]['total_assets'] == '10100.00'


def test_inspiration_cards_tags_and_stable_daily_choice(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(method: str, path: str, body: dict | None = None):
            nonlocal serial
            serial += 1
            response = client.request(method, path, json=body,
                                      headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'card-{serial}'})
            assert response.status_code == 200, response.text
            return response.json()['data']

        assert write('POST', '/api/v1/insights/daily',
                     {'day': (date.today() + timedelta(days=2)).isoformat()}) is None
        first = write('POST', '/api/v1/insights/cards',
                      {'content': '守住仓位纪律', 'tags': ['纪律', ' 仓位 ', '纪律']})
        second = write('POST', '/api/v1/insights/cards',
                       {'content': '等待信号', 'tags': ['策略']})
        assert first['tags'] == ['纪律', '仓位']
        assert len(client.get('/api/v1/insights/cards').json()['data']) == 2
        day = (date.today() + timedelta(days=3)).isoformat()
        choice = write('POST', '/api/v1/insights/daily', {'day': day})
        assert choice['id'] in (first['id'], second['id'])
        write('POST', '/api/v1/insights/cards', {'content': '新卡片', 'tags': []})
        assert write('POST', '/api/v1/insights/daily', {'day': day})['id'] == choice['id']
        write('DELETE', f"/api/v1/insights/cards/{choice['id']}?expected_revision=1")
        assert write('POST', '/api/v1/insights/daily', {'day': day}) is None
        assert len(client.get('/api/v1/insights/cards').json()['data']) == 2


def test_round_note_survives_trade_revision_with_link_warning(tmp_path: Path) -> None:
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(method: str, path: str, body: dict):
            nonlocal serial
            serial += 1
            response = client.request(method, path, json=body,
                                      headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'round-note-{serial}'})
            assert response.status_code == 200, response.text
            return response.json()['data']

        account = write('POST', '/api/v1/accounts', {'name': '回合摘要'})
        other = write('POST', '/api/v1/accounts', {'name': '其他账户'})
        base = f"/api/v1/accounts/{account['id']}"
        buy = write('POST', base + '/trades', {'trade_date': '2025-01-02', 'symbol': '600000',
                                              'side': 'buy', 'quantity': 100, 'price': '10'})
        write('POST', base + '/trades', {'trade_date': '2025-01-03', 'symbol': '600000',
                                        'side': 'sell', 'quantity': 100, 'price': '11'})
        assert process_one(app.state.db_factory)
        round_id = 'round:' + buy['id']
        note = write('PUT', base + '/round-notes/' + round_id,
                     {'expected_revision': 0, 'summary': '按计划止盈'})
        assert note['association_changed'] is False
        assert len(note['linked_trade_ids']) == 2
        unrelated = client.get(f"/api/v1/accounts/{other['id']}/round-notes/{round_id}")
        assert unrelated.status_code == 404
        write('POST', base + '/trades', {'trade_date': '2025-01-02', 'symbol': '600000',
                                        'side': 'buy', 'quantity': 100, 'price': '10'})
        assert process_one(app.state.db_factory)
        changed = client.get(base + '/round-notes/' + round_id).json()['data']
        assert changed['summary'] == '按计划止盈'
        assert changed['association_changed'] is True
        assert len(changed['current_trade_ids']) == 3
        updated = write('PUT', base + '/round-notes/' + round_id,
                        {'expected_revision': 1, 'summary': '补仓后需继续观察'})
        assert updated['association_changed'] is False
        assert len(updated['linked_trade_ids']) == 3


def test_security_and_invalid_nav(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        assert client.get('/api/v1/accounts').status_code == 401
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        headers = {'X-CSRF-Token': csrf, 'Idempotency-Key': 'account'}
        account = client.post('/api/v1/accounts', json={'name': 'Test'}, headers=headers).json()['data']
        root = f"/api/v1/accounts/{account['id']}"
        assert client.post(root + '/cash-flows', json={
            'flow_date': '2025-01-01', 'kind': 'withdraw', 'amount': '100'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'bad-flow'}).status_code == 400
        assert client.post(root + '/cash-flows', json={
            'flow_date': '2025-01-01', 'kind': 'initial', 'amount': '100'},
            headers={'Idempotency-Key': 'missing-csrf'}).status_code == 403
        assert client.get('/api/v1/accounts', headers={'Origin': 'https://evil.example'}).status_code == 403


def test_backup_restore_new_directory_and_single_instance(tmp_path: Path) -> None:
    source = tmp_path / 'source'
    with started_client(create_app(source, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        headers = {'X-CSRF-Token': csrf, 'Idempotency-Key': 'account'}
        account = client.post('/api/v1/accounts', json={'name': '恢复测试'}, headers=headers).json()['data']
        downloaded = client.get('/api/v1/backups/export')
        assert downloaded.status_code == 200
        archive = downloaded.content
        assert verify_backup(archive)['format'] == 'trade-rebuild-backup-v3'
        competing = InstanceLock(source)
        try:
            competing.acquire()
        except RuntimeError:
            pass
        else:
            competing.release()
            raise AssertionError('second instance acquired the same data directory')
    destination = tmp_path / 'restored'
    restore_to_new_directory(archive, destination)
    with started_client(create_app(destination, auto_rebuild=False)) as client:
        client.get('/api/v1/session')
        assert client.get('/api/v1/accounts').json()['data'][0]['id'] == account['id']
    try:
        restore_to_new_directory(archive, destination)
    except FileExistsError:
        pass
    else:
        raise AssertionError('restore overwrote populated destination')
    with zipfile.ZipFile(io.BytesIO(archive)) as current:
        db_bytes = current.read('trade.sqlite')
    legacy_manifest = {'format': 'trade-rebuild-backup-v1', 'database': 'trade.sqlite',
                       'sha256': hashlib.sha256(db_bytes).hexdigest(), 'bytes': len(db_bytes)}
    legacy_buffer = io.BytesIO()
    with zipfile.ZipFile(legacy_buffer, 'w') as legacy:
        legacy.writestr('manifest.json', json.dumps(legacy_manifest))
        legacy.writestr('trade.sqlite', db_bytes)
    assert verify_backup(legacy_buffer.getvalue())['format'] == 'trade-rebuild-backup-v1'


def test_period_notes_and_iso_week_boundary(tmp_path: Path) -> None:
    assert bounds('weekly', '2026-W01') == ('2025-12-29', '2026-01-04')
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        account = client.post('/api/v1/accounts', json={'name': '周期账户'},
                              headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'period-account'}).json()['data']
        root = f"/api/v1/accounts/{account['id']}"
        path = root + '/period-reviews/weekly/2026-W01'
        response = client.put(path, json={'expected_revision': 0,
                                         'sections': {'right_things': '控制仓位', 'core_goals': '复盘'}},
                              headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'week-1'})
        assert response.status_code == 200, response.text
        assert response.json()['data']['revision'] == 1
        derived = response.json()['data']['derived']
        assert derived['trade_count'] == 0
        assert derived['trades'] == [] and derived['rounds'] == []
        assert derived['closed_rounds'] == 0
        assert derived['return_pct'] is None
        assert derived['return_quality'] == 'missing_snapshot_or_baseline'
        assert derived['node_events'] == []
        assert client.get(path).json()['data']['sections']['right_things'] == '控制仓位'
        history = client.get(root + '/period-reviews/weekly').json()['data']
        assert [item['period_key'] for item in history] == ['2026-W01']
        bad = client.put(path, json={'expected_revision': 0, 'sections': {'right_things': '重写'}},
                         headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'week-conflict'})
        assert bad.status_code == 409
        delete_headers = {'X-CSRF-Token': csrf, 'Idempotency-Key': 'week-delete-stale'}
        assert client.delete(path + '?expected_revision=0', headers=delete_headers).status_code == 409
        delete_headers['Idempotency-Key'] = 'week-delete'
        assert client.delete(path + '?expected_revision=1', headers=delete_headers).status_code == 200
        assert client.get(root + '/period-reviews/weekly').json()['data'] == []
        assert client.get(path).json()['data']['revision'] == 0
        month = root + '/period-reviews/monthly/2026-01'
        month_response = client.put(month, json={'expected_revision': 0,
                                                 'sections': {'summary': '月度总结'}},
                                    headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'month-1'})
        assert month_response.status_code == 200, month_response.text
        assert client.get(root + '/period-reviews/monthly').json()['data'][0]['period_key'] == '2026-01'
        for number, side in enumerate(('buy', 'sell'), 1):
            response = client.post(root + '/trades', json={
                'trade_date': f'2026-01-0{number + 1}', 'symbol': '600000',
                'side': side, 'quantity': 100, 'price': str(10 + number)},
                headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'period-trade-{number}'})
            assert response.status_code == 200, response.text
        assert process_one(client.app.state.db_factory)
        updated = client.get(month).json()['data']['derived']
        assert updated['trade_count'] == 2 and len(updated['trades']) == 2
        assert updated['closed_rounds'] == 1 and len(updated['rounds']) == 1
        assert updated['rounds'][0]['symbol'] == '600000'
        assert client.get(root + '/period-reviews/weekly/2025-W53').status_code == 400


def test_review_markdown_export_includes_saved_content_and_respects_account(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        headers = {'X-CSRF-Token': csrf, 'Idempotency-Key': 'md-account'}
        account = client.post('/api/v1/accounts', json={'name': 'Markdown 账户'},
                              headers=headers).json()['data']
        root = f"/api/v1/accounts/{account['id']}"
        day = '2026-01-01'
        headers['Idempotency-Key'] = 'md-daily'
        saved = client.put(root + f'/daily-reviews/{day}', json={
            'expected_revision': 0, 'title': '元旦复盘', 'overall_summary': '只读导出正文',
            'tags': ['纪律'], 'next_watchlist': [{'code': '600000', 'name': '浦发',
                                              'condition': '突破', 'action': '观察'}]}, headers=headers)
        assert saved.status_code == 200, saved.text
        headers['Idempotency-Key'] = 'md-week'
        assert client.put(root + '/period-reviews/weekly/2026-W01', json={
            'expected_revision': 0, 'sections': {'core_goals': '控制回撤'}}, headers=headers).status_code == 200
        headers['Idempotency-Key'] = 'md-month'
        assert client.put(root + '/period-reviews/monthly/2026-01', json={
            'expected_revision': 0, 'sections': {'summary': '月度观察'}}, headers=headers).status_code == 200
        daily = client.get(root + f'/exports/review/daily/{day}.md')
        assert daily.status_code == 200
        assert 'text/markdown' in daily.headers['content-type']
        assert 'attachment' in daily.headers['content-disposition']
        assert '元旦复盘' in daily.text and '只读导出正文' in daily.text
        assert '600000' in daily.text and '纪律' in daily.text
        weekly = client.get(root + '/exports/review/weekly/2026-W01.md')
        assert '控制回撤' in weekly.text and '2025-12-29 至 2026-01-04' in weekly.text
        monthly = client.get(root + '/exports/review/monthly/2026-01.md')
        assert '月度观察' in monthly.text
        for kind, key in (('daily', day), ('weekly', '2026-W01'), ('monthly', '2026-01')):
            pdf = client.get(root + f'/exports/review/{kind}/{key}.pdf')
            assert pdf.status_code == 200, pdf.text[:200] if pdf.status_code != 200 else ''
            assert pdf.headers['content-type'] == 'application/pdf'
            assert pdf.content.startswith(b'%PDF-') and len(pdf.content) > 2000
        baseline_pdf_bytes = len(client.get(root + f'/exports/review/daily/{day}.pdf').content)
        picture = io.BytesIO()
        Image.new('RGB', (320, 180), '#16876b').save(picture, format='PNG')
        headers['Idempotency-Key'] = 'md-picture'
        uploaded = client.post(root + f'/daily-reviews/{day}/attachments',
                               files={'file': ('review.png', picture.getvalue(), 'image/png')},
                               headers=headers)
        assert uploaded.status_code == 200, uploaded.text
        pictured_pdf = client.get(root + f'/exports/review/daily/{day}.pdf')
        assert pictured_pdf.status_code == 200, pictured_pdf.text[:200] if pictured_pdf.status_code != 200 else ''
        assert len(pictured_pdf.content) > baseline_pdf_bytes
        assert b'/Subtype /Image' in pictured_pdf.content
        other_headers = {'X-CSRF-Token': csrf, 'Idempotency-Key': 'md-other'}
        other = client.post('/api/v1/accounts', json={'name': '其他账户'},
                            headers=other_headers).json()['data']
        assert client.get(f"/api/v1/accounts/{other['id']}/exports/review/daily/{day}.md").status_code == 404
        assert client.get(f"/api/v1/accounts/{other['id']}/exports/review/daily/{day}.pdf").status_code == 404
        assert client.get(root + '/exports/review/weekly/2026-W02.md').status_code == 404


def test_akshare_online_sync_freezes_units_incremental_and_unknown_availability(tmp_path: Path, monkeypatch) -> None:
    rows = [
        {'日期': '2026-01-02', '股票代码': '600000', '开盘': 10, '最高': 11,
         '最低': 9, '收盘': 10.5, '成交量': 12, '成交额': 12000},
        {'日期': '2026-01-05', '股票代码': '600000', '开盘': 10.5, '最高': 12,
         '最低': 10, '收盘': 11, '成交量': 15, '成交额': 16500},
    ]
    calls = []

    def fake_fetch(code, start, end):
        calls.append((code, start.isoformat(), end.isoformat()))
        return [row for row in rows if start.isoformat() <= row['日期'] <= end.isoformat()]

    monkeypatch.setattr(akshare_online, 'fetch_akshare_rows', fake_fetch)
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        def sync(mode: str, key: str):
            return client.post('/api/v1/market/akshare-sync', json={
                'symbol': 'sh600000', 'start_date': '2026-01-01',
                'end_date': '2026-01-08', 'mode': mode}, headers={
                    'X-CSRF-Token': csrf, 'Idempotency-Key': key})

        first = sync('full', 'ak-first')
        assert first.status_code == 200, first.text
        first_data = first.json()['data']
        assert first_data['fetched_count'] == 2 and first_data['changed'] is True
        dataset_id = first_data['dataset']['id']
        before_repeat_calls = len(calls)
        assert sync('full', 'ak-first').json()['data']['dataset']['id'] == dataset_id
        assert len(calls) == before_repeat_calls
        frozen = client.get('/api/v1/market/datasets/' + dataset_id).json()['data']
        assert frozen['provider'] == 'akshare_online'
        assert frozen['bars'][0]['volume'] == 1200
        assert frozen['bars'][0]['amount'] == '12000.00'
        assert frozen['bars'][0]['available_at'] is None
        assert frozen['source']['upstream_volume_unit'] == 'hands'
        assert sync('full', 'ak-repeat').json()['data']['dataset']['id'] == dataset_id
        rows.append({'日期': '2026-01-06', '股票代码': '600000', '开盘': 11,
                     '最高': 12, '最低': 10.5, '收盘': 11.5,
                     '成交量': 20, '成交额': 23000})
        updated = sync('incremental', 'ak-incremental')
        assert updated.status_code == 200, updated.text
        next_id = updated.json()['data']['dataset']['id']
        assert next_id != dataset_id and updated.json()['data']['changed'] is True
        assert client.get('/api/v1/market/datasets/' + next_id).json()['data']['bar_count'] == 3
        assert client.get('/api/v1/market/datasets/' + dataset_id).json()['data']['bar_count'] == 2
        assert sync('incremental', 'ak-same').json()['data']['dataset']['id'] == next_id
        assert calls[-1][1] == '2026-01-01'
        rows[0]['股票代码'] = '000001'
        rejected = sync('full', 'ak-invalid-symbol')
        assert rejected.status_code == 400
        assert rejected.json()['error']['code'] == 'MARKET_SYMBOL_MISMATCH'
        rows.clear()
        no_update = sync('incremental', 'ak-empty')
        assert no_update.status_code == 200
        assert no_update.json()['data']['dataset']['id'] == next_id
        assert no_update.json()['data']['fetched_count'] == 0
        conflict = client.post('/api/v1/market/akshare-sync', json={
            'symbol': 'sh600000', 'start_date': '2026-01-02',
            'end_date': '2026-01-08', 'mode': 'full'}, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': 'ak-first'})
        assert conflict.status_code == 409
        assert conflict.json()['error']['code'] == 'IDEMPOTENCY_CONFLICT'
        assert len(client.get('/api/v1/market/datasets').json()['data']) == 2


def test_baostock_online_sync_preserves_share_volume_and_provider_identity(tmp_path: Path, monkeypatch) -> None:
    def fake_fetch(symbol, start, end, data_dir):
        assert symbol == 'sh600000'
        assert start.isoformat() == '2026-01-01'
        assert end.isoformat() == '2026-01-08'
        return [
            {'date': '2026-01-02', 'code': 'sh.600000', 'open': '10', 'high': '11',
             'low': '9', 'close': '10.5', 'volume': '1200', 'amount': '12000'},
            {'date': '2026-01-05', 'code': 'sh.600000', 'open': '10.5', 'high': '12',
             'low': '10', 'close': '11', 'volume': '1500', 'amount': '16500'},
            {'date': '2026-01-06', 'code': 'sh.600000', 'open': '', 'high': '',
             'low': '', 'close': '', 'volume': '', 'amount': '', 'tradestatus': '0'},
        ]

    monkeypatch.setattr(baostock_online, 'fetch_baostock_rows', fake_fetch)
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        payload = {'provider': 'baostock', 'symbol': 'sh600000',
                   'start_date': '2026-01-01', 'end_date': '2026-01-08', 'mode': 'full'}
        response = client.post('/api/v1/market/online-sync', json=payload,
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'bao-one'})
        assert response.status_code == 200, response.text
        dataset_id = response.json()['data']['dataset']['id']
        dataset = client.get('/api/v1/market/datasets/' + dataset_id).json()['data']
        assert dataset['provider'] == 'baostock_online'
        assert dataset['bars'][0]['volume'] == 1200
        assert dataset['bar_count'] == 2
        assert dataset['bars'][0]['available_at'] is None
        assert dataset['source']['upstream_volume_unit'] == 'shares'
        assert dataset['source']['format'] == baostock_online.ADAPTER_VERSION
        assert client.post('/api/v1/market/online-sync', json=payload,
                           headers={'X-CSRF-Token': csrf,
                                    'Idempotency-Key': 'bao-two'}).json()['data']['dataset']['id'] == dataset_id


def test_online_sync_auto_falls_back_and_reports_actual_provider(tmp_path: Path, monkeypatch) -> None:
    def unavailable(*_args):
        raise TradeError('MARKET_PROVIDER_FAILED', 'BaoStock 暂不可用', 503)

    def ak_rows(_code, _start, _end):
        return [
            {'日期': '2026-01-02', '股票代码': '600000', '开盘': 10, '最高': 11,
             '最低': 9, '收盘': 10.5, '成交量': 12, '成交额': 12000},
            {'日期': '2026-01-05', '股票代码': '600000', '开盘': 10.5, '最高': 12,
             '最低': 10, '收盘': 11, '成交量': 15, '成交额': 16500},
        ]

    monkeypatch.setattr(baostock_online, 'fetch_baostock_rows', unavailable)
    monkeypatch.setattr(akshare_online, 'fetch_akshare_rows', ak_rows)
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        result = client.post('/api/v1/market/online-sync', json={
            'provider': 'auto', 'symbol': 'sh600000', 'start_date': '2026-01-01',
            'end_date': '2026-01-08', 'mode': 'full'}, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': 'auto-fallback'})
        assert result.status_code == 200, result.text
        data = result.json()['data']
        assert data['actual_provider'] == 'akshare'
        assert data['attempted_providers'] == ['baostock', 'akshare']
        assert data['fallback_errors'] == [{'provider': 'baostock',
                                            'code': 'MARKET_PROVIDER_FAILED'}]
        assert data['dataset']['provider'] == 'akshare_online'
        reversed_order = client.post('/api/v1/market/online-sync', json={
            'provider': 'auto', 'provider_order': ['akshare', 'baostock'],
            'symbol': 'sh600000', 'start_date': '2026-01-01',
            'end_date': '2026-01-08', 'mode': 'full'}, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': 'auto-reversed'})
        assert reversed_order.status_code == 200, reversed_order.text
        assert reversed_order.json()['data']['attempted_providers'] == ['akshare']
        invalid = client.post('/api/v1/market/online-sync', json={
            'provider': 'auto', 'provider_order': ['akshare', 'akshare'],
            'symbol': 'sh600000', 'start_date': '2026-01-01',
            'end_date': '2026-01-08', 'mode': 'full'}, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': 'auto-invalid-order'})
        assert invalid.status_code == 422


def test_market_batch_sync_progress_failure_cancel_and_restart(tmp_path: Path, monkeypatch) -> None:
    def fake_fetch(symbol, _start, _end, _data_dir):
        if symbol == 'sz000001':
            raise TradeError('MARKET_PROVIDER_FAILED', '测试来源失败', 503)
        return [
            {'date': '2026-01-02', 'code': 'sh.600000', 'open': '10', 'high': '11',
             'low': '9', 'close': '10.5', 'volume': '1200', 'amount': '12000'},
            {'date': '2026-01-05', 'code': 'sh.600000', 'open': '10.5', 'high': '12',
             'low': '10', 'close': '11', 'volume': '1500', 'amount': '16500'},
        ]

    monkeypatch.setattr(baostock_online, 'fetch_baostock_rows', fake_fetch)
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        def post(path, body, key):
            return client.post(path, json=body, headers={'X-CSRF-Token': csrf,
                                                         'Idempotency-Key': key})
        body = {'symbols': ['sh600000', 'sh600000', 'sz000001'], 'provider': 'baostock',
                'start_date': '2026-01-01', 'end_date': '2026-01-08', 'mode': 'full'}
        response = post('/api/v1/market/sync-jobs', body, 'batch-create')
        assert response.status_code == 200, response.text
        job = response.json()['data']
        assert job['total'] == 2 and job['state'] == 'queued'
        assert post('/api/v1/market/sync-jobs', body, 'batch-create').json()['data']['id'] == job['id']
        assert process_one_sync_symbol(app.state.db_factory, tmp_path)
        progress = client.get('/api/v1/market/sync-jobs/' + job['id']).json()['data']
        assert progress['completed'] == 1 and progress['state'] == 'queued'
        assert progress['results'][0]['dataset_id']
        with app.state.db_factory.begin() as session:
            session.get(MarketSyncJob, job['id']).state = 'running'
    restarted = create_app(tmp_path, auto_rebuild=False)
    with started_client(restarted) as client:
        client.get('/api/v1/session')
        assert client.get('/api/v1/market/sync-jobs/' + job['id']).json()['data']['state'] == 'queued'
        assert process_one_sync_symbol(restarted.state.db_factory, tmp_path)
        completed = client.get('/api/v1/market/sync-jobs/' + job['id']).json()['data']
        assert completed['state'] == 'partial_failed' and completed['completed'] == 2
        assert completed['results'][1]['errors'][0]['code'] == 'MARKET_PROVIDER_FAILED'
        assert len(client.get('/api/v1/market/datasets').json()['data']) == 1
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        retry = client.post('/api/v1/market/sync-jobs/' + job['id'] + '/retry-failed',
                            json={}, headers={'X-CSRF-Token': csrf,
                                              'Idempotency-Key': 'batch-retry-failed'})
        assert retry.status_code == 200, retry.text
        retried = retry.json()['data']
        assert retried['id'] != job['id'] and retried['symbols'] == ['sz000001']
        assert client.post('/api/v1/market/sync-jobs/' + job['id'] + '/retry-failed',
                           json={}, headers={'X-CSRF-Token': csrf,
                                             'Idempotency-Key': 'batch-retry-failed'}).json()['data']['id'] == retried['id']
        assert process_one_sync_symbol(restarted.state.db_factory, tmp_path)
        assert client.get('/api/v1/market/sync-jobs/' + retried['id']).json()['data']['state'] == 'partial_failed'
        assert client.get('/api/v1/market/sync-jobs/' + job['id']).json()['data']['completed'] == 2
        queued = client.post('/api/v1/market/sync-jobs', json={**body, 'symbols': ['sh600000']},
                             headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'batch-cancel'})
        cancelled = client.post('/api/v1/market/sync-jobs/' + queued.json()['data']['id'] + '/cancel',
                                json={}, headers={'X-CSRF-Token': csrf,
                                                  'Idempotency-Key': 'batch-cancel-action'})
        assert cancelled.json()['data']['state'] == 'cancelled'
        assert not process_one_sync_symbol(restarted.state.db_factory, tmp_path)


def test_market_batch_sync_does_not_fallback_after_local_save_failure(tmp_path: Path, monkeypatch) -> None:
    attempted = []
    def prepare_first(_session, _data_dir, _body):
        attempted.append('first')
        return {}
    def save_first(_session, _data_dir, _prepared):
        raise OSError('disk failure')
    def prepare_second(_session, _data_dir, _body):
        attempted.append('second')
        return {}
    monkeypatch.setattr(sync_jobs, '_PROVIDERS', {
        'baostock': (prepare_first, save_first),
        'akshare': (prepare_second, save_first),
    })
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        response = client.post('/api/v1/market/sync-jobs', json={
            'symbols': ['sh600000'], 'provider': 'auto',
            'start_date': '2026-01-01', 'end_date': '2026-01-08'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'batch-save-fail'})
        assert response.status_code == 200, response.text
        assert process_one_sync_symbol(app.state.db_factory, tmp_path)
        result = client.get('/api/v1/market/sync-jobs/' + response.json()['data']['id']).json()['data']
        assert result['state'] == 'partial_failed'
        assert result['results'][0]['errors'][0]['code'] == 'LOCAL_STORAGE_FAILED'
        assert attempted == ['first']


def test_four_step_funnel_matches_legacy_filters_and_preserves_missing_turnover() -> None:
    base = {'dataset_id': 'a' * 64, 'as_of_date': '2026-01-08', 'score': 75,
            'ret40': 0.28, 'turnover20': 0.08, 'amount20': 8e8,
            'amplitude20': 0.05, 'retrace20': 0.12, 'pullback_days': 2,
            'ma10_above_ma20_days': 8, 'ma5_above_ma10_days': 6,
            'price_vs_ma20': 0.05, 'vol_slope20': 0.08,
            'up_down_volume_ratio': 1.35, 'pullback_volume_ratio': 0.8,
            'has_blowoff_top': False, 'has_divergence_5d': False,
            'has_upper_shadow_risk': False, 'ai_confidence': 0.7,
            'theme_stage': '发酵中', 'trend_class': 'A'}
    candidates = [ScreenerCandidate(**{**base, 'dataset_id': '1' * 64, 'symbol': 'sh600001', 'ret40': 0.9,
                                        'turnover20': None, 'quality_flags': ['FLOAT_SHARES_NOT_FOUND']}),
                  ScreenerCandidate(**{**base, 'dataset_id': '2' * 64, 'symbol': 'sh600002', 'ret40': 0.8}),
                  ScreenerCandidate(**{**base, 'dataset_id': '3' * 64, 'symbol': 'sh600003', 'ret40': 0.7,
                                        'has_blowoff_top': True}),
                  ScreenerCandidate(**{**base, 'dataset_id': '4' * 64, 'symbol': 'sh600004', 'ret40': 0.6,
                                        'retrace20': 0.26})]
    config = FunnelConfig.model_validate({'step1': {'rank_start': 1, 'top_n': 10}})
    actual = run_funnel(candidates, config)
    legacy_rows = [ScreenerResult(**{**row.model_dump(), 'latest_price': 10,
                                     'day_change': 0, 'day_change_pct': 0,
                                     'turnover20': row.turnover20 or 0,
                                     'stage': 'Mid', 'labels': [], 'reject_reasons': []})
                   for row in candidates]
    legacy_config = ScreenerStepConfigs.model_validate({
        **config.model_dump(exclude={'mode'}), 'step1': {**config.step1.model_dump(), 'top_n': 10}})
    legacy = InMemoryStore._run_screener_filters_for_backtest(
        legacy_rows, mode='strict', step_configs=legacy_config)
    for stage, old_rows in zip(('step1', 'step2', 'step3', 'step4'), legacy):
        assert [item['symbol'] for item in actual['pools'][stage]] == [item.symbol for item in old_rows]
    assert actual['summary'] == {'input': 4, 'step1': 3, 'step2': 2,
                                 'step3': 1, 'step4': 1}
    assert 'sh600001' not in [item['symbol'] for item in actual['pools']['step1']]
    assert actual['rejections']['1' * 64]['reasons'] == ['TURNOVER_MISSING']
    loose = FunnelConfig.model_validate({'mode': 'loose', 'step1': {'top_n': 10},
                                          'step3': {'allow_blowoff_top': True},
                                          'step4': {'final_top_n': 2}})
    loose_actual = run_funnel(candidates, loose)
    loose_legacy_config = ScreenerStepConfigs.model_validate(loose.model_dump(exclude={'mode'}))
    loose_expected = InMemoryStore._run_screener_filters_for_backtest(
        legacy_rows, mode='loose', step_configs=loose_legacy_config)
    for stage, old_rows in zip(('step1', 'step2', 'step3', 'step4'), loose_expected):
        assert [item['symbol'] for item in loose_actual['pools'][stage]] == [item.symbol for item in old_rows]


def test_frozen_screener_metrics_match_legacy_tdx_formula_without_faking_missing_fields() -> None:
    bars = []
    for index in range(300):
        day = (date(2025, 1, 1) + timedelta(days=index)).isoformat()
        close = 10 + index * 0.035 + (index % 7) * 0.02
        bars.append({'event_date': day, 'open': f'{close - 0.08:.4f}',
                     'high': f'{close + 0.35:.4f}', 'low': f'{close - 0.4:.4f}',
                     'close': f'{close:.4f}', 'volume': 1_000_000 + index * 1800,
                     'amount': f'{close * (1_000_000 + index * 1800):.4f}',
                     'available_at': None})
    dataset = {'id': 'b' * 64, 'symbol': 'sh600000', 'bar_count': len(bars), 'bars': bars}
    cutoff = bars[-10]['event_date']
    actual = build_candidate(dataset, cutoff, 40, float_shares=20_000_000,
                             float_shares_as_of_date=cutoff)
    assert actual is not None and actual.as_of_date == cutoff
    old_series = {'symbol': 'sh600000', 'total_bars': len(bars),
                  'dates': [item['event_date'] for item in bars],
                  **{key: [float(item[key]) for item in bars]
                     for key in ('open', 'high', 'low', 'close', 'amount')},
                  'volume': [item['volume'] for item in bars]}
    expected = legacy_build_screener_row(old_series, 40, 20_000_000, cutoff)
    assert expected is not None
    for field in ('score', 'ret40', 'turnover20', 'amount20', 'amplitude20',
                  'retrace20', 'pullback_days', 'ma10_above_ma20_days',
                  'ma5_above_ma10_days', 'price_vs_ma20', 'vol_slope20',
                  'up_down_volume_ratio', 'pullback_volume_ratio',
                  'has_blowoff_top', 'has_divergence_5d', 'has_upper_shadow_risk',
                  'ai_confidence', 'theme_stage', 'trend_class'):
        assert getattr(actual, field) == getattr(expected, field), field
    assert 'HISTORICAL_AVAILABLE_AT_UNKNOWN' in actual.quality_flags
    without_amount = {**dataset, 'bars': [{**item, 'amount': None} for item in bars]}
    missing = build_candidate(without_amount, cutoff, 40)
    assert missing is not None and missing.amount20 is None and missing.turnover20 is None
    assert missing.degraded and 'FLOAT_SHARES_NOT_FOUND' in missing.quality_flags
    future_shares = build_candidate(dataset, cutoff, 40, float_shares=20_000_000,
                                    float_shares_as_of_date=bars[-1]['event_date'])
    assert future_shares is not None and future_shares.turnover20 is None
    assert 'FLOAT_SHARES_FROM_FUTURE' in future_shares.quality_flags
    future_bar = {**bars[-10], 'available_at': '2026-01-01T00:00:00+08:00'}
    delayed_dataset = {**dataset, 'bars': [*bars[:-10], future_bar, *bars[-9:]]}
    delayed = build_candidate(delayed_dataset, cutoff, 40, float_shares=20_000_000,
                              float_shares_as_of_date=cutoff)
    assert delayed is not None and delayed.as_of_date == bars[-11]['event_date']
    assert 'BARS_AFTER_DECISION_EXCLUDED' in delayed.quality_flags


def test_four_step_screener_run_freezes_universe_and_survives_restart(tmp_path: Path) -> None:
    bars = []
    for index in range(300):
        close = 10 + index * 0.04
        bars.append({'event_date': (date(2025, 1, 1) + timedelta(days=index)).isoformat(),
                     'open': f'{close - 0.1:.4f}', 'high': f'{close + 0.35:.4f}',
                     'low': f'{close - 0.4:.4f}', 'close': f'{close:.4f}',
                     'volume': 10_000_000 + index * 20_000,
                     'amount': f'{close * (10_000_000 + index * 20_000):.4f}'})
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        def post(path, body, key):
            return client.post(path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': key})
        ids = []
        for symbol in ('600000', '000001'):
            response = post('/api/v1/market/datasets', {'symbol': symbol, 'bars': bars},
                            f'screener-dataset-{symbol}')
            assert response.status_code == 200, response.text
            ids.append(response.json()['data']['id'])
        payload = {'datasets': [
            {'dataset_id': ids[0], 'float_shares': 100_000_000,
             'float_shares_as_of_date': bars[-1]['event_date']},
            {'dataset_id': ids[1]}],
            'as_of_date': bars[-1]['event_date'], 'return_window_days': 40,
            'config': {'step1': {'amount_threshold': 5e7},
                       'step2': {'retrace_min': 0, 'max_pullback_days': 20}}}
        response = post('/api/v1/research/screener-runs', payload, 'screener-create')
        assert response.status_code == 200, response.text
        run = response.json()['data']
        assert run['result']['summary']['input'] == 2
        assert run['result']['summary']['step1'] == 1
        assert run['result']['pools']['step1'][0]['dataset_id'] == ids[0]
        assert run['result']['pools']['input'][1]['turnover20'] is None
        assert 'FLOAT_SHARES_NOT_FOUND' in run['quality_flags']
        assert run['result']['metric_label'] == 'legacy_formula_candidate_confidence_not_ai_model'
        assert post('/api/v1/research/screener-runs', payload,
                    'screener-create').json()['data']['id'] == run['id']
        assert client.get('/api/v1/research/screener-runs').json()['data'][0]['id'] == run['id']
        duplicate = post('/api/v1/research/screener-runs', {
            **payload, 'datasets': [payload['datasets'][0], payload['datasets'][0]]}, 'screener-duplicate')
        assert duplicate.status_code == 400 and duplicate.json()['error']['code'] == 'DUPLICATE_DATASET'
        lenient = post('/api/v1/research/screener-runs', {**payload, 'config': {
            'step1': {'amount_threshold': 5e7, 'amplitude_threshold': 0.01},
            'step2': {'retrace_min': 0, 'retrace_max': 0.8, 'max_pullback_days': 30,
                      'min_ma10_above_ma20_days': 0, 'min_ma5_above_ma10_days': 0,
                      'max_price_vs_ma20': 0.5, 'require_above_ma20': False,
                      'allow_b_trend': True},
            'step3': {'min_vol_slope20': 0, 'min_up_down_volume_ratio': 0,
                      'max_pullback_volume_ratio': 5, 'allow_blowoff_top': True,
                      'allow_divergence_5d': True, 'allow_upper_shadow_risk': True,
                      'allow_degraded': True},
            'step4': {'min_ai_confidence': 0,
                      'allowed_theme_stages': ['发酵中', '高潮', '退潮', 'Unknown']}}}, 'screener-lenient')
        assert lenient.status_code == 200, lenient.text
        final_run = lenient.json()['data']
        assert final_run['result']['summary']['step4'] == 1
        promoted = post(f"/api/v1/research/screener-runs/{final_run['id']}/signals/{ids[0]}",
                        {}, 'screener-promote')
        assert promoted.status_code == 200, promoted.text
        signal = promoted.json()['data']
        assert signal['strategy_id'] == 'four_step_funnel_v1'
        assert signal['result']['candidate']['screener_run_id'] == final_run['id']
        assert post(f"/api/v1/research/screener-runs/{final_run['id']}/signals/{ids[1]}",
                    {}, 'screener-non-finalist').status_code == 404
        sim = post('/api/v1/sim-accounts', {
            'name': '漏斗草稿账户', 'initial_capital': '10000',
            'start_date': bars[-1]['event_date']}, 'screener-sim')
        assert sim.status_code == 200, sim.text
        draft = post(f"/api/v1/sim-accounts/{sim.json()['data']['id']}/drafts", {
            'source_run_id': signal['id'], 'quantity': 100,
            'limit_price': bars[-1]['close']}, 'screener-draft')
        assert draft.status_code == 200, draft.text
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        client.get('/api/v1/session')
        persisted = client.get('/api/v1/research/screener-runs/' + run['id']).json()['data']
        assert persisted['result']['summary'] == run['result']['summary']
        assert persisted['request']['datasets'][0]['dataset_id'] == ids[0]


def test_tdx_float_shares_lookup_records_source_without_modifying_dbf(tmp_path: Path, monkeypatch) -> None:
    root = tmp_path / 'tdx'
    folder = root / 'T0002' / 'hq_cache'
    folder.mkdir(parents=True)
    fields = [('SC', 1), ('GPDM', 6), ('LTAG', 10)]
    header_len = 32 + len(fields) * 32 + 1
    record_len = 1 + sum(length for _, length in fields)
    header = struct.pack('<BBBBIHH20x', 3, 0, 0, 0, 2, header_len, record_len)
    descriptors = []
    for name, length in fields:
        descriptor = bytearray(32)
        descriptor[:len(name)] = name.encode('ascii')
        descriptor[11] = ord('C')
        descriptor[16] = length
        descriptors.append(bytes(descriptor))
    row1 = b' ' + b'1' + b'600000' + b'     10000'
    row2 = b' ' + b'0' + b'000001' + b'      5000'
    contents = header + b''.join(descriptors) + b'\x0d' + row1 + row2
    path = folder / 'base.dbf'
    path.write_bytes(contents)
    monkeypatch.setenv('TRADE_TDX_ROOT', str(root))
    with started_client(create_app(tmp_path / 'data', auto_rebuild=False)) as client:
        client.get('/api/v1/session')
        response = client.get('/api/v1/market/tdx-float-shares', params=[
            ('symbol', 'sh600000'), ('symbol', 'sz000001'), ('symbol', 'bj830001')])
        assert response.status_code == 200, response.text
        data = response.json()['data']
        assert data['requested_values'] == {'sh600000': 100_000_000,
                                            'sz000001': 50_000_000}
        assert data['missing'] == ['bj830001']
        assert data['source']['sha256'] == hashlib.sha256(contents).hexdigest()
        assert data['source']['as_of_quality'] == 'unknown'
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        bars = [{'event_date': (date(2025, 1, 1) + timedelta(days=index)).isoformat(),
                 'open': '10', 'high': '11', 'low': '9', 'close': '10',
                 'volume': 1_000_000, 'amount': '10000000'} for index in range(300)]
        imported = client.post('/api/v1/market/datasets', json={'symbol': '600000', 'bars': bars},
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'dbf-screener-market'})
        assert imported.status_code == 200, imported.text
        run_body = {'datasets': [{'dataset_id': imported.json()['data']['id'],
                                  'float_shares': 100_000_000,
                                  'float_shares_source_sha256': data['source']['sha256']}],
                    'as_of_date': bars[-1]['event_date']}
        accepted = client.post('/api/v1/research/screener-runs', json=run_body,
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'dbf-screener-run'})
        assert accepted.status_code == 200, accepted.text
        assert accepted.json()['data']['request']['datasets'][0]['float_shares_source_sha256'] == data['source']['sha256']
        forged_value = client.post('/api/v1/research/screener-runs', json={
            **run_body, 'datasets': [{**run_body['datasets'][0], 'float_shares': 200_000_000}]},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'dbf-screener-forged'})
        assert forged_value.status_code == 409
        assert forged_value.json()['error']['code'] == 'FLOAT_SHARES_SOURCE_CHANGED'
        path.write_bytes(contents.replace(b'     10000', b'     20000'))
        rejected = client.post('/api/v1/research/screener-runs', json=run_body,
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'dbf-screener-stale'})
        assert rejected.status_code == 409
        assert rejected.json()['error']['code'] == 'FLOAT_SHARES_SOURCE_CHANGED'
        path.write_bytes(contents)
    assert path.read_bytes() == contents


def test_tdx_full_universe_job_freezes_sources_recovers_and_reports_errors(tmp_path: Path, monkeypatch) -> None:
    tdx = tmp_path / 'tdx'
    sources = {}
    for market, code in (('sh', '600000'), ('sh', '600001'), ('sz', '000001')):
        folder = tdx / 'vipdoc' / market / 'lday'
        folder.mkdir(parents=True, exist_ok=True)
        records = []
        for index in range(300):
            day = date(2025, 1, 1) + timedelta(days=index)
            close = int((10 + index * 0.04) * 100)
            volume = 10_000_000 + index * 20_000
            records.append(struct.pack('<IIIIIfII', int(day.strftime('%Y%m%d')),
                                       close - 10, close + 35, close - 40, close,
                                       float(close / 100 * volume), volume, 0))
        path = folder / f'{market}{code}.day'
        sources[path] = b''.join(records)
        path.write_bytes(sources[path])
    dbf_dir = tdx / 'T0002' / 'hq_cache'
    dbf_dir.mkdir(parents=True)
    fields = [('SC', 1), ('GPDM', 6), ('LTAG', 10)]
    header_len, record_len = 32 + len(fields) * 32 + 1, 1 + sum(length for _, length in fields)
    descriptors = []
    for name, length in fields:
        descriptor = bytearray(32)
        descriptor[:len(name)] = name.encode('ascii')
        descriptor[11], descriptor[16] = ord('C'), length
        descriptors.append(bytes(descriptor))
    dbf = (struct.pack('<BBBBIHH20x', 3, 0, 0, 0, 3, header_len, record_len)
           + b''.join(descriptors) + b'\x0d'
           + b' ' + b'1' + b'600000' + b'     10000'
           + b' ' + b'1' + b'600001' + b'     10000'
           + b' ' + b'0' + b'000001' + b'     10000')
    (dbf_dir / 'base.dbf').write_bytes(dbf)
    name_record = bytearray(360)
    name_record[:6] = b'600000'
    name_record[31:31 + len('浦发银行'.encode('gbk'))] = '浦发银行'.encode('gbk')
    (dbf_dir / 'shs.tnf').write_bytes(bytes(50) + bytes(name_record))
    monkeypatch.setenv('TRADE_TDX_ROOT', str(tdx))
    data_dir = tmp_path / 'data'
    app = create_app(data_dir, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        response = client.post('/api/v1/research/tdx-universe-jobs', json={
            'markets': ['sh', 'sz'], 'as_of_date': '2025-10-27',
            'max_bars': 360, 'return_window_days': 40,
            'config': {'step1': {'amount_threshold': 5e7},
                       'step2': {'retrace_min': 0, 'max_pullback_days': 20}}},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'universe-create'})
        assert response.status_code == 200, response.text
        job = response.json()['data']
        assert job['total'] == 3 and job['state'] == 'queued'
        assert client.post('/api/v1/research/tdx-universe-jobs', json={
            'markets': ['sh', 'sz'], 'as_of_date': '2025-10-27',
            'max_bars': 360, 'return_window_days': 40,
            'config': {'step1': {'amount_threshold': 5e7},
                       'step2': {'retrace_min': 0, 'max_pullback_days': 20}}},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'universe-create'}).json()['data']['id'] == job['id']
        assert process_one_universe_symbol(app.state.db_factory, data_dir, tdx)
        assert client.get('/api/v1/research/tdx-universe-jobs/' + job['id']).json()['data']['processed'] == 1
        with app.state.db_factory.begin() as session:
            session.get(TdxUniverseJob, job['id']).state = 'running'
        (tdx / 'vipdoc' / 'sh' / 'lday' / 'sh600001.day').unlink()
    restarted = create_app(data_dir, auto_rebuild=False)
    with started_client(restarted) as client:
        client.get('/api/v1/session')
        assert client.get('/api/v1/research/tdx-universe-jobs/' + job['id']).json()['data']['state'] == 'queued'
        assert process_one_universe_symbol(restarted.state.db_factory, data_dir, tdx)
        assert process_one_universe_symbol(restarted.state.db_factory, data_dir, tdx)
        assert process_one_universe_symbol(restarted.state.db_factory, data_dir, tdx)
        complete = client.get('/api/v1/research/tdx-universe-jobs/' + job['id']).json()['data']
        assert complete['state'] == 'partial_failed'
        assert complete['processed'] == 3 and complete['candidates'] == 2
        assert complete['recent_errors'][0]['code'] == 'TDX_DAY_NOT_FOUND'
        run = client.get('/api/v1/research/screener-runs/' + complete['run_id']).json()['data']
        assert run['result']['summary']['input'] == 2
        assert next(item for item in run['result']['pools']['input'] if item['symbol'] == '600000')['name'] == '浦发银行'
        assert run['request']['source'] == 'tdx_universe'
        assert run['request']['symbol_name_sources'][0]['file'] == 'shs.tnf'
        assert run['request']['float_shares_source']['sha256'] == hashlib.sha256(dbf).hexdigest()
        assert {item['symbol'] for item in run['result']['excluded']} == {'sh600001'}
        assert len(client.get('/api/v1/market/datasets').json()['data']) == 2
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        queued = client.post('/api/v1/research/tdx-universe-jobs', json={
            'markets': ['sz'], 'as_of_date': '2025-10-27'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'universe-cancel'})
        assert queued.status_code == 200, queued.text
        cancelled = client.post('/api/v1/research/tdx-universe-jobs/' + queued.json()['data']['id'] + '/cancel',
                                json={}, headers={'X-CSRF-Token': csrf,
                                                  'Idempotency-Key': 'universe-cancel-action'})
        assert cancelled.status_code == 200 and cancelled.json()['data']['state'] == 'cancelled'
        assert not process_one_universe_symbol(restarted.state.db_factory, data_dir, tdx)
    for path, contents in sources.items():
        if path.is_file():
            assert path.read_bytes() == contents


def test_account_excel_export_keeps_accounts_and_formula_text_separate(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']

        def post(path: str, body: dict, key: str):
            return client.post(path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': key})

        real = post('/api/v1/accounts', {'name': '工作簿账户'}, 'excel-real').json()['data']
        sim = post('/api/v1/sim-accounts', {'name': '工作簿模拟',
                   'initial_capital': '10000', 'start_date': '2026-01-02'},
                   'excel-sim').json()['data']
        assert post(f"/api/v1/accounts/{real['id']}/trades", {
            'trade_date': '2026-01-02', 'symbol': '600000', 'name': '=1+1',
            'side': 'buy', 'quantity': 100, 'price': '10', 'note': '@unsafe'},
            'excel-trade').status_code == 200
        real_file = client.get(f"/api/v1/accounts/{real['id']}/exports/account.xlsx")
        assert real_file.status_code == 200
        assert real_file.content.startswith(b'PK')
        real_book = load_workbook(io.BytesIO(real_file.content), read_only=True, data_only=False)
        assert real_book.sheetnames == ['Meta', 'Trades', 'Flows', 'Snapshots', 'Rounds', 'NAV']
        rows = list(real_book['Trades'].values)
        assert rows[0][3:5] == ('name', 'side')
        assert rows[1][3] == "'=1+1" and rows[1][3] != '=1+1'
        assert rows[1][11] == "'@unsafe"
        assert real_book['Trades']['D2'].data_type == 's'
        order = post(f"/api/v1/sim-accounts/{sim['id']}/orders", {
            'symbol': '000001', 'side': 'buy', 'quantity': 100, 'limit_price': '10',
            'signal_date': '2026-01-02', 'submit_date': '2026-01-02'}, 'excel-order').json()['data']
        assert post(f"/api/v1/sim-accounts/{sim['id']}/orders/{order['id']}/fill", {
            'expected_revision': 1, 'fill_date': '2026-01-02', 'fill_price': '9'},
            'excel-fill').status_code == 200
        sim_file = client.get(f"/api/v1/accounts/{sim['id']}/exports/account.xlsx")
        assert sim_file.status_code == 200
        sim_book = load_workbook(io.BytesIO(sim_file.content), read_only=True)
        assert sim_book.sheetnames == ['Meta', 'Orders', 'Fills', 'Positions']
        assert list(sim_book['Fills'].values)[1][3:6] == ('000001', 'buy', 100)
        assert client.get('/api/v1/accounts/00000000000000000000000000000000/exports/account.xlsx').status_code == 404


def test_account_csv_exports_are_scoped_stable_and_formula_safe(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        def post(path: str, body: dict, key: str):
            return client.post(path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': key})
        account = post('/api/v1/accounts', {'name': 'CSV 实盘'}, 'csv-account').json()['data']
        other = post('/api/v1/accounts', {'name': 'CSV 其他'}, 'csv-other').json()['data']
        root = f"/api/v1/accounts/{account['id']}"
        assert post(root + '/trades', {'trade_date': '2026-01-02', 'symbol': '600000',
                    'name': '=1+1', 'side': 'buy', 'quantity': 100, 'price': '10',
                    'note': '@unsafe'}, 'csv-trade').status_code == 200
        assert post(root + '/cash-flows', {'flow_date': '2026-01-02', 'kind': 'initial',
                    'amount': '10000'}, 'csv-flow').status_code == 200
        response = client.get(root + '/exports/trades.csv')
        assert response.status_code == 200
        assert response.content.startswith(b'\xef\xbb\xbf')
        assert 'attachment' in response.headers['content-disposition']
        rows = list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
        assert len(rows) == 1
        assert rows[0]['name'] == "'=1+1" and rows[0]['note'] == "'@unsafe"
        assert rows[0]['price'] == '10.0000' and rows[0]['quantity'] == '100'
        assert len(list(csv.DictReader(io.StringIO(client.get(root + '/exports/flows.csv').content.decode('utf-8-sig'))))) == 1
        assert client.get(root + '/exports/snapshots.csv').status_code == 200
        assert client.get(root + '/exports/rounds.csv').status_code == 200
        assert client.get(f"/api/v1/accounts/{other['id']}/exports/trades.csv").text.find('600000') == -1
        assert client.get(root + '/exports/sim-fills.csv').status_code == 404
        sim = post('/api/v1/sim-accounts', {'name': 'CSV 模拟', 'initial_capital': '1000',
                                          'start_date': '2026-01-01'}, 'csv-sim').json()['data']
        assert client.get(f"/api/v1/accounts/{sim['id']}/exports/sim-fills.csv").status_code == 200
        assert client.get(f"/api/v1/accounts/{sim['id']}/exports/trades.csv").status_code == 404


def test_tdx_day_import_freezes_source_and_marks_historical_availability_unknown(
        tmp_path: Path, monkeypatch) -> None:
    assert normalize_tdx_symbol('510300') == ('sh', '510300')
    tdx_root = tmp_path / 'tdx' / 'vipdoc' / 'sh' / 'lday'
    tdx_root.mkdir(parents=True)
    source = tdx_root / 'sh600000.day'
    record = struct.Struct('<IIIIIfII')
    original = b''.join([
        record.pack(20250102, 1000, 1100, 900, 1050, 12345.0, 1000, 0),
        record.pack(20250103, 1050, 1200, 1000, 1150, 23456.0, 1200, 0),
    ])
    source.write_bytes(original)
    minute_folder = tdx_root.parent / 'minline'
    minute_folder.mkdir()
    minute_source = minute_folder / 'sh600000.lc1'
    minute_record = struct.Struct('<HHfffffII')
    encoded_date = ((2025 - 2004) << 11) | 102
    minute_bytes = b''.join([
        minute_record.pack(encoded_date, 571, 10.0, 10.2, 9.9, 10.1, 1234.0, 100, 0),
        minute_record.pack(encoded_date, 572, 10.1, 10.3, 10.0, 10.2, 2345.0, 200, 0),
    ])
    minute_source.write_bytes(minute_bytes)
    monkeypatch.setenv('TRADE_TDX_ROOT', str(tmp_path / 'tdx'))
    with started_client(create_app(tmp_path / 'new-data', auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        endpoint = '/api/v1/market/tdx-import'
        first = client.post(endpoint, json={'symbol': 'SH600000'}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'tdx-first'})
        assert first.status_code == 200, first.text
        item = first.json()['data']
        assert item['provider'] == 'tdx_local' and item['symbol'] == '600000'
        assert item['bar_count'] == 2 and item['availability_quality'] == 'historical_availability_unknown'
        frozen = client.get('/api/v1/market/datasets/' + item['id']).json()['data']
        assert frozen['source']['sha256'] == hashlib.sha256(original).hexdigest()
        assert frozen['bars'][0]['open'] == '10.0000' and frozen['bars'][1]['close'] == '11.5000'
        assert frozen['bars'][0]['amount'] == '12345.00'
        assert frozen['bars'][0]['available_at'] is None
        close = client.get('/api/v1/market/close', params={
            'dataset_id': item['id'], 'day': '2025-01-02'})
        assert close.status_code == 200, close.text
        assert close.json()['data']['close'] == '10.5000'
        assert close.json()['data']['trading_date'] == '2025-01-02'
        assert close.json()['data']['available_at'] is None
        assert client.get('/api/v1/market/close', params={
            'dataset_id': item['id'], 'day': '2025-01-04'}).status_code == 404
        intraday = client.get('/api/v1/market/tdx-intraday', params={
            'symbol': '600000', 'day': '2025-01-02'})
        assert intraday.status_code == 200, intraday.text
        minutes = intraday.json()['data']
        assert [point['time'] for point in minutes['points']] == ['09:31', '09:32']
        assert minutes['points'][1]['close'] == '10.2000'
        assert minutes['points'][0]['amount'] == '1234.00'
        assert minutes['source']['sha256'] == hashlib.sha256(minute_bytes).hexdigest()
        assert minutes['availability_quality'] == 'historical_availability_unknown'
        assert minute_source.read_bytes() == minute_bytes
        assert client.get('/api/v1/market/tdx-intraday', params={
            'symbol': '600000', 'day': '2025-01-03'}).status_code == 404
        assert eligible_bars(frozen['bars'], '2025-01-04T00:00:00+00:00', True)[0] == []
        assert len(eligible_bars(frozen['bars'], '2025-01-04T00:00:00+00:00', False)[0]) == 2
        assert source.read_bytes() == original
        assert client.post(endpoint, json={'symbol': '../bad'}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'tdx-invalid-symbol'}).status_code in (400, 422)
        changed = original + record.pack(20250106, 1150, 1300, 1100, 1250, 34567.0, 1400, 0)
        source.write_bytes(changed)
        second = client.post(endpoint, json={'symbol': '600000'}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'tdx-second'})
        assert second.status_code == 200, second.text
        assert second.json()['data']['id'] != item['id']
        assert client.get('/api/v1/market/datasets/' + item['id']).json()['data']['bar_count'] == 2
        source.write_bytes(changed + b'bad')
        broken = client.post(endpoint, json={'symbol': '600000'}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'tdx-broken'})
        assert broken.status_code == 400
    monkeypatch.delenv('TRADE_TDX_ROOT')
    with started_client(create_app(tmp_path / 'no-tdx', auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        unavailable = client.post(endpoint, json={'symbol': '600000'}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'tdx-unconfigured'})
        assert unavailable.status_code == 409


def test_manual_stock_annotation_versions_survive_backup_and_do_not_touch_old_data(tmp_path: Path) -> None:
    source = tmp_path / 'annotation-source'
    with started_client(create_app(source, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        endpoint = '/api/v1/market/annotations/600000'
        assert client.get(endpoint).json()['data'] is None
        body = {'expected_revision': 0, 'start_date': '2025-01-02', 'stage': 'Early',
                'trend_class': 'A_B', 'decision': '保留', 'notes': '人工判断'}
        saved = client.put(endpoint, json=body, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'annotation-first'})
        assert saved.status_code == 200, saved.text
        assert saved.json()['data']['symbol'] == 'sh600000'
        assert saved.json()['data']['updated_by'] == 'manual'
        assert client.get('/api/v1/market/annotations/sh600000').json()['data']['revision'] == 1
        assert len(client.get('/api/v1/market/annotations').json()['data']) == 1
        stale = client.put(endpoint, json={**body, 'decision': '排除'}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'annotation-stale'})
        assert stale.status_code == 409
        changed = client.put(endpoint, json={**body, 'expected_revision': 1,
                                             'decision': '排除'}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'annotation-second'})
        assert changed.status_code == 200 and changed.json()['data']['revision'] == 2
        assert changed.json()['data']['decision'] == '排除'
        archive = client.get('/api/v1/backups/export').content
    destination = tmp_path / 'annotation-restored'
    restore_to_new_directory(archive, destination)
    with started_client(create_app(destination, auto_rebuild=False)) as client:
        client.get('/api/v1/session')
        assert client.get(endpoint).json()['data']['notes'] == '人工判断'
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        assert client.delete(endpoint + '?expected_revision=1', headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'annotation-delete-stale'}).status_code == 409
        deleted = client.delete(endpoint + '?expected_revision=2', headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'annotation-delete'})
        assert deleted.status_code == 200
        assert client.get(endpoint).json()['data'] is None


def test_legacy_provider_cache_csv_import_tracks_source_and_provider(
        tmp_path: Path, monkeypatch) -> None:
    cache = tmp_path / 'legacy-cache'
    cache.mkdir()
    path = cache / 'sh600000.csv'
    path.write_text('date,open,high,low,close,volume,amount,symbol\n'
                    '2025-01-02,10,11,9,10.5,1000,12345,sh600000\n'
                    '2025-01-03,10.5,12,10,11.5,1200,23456,sh600000\n', encoding='utf-8')
    monkeypatch.setenv('TRADE_AKSHARE_CACHE_DIR', str(cache))
    monkeypatch.setenv('TRADE_BAOSTOCK_CACHE_DIR', str(cache))
    with started_client(create_app(tmp_path / 'new-cache-data', auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        endpoint = '/api/v1/market/cache-import'
        headers = {'X-CSRF-Token': csrf, 'Idempotency-Key': 'cache-ak'}
        ak = client.post(endpoint, json={'provider': 'akshare', 'symbol': '600000'}, headers=headers)
        assert ak.status_code == 200, ak.text
        assert ak.json()['data']['provider'] == 'akshare_cache'
        detail = client.get('/api/v1/market/datasets/' + ak.json()['data']['id']).json()['data']
        assert detail['bars'][1]['amount'] == '23456.00'
        assert detail['source']['sha256'] == hashlib.sha256(path.read_bytes()).hexdigest()
        assert detail['source']['source_row_count'] == 2
        headers['Idempotency-Key'] = 'cache-bs'
        bs = client.post(endpoint, json={'provider': 'baostock', 'symbol': 'sh600000'}, headers=headers)
        assert bs.status_code == 200 and bs.json()['data']['provider'] == 'baostock_cache'
        assert bs.json()['data']['id'] != ak.json()['data']['id']
        frozen_id = ak.json()['data']['id']
        archive = create_backup(tmp_path / 'new-cache-data')
        assert f'market/{frozen_id}.json' in verify_backup(archive)['files']
        headers['Idempotency-Key'] = 'cache-no-file'
        assert client.post(endpoint, json={'provider': 'akshare', 'symbol': '000001'},
                           headers=headers).status_code == 404
        path.write_text(path.read_text(encoding='utf-8').replace('sh600000\n', 'sz000001\n'),
                        encoding='utf-8')
        headers['Idempotency-Key'] = 'cache-mismatch'
        assert client.post(endpoint, json={'provider': 'akshare', 'symbol': '600000'},
                           headers=headers).status_code == 400
    restored = tmp_path / 'cache-restored'
    restore_to_new_directory(archive, restored)
    with started_client(create_app(restored, auto_rebuild=False)) as client:
        client.get('/api/v1/session')
        assert client.get('/api/v1/market/datasets/' + frozen_id).json()['data']['source']['filename'] == 'sh600000.csv'


def test_nav_zero_shares_new_segment_and_price_limit() -> None:
    flows = [FlowFact('initial', '2025-01-01', 'initial', 1000000, '1'),
             FlowFact('withdraw', '2025-01-03', 'withdraw', 1100000, '2'),
             FlowFact('new', '2025-01-04', 'initial', 500000, '3')]
    snapshots = [SnapshotFact('2025-01-02', 1100000), SnapshotFact('2025-01-03', 0),
                 SnapshotFact('2025-01-04', 500000)]
    result = calculate_nav(flows, snapshots)
    assert result['points'][2]['nav'] is None
    assert result['points'][3]['segment'] == 2
    assert result['current']['nav'] == '1.00000000'
    try:
        price_units('1000000000000000')
    except TradeError as exc:
        assert exc.code == 'PRICE_TOO_LARGE'
    else:
        raise AssertionError('price overflow was accepted')


def test_superseded_projection_never_publishes_old_inputs(tmp_path: Path, monkeypatch) -> None:
    import trade_app.analytics.service as analytics_service

    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        account = client.post('/api/v1/accounts', json={'name': '版本测试'},
                              headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'account'}).json()['data']
        root = f"/api/v1/accounts/{account['id']}"
        client.post(root + '/cash-flows', json={'flow_date': '2025-01-01', 'kind': 'initial', 'amount': '10000'},
                    headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'initial'})
        original = analytics_service.project_account

        def mutate_during_calculation(flows, snapshots, trades, targets):
            with app.state.db_factory.begin() as session:
                create_trade(session, account['id'], {'trade_date': '2025-01-01', 'symbol': '600000',
                            'side': 'buy', 'quantity': 100, 'price': '10', 'fee': '0'})
            return original(flows, snapshots, trades, targets)

        monkeypatch.setattr(analytics_service, 'project_account', mutate_during_calculation)
        assert process_one(app.state.db_factory)
        status = client.get(root + '/analytics').json()['data']
        assert status['status'] == 'recalculating'
        assert status['projection_version'] is None
        monkeypatch.setattr(analytics_service, 'project_account', original)
        assert process_one(app.state.db_factory)
        status = client.get(root + '/analytics').json()['data']
        assert status['status'] == 'fresh'
        assert status['result']['positions'][0]['quantity'] == 100


def test_simulation_order_reservation_t_plus_one_and_fifo(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(path: str, body: dict | None = None, method: str = 'POST'):
            nonlocal serial
            serial += 1
            response = client.request(method, path, json=body,
                                      headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'sim-{serial}'})
            return response

        created = write('/api/v1/sim-accounts', {'name': '模拟账户', 'initial_capital': '10000',
                                                  'start_date': '2025-01-01'})
        assert created.status_code == 200, created.text
        account_id = created.json()['data']['id']
        root = f'/api/v1/sim-accounts/{account_id}'
        buy = write(root + '/orders', {'symbol': '600000', 'side': 'buy', 'quantity': 100,
                                      'limit_price': '10', 'signal_date': '2025-01-01',
                                      'submit_date': '2025-01-01'})
        assert buy.status_code == 200, buy.text
        order = buy.json()['data']
        too_large = write(root + '/orders', {'symbol': '600000', 'side': 'buy', 'quantity': 1000,
                                              'limit_price': '10', 'signal_date': '2025-01-01',
                                              'submit_date': '2025-01-01'})
        assert too_large.status_code == 400
        filled = write(root + f"/orders/{order['id']}/fill", {
            'expected_revision': 1, 'fill_date': '2025-01-01', 'fill_price': '9'})
        assert filled.status_code == 200, filled.text
        portfolio = filled.json()['data']['portfolio']
        assert portfolio['positions'][0]['quantity'] == 100
        assert portfolio['positions'][0]['sellable_quantity'] == 0
        same_day_sell = write(root + '/orders', {'symbol': '600000', 'side': 'sell', 'quantity': 100,
                                                 'limit_price': '11', 'signal_date': '2025-01-01',
                                                 'submit_date': '2025-01-01'})
        assert same_day_sell.status_code == 400
        assert write(root + '/settle', {'to_date': '2025-01-02'}).status_code == 200
        sell = write(root + '/orders', {'symbol': '600000', 'side': 'sell', 'quantity': 100,
                                       'limit_price': '11', 'signal_date': '2025-01-02',
                                       'submit_date': '2025-01-02'})
        assert sell.status_code == 200, sell.text
        sell_order = sell.json()['data']
        duplicate_sell = write(root + '/orders', {'symbol': '600000', 'side': 'sell', 'quantity': 1,
                                                  'limit_price': '11', 'signal_date': '2025-01-02',
                                                  'submit_date': '2025-01-02'})
        assert duplicate_sell.status_code == 400
        result = write(root + f"/orders/{sell_order['id']}/fill", {
            'expected_revision': 1, 'fill_date': '2025-01-02', 'fill_price': '12'})
        assert result.status_code == 200, result.text
        assert Decimal(result.json()['data']['fill']['realized_pnl']) > 0
        assert result.json()['data']['portfolio']['positions'] == []
        assert client.get(root + '/fills').json()['data'][-1]['price_source'] == 'manual'
        performance = client.get(root + '/performance').json()['data']
        assert performance['buy_fill_count'] == 1
        assert performance['sell_fill_count'] == 1
        assert performance['win_rate_pct'] == '100.00'
        assert Decimal(performance['realized_pnl']) > 0
        assert performance['monthly'][0]['month'] == '2025-01'
        assert performance['closed_fills'][0]['buy_allocations'][0]['buy_date'] == '2025-01-01'
        assert performance['closed_fills'][0]['buy_allocations'][0]['quantity'] == 100
        assert performance['realized_curve'][0]['cumulative_realized_pnl'] == performance['realized_pnl']
        assert client.get(f'/api/v1/accounts/{account_id}/trades').status_code == 404
        before_reset = client.get(root + '/portfolio').json()['data']
        reset = write(root + '/reset', {
            'expected_wallet_revision': before_reset['wallet_revision'],
            'start_date': '2025-02-01', 'reset_config': False})
        assert reset.status_code == 200, reset.text
        successor_id = reset.json()['data']['new_account']['id']
        successor_root = f'/api/v1/sim-accounts/{successor_id}'
        assert reset.json()['data']['history_preserved'] is True
        assert client.get(root + '/portfolio').json()['data']['frozen'] is True
        assert len(client.get(root + '/fills').json()['data']) == 2
        assert client.get(root + '/performance').json()['data']['sell_fill_count'] == 1
        assert client.get(successor_root + '/portfolio').json()['data']['cash'] == '10000.00'
        assert client.get(successor_root + '/orders').json()['data'] == []
        assert write(root + '/orders', {'symbol': '600000', 'side': 'buy', 'quantity': 100,
                     'limit_price': '10', 'signal_date': '2025-01-02',
                     'submit_date': '2025-01-02'}).status_code == 409
        old_revision = client.get(root + '/portfolio').json()['data']['wallet_revision']
        activated = write(root + '/activate-recovery', {'expected_wallet_revision': old_revision})
        assert activated.status_code == 200, activated.text
        assert activated.json()['data']['active_account_id'] == account_id
        assert client.get(root + '/portfolio').json()['data']['frozen'] is False
        assert client.get(successor_root + '/portfolio').json()['data']['frozen'] is True
        assert write(successor_root + '/settle', {'to_date': '2025-02-02'}).status_code == 409
        accounts = client.get('/api/v1/accounts').json()['data']
        assert next(item for item in accounts if item['id'] == successor_id)['frozen'] is True


def test_sim_fill_review_tags_account_scope_history_and_realized_stats(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(path: str, body: dict | None = None, method: str = 'POST'):
            nonlocal serial
            serial += 1
            return client.request(method, path, json=body,
                headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'tags-{serial}'})

        account = write('/api/v1/sim-accounts', {'name': '标签账户',
            'initial_capital': '10000', 'start_date': '2025-01-01'}).json()['data']
        root = f"/api/v1/sim-accounts/{account['id']}"
        other = write('/api/v1/sim-accounts', {'name': '隔离账户',
            'initial_capital': '10000', 'start_date': '2025-01-01'}).json()['data']
        other_root = f"/api/v1/sim-accounts/{other['id']}"
        emotion = write(root + '/review-tags', {'type': 'emotion', 'name': '冷静'}).json()['data']
        reason = write(root + '/review-tags', {'type': 'reason', 'name': '突破'}).json()['data']
        extra = write(root + '/review-tags', {'type': 'reason', 'name': '量能'}).json()['data']
        assert write(root + '/review-tags', {'type': 'reason', 'name': '突破'}).status_code == 409
        assert client.get(other_root + '/review-tags').json()['data'] == []
        buy = write(root + '/orders', {'symbol': '600000', 'side': 'buy', 'quantity': 100,
            'limit_price': '10', 'signal_date': '2025-01-01',
            'submit_date': '2025-01-01'}).json()['data']
        write(root + f"/orders/{buy['id']}/fill", {
            'expected_revision': 1, 'fill_date': '2025-01-01', 'fill_price': '9'})
        buy_fill = client.get(root + '/fills').json()['data'][0]
        payload = {'expected_revision': 0, 'emotion_tag_id': emotion['id'],
                   'reason_tag_ids': [reason['id'], extra['id']]}
        assignment = write(root + f"/fills/{buy_fill['id']}/tags", payload, 'PUT')
        assert assignment.status_code == 200, assignment.text
        assert write(root + f"/fills/{buy_fill['id']}/tags", payload, 'PUT').status_code == 409
        assert write(other_root + f"/fills/{buy_fill['id']}/tags", payload, 'PUT').status_code == 404
        assert write(root + f"/fills/{buy_fill['id']}/tags", {
            'expected_revision': 1, 'emotion_tag_id': None,
            'reason_tag_ids': [reason['id'], reason['id']]}, 'PUT').status_code == 400
        buy_stats = client.get(root + '/tag-stats').json()['data']
        by_id = {item['id']: item for item in buy_stats['data']}
        assert by_id[reason['id']]['fill_count'] == 1
        assert by_id[reason['id']]['sell_count'] == 0
        assert by_id[reason['id']]['win_rate_pct'] is None
        assert by_id[reason['id']]['realized_pnl'] == '0.00'
        write(root + '/settle', {'to_date': '2025-01-02'})
        sell = write(root + '/orders', {'symbol': '600000', 'side': 'sell', 'quantity': 100,
            'limit_price': '11', 'signal_date': '2025-01-02',
            'submit_date': '2025-01-02'}).json()['data']
        write(root + f"/orders/{sell['id']}/fill", {
            'expected_revision': 1, 'fill_date': '2025-01-02', 'fill_price': '12'})
        sell_fill = client.get(root + '/fills').json()['data'][-1]
        write(root + f"/fills/{sell_fill['id']}/tags", {
            'expected_revision': 0, 'emotion_tag_id': emotion['id'],
            'reason_tag_ids': [reason['id']]}, 'PUT')
        stats = client.get(root + '/tag-stats').json()['data']['data']
        reason_stats = next(item for item in stats if item['id'] == reason['id'])
        assert reason_stats['fill_count'] == 2
        assert reason_stats['sell_count'] == 1
        assert reason_stats['win_rate_pct'] == '100.00'
        assert Decimal(reason_stats['realized_pnl']) > 0
        assert Decimal(reason_stats['realized_return_pct']) > 0
        assert client.get(root + '/tag-stats?date_from=2025-01-02').json()['data']['data'][1]['fill_count'] == 1
        deleted = write(root + f"/review-tags/{reason['id']}?expected_revision=1", {}, 'DELETE')
        assert deleted.status_code == 200 and not deleted.json()['data']['active']
        assert len(client.get(root + '/fills').json()['data']) == 2
        assert next(item for item in client.get(root + '/tag-stats').json()['data']['data']
                    if item['id'] == reason['id'])['fill_count'] == 2
        assert write(root + f"/fills/{buy_fill['id']}/tags", {
            'expected_revision': 1, 'emotion_tag_id': emotion['id'],
            'reason_tag_ids': [reason['id'], extra['id']]}, 'PUT').status_code == 200
        assert write(root + '/reset', {'expected_wallet_revision': client.get(root + '/portfolio').json()['data']['wallet_revision'],
            'start_date': '2025-02-01', 'reset_config': False}).status_code == 200
        assert write(root + '/review-tags', {'type': 'reason', 'name': '新标签'}).status_code == 409
        assert client.get(root + '/fill-tags').json()['data'][0]['reason_tag_ids'] == [reason['id'], extra['id']]


def test_frozen_opening_match_respects_limit_date_and_source(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(path: str, body: dict):
            nonlocal serial
            serial += 1
            return client.post(path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': f'open-{serial}'})

        account = write('/api/v1/sim-accounts', {'name': '开盘撮合账户',
                         'initial_capital': '10000', 'start_date': '2025-01-01'}).json()['data']
        root = f"/api/v1/sim-accounts/{account['id']}"
        dataset = write('/api/v1/market/datasets', {'symbol': '600000', 'bars': [
            {'event_date': '2025-01-01', 'open': '10', 'high': '11', 'low': '9',
             'close': '10.5', 'volume': 1000},
            {'event_date': '2025-01-02', 'open': '9.5', 'high': '10', 'low': '9',
             'close': '9.8', 'volume': 1200},
        ]}).json()['data']
        order = write(root + '/orders', {'symbol': '600000', 'side': 'buy',
                      'quantity': 100, 'limit_price': '9', 'signal_date': '2025-01-01',
                      'submit_date': '2025-01-01'}).json()['data']
        endpoint = root + f"/orders/{order['id']}/match-open"
        body = {'expected_revision': 1, 'dataset_id': dataset['id']}
        assert write(endpoint, body).status_code == 409  # Same-day open was already known.
        assert write(root + '/settle', {'to_date': '2025-01-02'}).status_code == 200
        missed = write(endpoint, body)
        assert missed.status_code == 200 and missed.json()['data']['status'] == 'limit_not_met'
        assert client.get(root + '/orders').json()['data'][0]['status'] == 'pending'
        # A second frozen sample can represent a corrected vendor file, retaining its own hash.
        revised = write('/api/v1/market/datasets', {'symbol': '600000', 'bars': [
            {'event_date': '2025-01-01', 'open': '10', 'high': '11', 'low': '9',
             'close': '10.5', 'volume': 1000},
            {'event_date': '2025-01-02', 'open': '8.5', 'high': '10', 'low': '8',
             'close': '9.8', 'volume': 1200},
        ]}).json()['data']
        filled = write(endpoint, {'expected_revision': 1, 'dataset_id': revised['id']})
        assert filled.status_code == 200, filled.text
        assert filled.json()['data']['status'] == 'filled'
        assert filled.json()['data']['fill']['fill_price'] == '8.5000'
        assert filled.json()['data']['fill']['price_source'] == f"market_open:{revised['id']}"
        assert filled.json()['data']['portfolio']['positions'][0]['sellable_quantity'] == 0
        assert write(endpoint, {'expected_revision': 1, 'dataset_id': revised['id']}).status_code == 409


def test_market_day_advance_matches_pending_orders_atomically(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        serial = 0

        def write(path: str, body: dict):
            nonlocal serial
            serial += 1
            return client.post(path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': f'batch-{serial}'})

        account = write('/api/v1/sim-accounts', {'name': '批量撮合',
                         'initial_capital': '10000', 'start_date': '2025-01-01'}).json()['data']
        root = f"/api/v1/sim-accounts/{account['id']}"
        dataset = write('/api/v1/market/datasets', {'symbol': '600000', 'bars': [
            {'event_date': '2025-01-01', 'open': '10', 'high': '11', 'low': '9',
             'close': '10.5', 'volume': 1000},
            {'event_date': '2025-01-02', 'open': '9', 'high': '10', 'low': '8.8',
             'close': '9.5', 'volume': 1100},
        ]}).json()['data']
        for limit in ('10', '8'):
            assert write(root + '/orders', {'symbol': '600000', 'side': 'buy',
                         'quantity': 100, 'limit_price': limit, 'signal_date': '2025-01-01',
                         'submit_date': '2025-01-01'}).status_code == 200
        revision = client.get(root + '/portfolio').json()['data']['wallet_revision']
        endpoint = root + '/advance-market-day'
        missing = write(endpoint, {'expected_wallet_revision': revision,
                       'to_date': '2025-01-02', 'datasets': {}})
        assert missing.status_code == 400
        assert client.get(root + '/portfolio').json()['data']['as_of_date'] == '2025-01-01'
        result = write(endpoint, {'expected_wallet_revision': revision,
                       'to_date': '2025-01-02', 'datasets': {'600000': dataset['id']}})
        assert result.status_code == 200, result.text
        statuses = [item['status'] for item in result.json()['data']['outcomes']]
        assert sorted(statuses) == ['filled', 'limit_not_met']
        assert result.json()['data']['portfolio']['positions'][0]['sellable_quantity'] == 0
        csv_rows = list(csv.DictReader(io.StringIO(client.get(
            f"/api/v1/accounts/{account['id']}/exports/sim-fills.csv").content.decode('utf-8-sig'))))
        assert len(csv_rows) == 1 and csv_rows[0]['symbol'] == '600000'
        assert csv_rows[0]['side'] == 'buy' and csv_rows[0]['quantity'] == '100'
        valuation_url = root + '/valuation'
        valuation_params = {'decision_at': '2025-01-02T23:59:59+00:00',
                            'dataset_id': dataset['id']}
        strict = client.get(valuation_url, params=valuation_params)
        assert strict.status_code == 200, strict.text
        assert strict.json()['data']['total_assets'] is None
        assert 'historical_availability_unknown' in strict.json()['data']['quality_flags']
        assumed = client.get(valuation_url, params={**valuation_params, 'strict': 'false'})
        assert assumed.status_code == 200, assumed.text
        assert assumed.json()['data']['positions'][0]['market_value'] == '950.00'
        assert assumed.json()['data']['total_assets'] is not None
        assert assumed.json()['data']['valuation_quality'] == 'qualified'
        position = assumed.json()['data']['positions'][0]
        assert position['cost_basis'] == client.get(root + '/portfolio').json()['data']['positions'][0]['cost_basis']
        assert Decimal(position['unrealized_pnl']) == Decimal(position['market_value']) - Decimal(position['cost_basis'])
        assert assumed.json()['data']['known_unrealized_pnl'] == position['unrealized_pnl']
        assert strict.json()['data']['positions'][0]['unrealized_pnl'] is None
        assert write(endpoint, {'expected_wallet_revision': revision,
                     'to_date': '2025-01-03', 'datasets': {'600000': dataset['id']}}).status_code == 409


def test_immutable_market_snapshot_point_in_time_and_backup(tmp_path: Path) -> None:
    source = tmp_path / 'market-source'
    with started_client(create_app(source, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        bars = [
            {'event_date': '2025-01-01', 'open': '10', 'high': '11', 'low': '9',
             'close': '10.5', 'volume': 1000, 'available_at': '2025-01-02T00:00:00+00:00'},
            {'event_date': '2025-01-02', 'open': '10.5', 'high': '12', 'low': '10',
             'close': '11.5', 'volume': 1200, 'available_at': '2025-01-03T00:00:00+00:00'},
        ]
        response = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'adjustment': 'none', 'bars': bars},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'bars'})
        assert response.status_code == 200, response.text
        dataset = response.json()['data']
        repeated = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'adjustment': 'none', 'bars': bars},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'bars-again'})
        assert repeated.json()['data']['id'] == dataset['id']
        frozen = client.get('/api/v1/market/datasets/' + dataset['id']).json()['data']
        assert frozen['bar_count'] == 2
        strict, quality = eligible_bars(frozen['bars'], '2025-01-02T12:00:00+00:00', True)
        assert len(strict) == 1 and quality == []
        unknown = [{**bar, 'available_at': None} for bar in frozen['bars']]
        assert eligible_bars(unknown, '2025-01-03T12:00:00+00:00', True)[0] == []
        assert eligible_bars(unknown, '2025-01-03T12:00:00+00:00', False)[1] == ['historical_availability_unknown']
        archive = client.get('/api/v1/backups/export').content
        assert f"market/{dataset['id']}.json" in verify_backup(archive)['files']
    destination = tmp_path / 'market-restored'
    restore_to_new_directory(archive, destination)
    with started_client(create_app(destination, auto_rebuild=False)) as client:
        client.get('/api/v1/session')
        assert client.get('/api/v1/market/datasets/' + dataset['id']).json()['data']['bars'] == frozen['bars']
        health = client.get('/api/v1/market/diagnostics?verify=true').json()['data']
        assert health['dataset_count'] == 1 and health['verified_contents'] == 1
        assert health['missing_or_bad_count'] == 0 and health['provider_counts'] == {'manual_import': 1}
        (destination / 'market' / f"{dataset['id']}.json").write_text('corrupted', encoding='utf-8')
        assert client.get('/api/v1/market/datasets/' + dataset['id']).status_code == 409
        damaged = client.get('/api/v1/market/diagnostics?verify=true').json()['data']
        assert damaged['missing_or_bad_count'] == 1
        assert damaged['items'][0]['content_status'] == 'hash_mismatch'
        assert client.get('/api/v1/backups/export').status_code == 409


def test_relative_strength_signal_matches_legacy_plugin_on_frozen_bars(tmp_path: Path) -> None:
    with started_client(create_app(tmp_path, auto_rebuild=False)) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        catalog_response = client.get('/api/v1/research/strategies')
        assert catalog_response.status_code == 200, catalog_response.text
        catalog = catalog_response.json()['data']
        assert len(catalog) == 19
        assert len({row['id'] for row in catalog}) == 19
        assert sum(row['origin'] == 'final_trade' for row in catalog) == 16
        assert sum(row['origin'] == 'classic_reference' for row in catalog) == 3
        from app.core.strategy_registry import StrategyRegistry
        from trade_app.research.classic_domain import CLASSIC_STRATEGY_IDS
        assert {row['id'] for row in catalog if row['origin'] == 'final_trade'} == {row.strategy_id for row in StrategyRegistry().list()}
        assert {row['id'] for row in catalog if row['origin'] == 'classic_reference'} == set(CLASSIC_STRATEGY_IDS)
        active = next(row for row in catalog if row['id'] == 'relative_strength_breakout_v1')
        assert active['status'] == 'partial_signal' and active['signal_params']['min_ret40'] == '0.12'
        assert sum(row['status'] == 'not_migrated' for row in catalog) == 0
        assert next(row for row in catalog if row['id'] == 'b1_mtf_v1')['status'] == 'b1_screener'
        start = date(2025, 1, 1)
        bars = []
        for index in range(45):
            day = start + timedelta(days=index)
            close = 10 + index * 0.2
            bars.append({'event_date': day.isoformat(), 'open': f'{close - 0.1:.4f}',
                         'high': f'{close + 0.5:.4f}', 'low': f'{close - 0.5:.4f}',
                         'close': f'{close:.4f}', 'volume': 1000 + index * 30,
                         'available_at': datetime.combine(day + timedelta(days=1),
                                                           datetime.min.time(), timezone.utc).isoformat()})
        dataset = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'bars': bars, 'adjustment': 'none'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'market'}).json()['data']
        body = {'dataset_id': dataset['id'], 'strategy_id': 'relative_strength_breakout_v1',
                'decision_at': '2025-02-16T12:00:00+00:00', 'strict': True,
                'params': {'min_vol_slope20': '0'}}
        created = client.post('/api/v1/research/runs', json=body,
                              headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'research'})
        assert created.status_code == 200, created.text
        run = created.json()['data']
        candidate = run['result']['candidate']
        assert run['result']['status'] == 'computed'
        assert run['result']['signal'] is True
        sim = client.post('/api/v1/sim-accounts', json={
            'name': '研究草稿账户', 'initial_capital': '10000', 'start_date': '2025-02-16'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'research-sim'})
        assert sim.status_code == 200, sim.text
        sim_root = f"/api/v1/sim-accounts/{sim.json()['data']['id']}"
        sizing = sim_root + '/drafts/size'
        sizing_headers = {'X-CSRF-Token': csrf}
        by_lots = client.post(sizing, json={'mode': 'lots', 'value': '5', 'limit_price': '20'}, headers=sizing_headers).json()['data']
        assert by_lots['quantity'] == 500 and by_lots['cash_gap'] != '0.00'
        assert by_lots['max_affordable_quantity'] == 400
        by_amount = client.post(sizing, json={'mode': 'amount', 'value': '2500', 'limit_price': '20'}, headers=sizing_headers).json()['data']
        assert by_amount['quantity'] == 100 and by_amount['can_create']
        too_small = client.post(sizing, json={'mode': 'amount', 'value': '1', 'limit_price': '20'}, headers=sizing_headers).json()['data']
        assert too_small['quantity'] == 0 and too_small['required_cash'] == '0.00'
        assert not too_small['can_create']
        by_percent = client.post(sizing, json={'mode': 'cash_percent', 'value': '25', 'limit_price': '20'}, headers=sizing_headers).json()['data']
        assert by_percent['quantity'] == 100
        assert client.post(sizing, json={'mode': 'cash_percent', 'value': '101', 'limit_price': '20'}, headers=sizing_headers).status_code == 400
        assert client.post(sizing, json={'mode': 'lots', 'value': '1.5', 'limit_price': '20'}, headers=sizing_headers).status_code == 400
        draft = client.post(sim_root + '/drafts', json={
            'source_run_id': run['id'], 'quantity': 100, 'limit_price': '20'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'research-draft'})
        assert draft.status_code == 200, draft.text
        draft_id = draft.json()['data']['id']
        preview = client.get(sim_root + f'/drafts/{draft_id}/preview').json()['data']
        assert preview['can_submit'] and preview['cash_gap'] == '0.00'
        assert preview['note'].startswith('观察信号')
        foreign = client.post('/api/v1/accounts', json={'name': '隔离实盘账户'},
                              headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'research-real'})
        assert foreign.status_code == 200
        assert client.post(f"/api/v1/sim-accounts/{foreign.json()['data']['id']}/drafts", json={
            'source_run_id': run['id'], 'quantity': 100, 'limit_price': '20'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'research-real-draft'}).status_code == 404
        submit = client.post(sim_root + f"/drafts/{draft_id}/submit?expected_revision=1&expected_wallet_revision={preview['wallet_revision']}&expected_config_version={preview['config_version']}",
                             json={}, headers={'X-CSRF-Token': csrf,
                                               'Idempotency-Key': 'research-draft-submit'})
        assert submit.status_code == 200, submit.text
        assert submit.json()['data']['draft']['status'] == 'submitted'
        assert client.get(sim_root + '/orders').json()['data'][0]['symbol'] == '600000'
        assert client.post(sim_root + f"/drafts/{draft_id}/submit?expected_revision=2&expected_wallet_revision={preview['wallet_revision']}&expected_config_version={preview['config_version']}",
                           json={}, headers={'X-CSRF-Token': csrf,
                                             'Idempotency-Key': 'research-draft-resubmit'}).status_code == 409
        second_run = client.post('/api/v1/research/runs', json={
            **body, 'decision_at': '2025-02-15T12:00:00+00:00'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'research-second'})
        assert second_run.status_code == 200 and second_run.json()['data']['result']['signal']
        batch_sim = client.post('/api/v1/sim-accounts', json={
            'name': '批量草稿账户', 'initial_capital': '10000', 'start_date': '2025-02-16'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'research-batch-sim'})
        batch_root = f"/api/v1/sim-accounts/{batch_sim.json()['data']['id']}"
        assert client.get(batch_root + f'/drafts/{draft_id}/preview').status_code == 404
        draft_ids = []
        for index, source_id in enumerate((run['id'], second_run.json()['data']['id'])):
            response = client.post(batch_root + '/drafts', json={
                'source_run_id': source_id, 'quantity': 300, 'limit_price': '20'},
                headers={'X-CSRF-Token': csrf, 'Idempotency-Key': f'batch-draft-{index}'})
            assert response.status_code == 200, response.text
            draft_ids.append(response.json()['data']['id'])
        params = [('draft_id', item) for item in draft_ids]
        batch_preview = client.get(batch_root + '/drafts/preview-batch', params=params)
        assert batch_preview.status_code == 200, batch_preview.text
        assert not batch_preview.json()['data']['can_submit']
        overfunded = {'drafts': [{'id': item, 'expected_revision': 1} for item in draft_ids],
                      'expected_wallet_revision': batch_preview.json()['data']['wallet_revision'],
                      'expected_config_version': batch_preview.json()['data']['config_version']}
        rejected = client.post(batch_root + '/drafts/submit-batch', json=overfunded,
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'batch-overfunded'})
        assert rejected.status_code == 409
        assert client.get(batch_root + '/orders').json()['data'] == []
        resized = client.put(batch_root + '/drafts/' + draft_ids[1], json={
            'expected_revision': 1, 'quantity': 100, 'limit_price': '20'},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'batch-resize'})
        assert resized.status_code == 200, resized.text
        ready_preview = client.get(batch_root + '/drafts/preview-batch', params=params).json()['data']
        assert ready_preview['can_submit']
        stale_second = client.post(batch_root + '/drafts/submit-batch', json={
            'drafts': [{'id': draft_ids[0], 'expected_revision': 1},
                       {'id': draft_ids[1], 'expected_revision': 1}],
            'expected_wallet_revision': ready_preview['wallet_revision'],
            'expected_config_version': ready_preview['config_version']},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'batch-stale-second'})
        assert stale_second.status_code == 409
        assert client.get(batch_root + '/orders').json()['data'] == []
        assert all(row['status'] == 'draft' for row in client.get(batch_root + '/drafts').json()['data'])
        accepted = client.post(batch_root + '/drafts/submit-batch', json={
            'drafts': [{'id': draft_ids[0], 'expected_revision': 1},
                       {'id': draft_ids[1], 'expected_revision': 2}],
            'expected_wallet_revision': ready_preview['wallet_revision'],
            'expected_config_version': ready_preview['config_version']},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'batch-submit'})
        assert accepted.status_code == 200, accepted.text
        assert len(accepted.json()['data']['orders']) == 2
        assert len(client.get(batch_root + '/orders').json()['data']) == 2
        plugin = RelativeStrengthBreakoutPlugin()
        legacy_row = SimpleNamespace(**candidate)
        assert bool(plugin.build_universe(candidates=[legacy_row], params={
            'min_ret40': 0.12, 'max_retrace20': 0.22,
            'min_up_down_volume_ratio': 1.15, 'min_vol_slope20': 0,
            'min_ai_confidence': 0}, mode='signals')) == run['result']['signal']
        assert plugin.generate_signals(row=legacy_row, snapshot={}, params={
            'min_ret40': 0.12, 'max_retrace20': 0.22,
            'min_up_down_volume_ratio': 1.15}) == run['result']['signal']
        repeated = client.post('/api/v1/research/runs', json=body,
                               headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'research-again'})
        assert repeated.json()['data']['id'] == run['id']
        earlier = client.post('/api/v1/research/runs', json={**body, 'decision_at': '2025-01-20T12:00:00+00:00'},
                              headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'research-earlier'})
        assert earlier.json()['data']['result']['status'] == 'insufficient_data'


def test_durable_single_symbol_backtest_uses_frozen_bars_and_shared_fees(tmp_path: Path) -> None:
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        headers = {'X-CSRF-Token': csrf, 'Idempotency-Key': 'backtest-market'}
        start = date(2025, 1, 1)
        bars = []
        for index in range(60):
            day = start + timedelta(days=index)
            close = 10 + index * 0.15
            bars.append({'event_date': day.isoformat(), 'open': f'{close - 0.1:.4f}',
                         'high': f'{close + 0.5:.4f}', 'low': f'{close - 0.5:.4f}',
                         'close': f'{close:.4f}', 'volume': 1000 + index * 30,
                         'available_at': datetime.combine(day + timedelta(days=1),
                                                           datetime.min.time(), timezone.utc).isoformat()})
        dataset_response = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'bars': bars}, headers=headers)
        assert dataset_response.status_code == 200, dataset_response.text
        dataset_id = dataset_response.json()['data']['id']
        body = {'dataset_id': dataset_id, 'strategy_id': 'relative_strength_breakout_v1',
                'initial_capital': '10000', 'holding_bars': 3,
                'max_position_pct': '0.95', 'strict': True,
                'params': {'min_vol_slope20': '0'}}
        queued = client.post('/api/v1/backtests', json=body, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'backtest-create'})
        assert queued.status_code == 200, queued.text
        run_id = queued.json()['data']['id']
        assert queued.json()['data']['state'] == 'queued'
        assert client.get('/api/v1/backtests/' + run_id + '/export.xlsx').status_code == 409
        assert process_one_backtest(app.state.db_factory, tmp_path)
        completed = client.get('/api/v1/backtests/' + run_id)
        assert completed.status_code == 200, completed.text
        run = completed.json()['data']
        assert run['state'] == 'succeeded'
        assert run['result']['trade_count'] > 0
        workbook_response = client.get('/api/v1/backtests/' + run_id + '/export.xlsx')
        assert workbook_response.status_code == 200
        workbook = load_workbook(io.BytesIO(workbook_response.content), read_only=True)
        assert workbook.sheetnames == ['Summary', 'Params', 'Config', 'Trades', 'Equity', 'Decisions', 'Analysis']
        assert dict(workbook['Analysis'].values).get('status') == 'not_generated'
        assert list(workbook['Trades'].values)[1][0] == run['result']['trades'][0]['date']
        assert list(workbook['Equity'].values)[1][0] == run['result']['equity'][0]['date']
        assert len(run['result']['equity']) == len(bars)
        assert run['result']['trades'][0]['signal_date'] < run['result']['trades'][0]['date']
        first_buy = run['result']['trades'][0]
        shared_fees = calculate_fees(Decimal(first_buy['price']), first_buy['quantity'],
                                     'buy', fee_rule(DEFAULT_CONFIG))
        assert Decimal(first_buy['fees']) == shared_fees.total
        assert 'single_symbol_selection_unverified' in run['result']['quality_flags']
        assert client.post('/api/v1/backtests', json=body, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'backtest-repeat'}).json()['data']['id'] == run_id
        assert client.get('/api/v1/backtests').json()['data'][0]['summary']['trade_count'] > 0
        assert client.post('/api/v1/backtests/' + run_id + '/cancel', json={}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'backtest-cancel'}).status_code == 409
        unknown_bars = [{**bar, 'available_at': None} for bar in bars]
        unknown = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'bars': unknown_bars}, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'backtest-unknown-market'})
        assert unknown.status_code == 200, unknown.text
        unknown_id = unknown.json()['data']['id']
        unknown_run = client.post('/api/v1/backtests', json={**body, 'dataset_id': unknown_id},
                                  headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'backtest-unknown'})
        assert unknown_run.status_code == 200, unknown_run.text
        assert process_one_backtest(app.state.db_factory, tmp_path)
        unknown_result = client.get('/api/v1/backtests/' + unknown_run.json()['data']['id']).json()['data']['result']
        assert unknown_result['trade_count'] == 0
        assert 'historical_availability_unknown' in unknown_result['quality_flags']


def test_backtest_queue_cancel_retry_and_restart_recovery(tmp_path: Path) -> None:
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        def post(path: str, body: dict, key: str):
            return client.post(path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': key})
        bars = [{'event_date': (date(2025, 1, 1) + timedelta(days=index)).isoformat(),
                 'open': '10', 'high': '11', 'low': '9', 'close': '10', 'volume': 1000}
                for index in range(32)]
        dataset = post('/api/v1/market/datasets', {'symbol': '600000', 'bars': bars}, 'recovery-market')
        assert dataset.status_code == 200, dataset.text
        queued = post('/api/v1/backtests', {
            'dataset_id': dataset.json()['data']['id'],
            'strategy_id': 'relative_strength_breakout_v1'}, 'recovery-run')
        assert queued.status_code == 200, queued.text
        run_id = queued.json()['data']['id']
        cancelled = post('/api/v1/backtests/' + run_id + '/cancel', {}, 'recovery-cancel')
        assert cancelled.status_code == 200 and cancelled.json()['data']['state'] == 'cancelled'
        assert not process_one_backtest(app.state.db_factory, tmp_path)
        retried = post('/api/v1/backtests/' + run_id + '/retry', {}, 'recovery-retry')
        assert retried.status_code == 200 and retried.json()['data']['state'] == 'queued'
        with app.state.db_factory.begin() as session:
            session.get(BacktestRun, run_id).state = 'running'
    restarted = create_app(tmp_path, auto_rebuild=False)
    with started_client(restarted) as client:
        client.get('/api/v1/session')
        recovered = client.get('/api/v1/backtests/' + run_id).json()['data']
        assert recovered['state'] == 'queued'
        assert recovered['error'] == 'RECOVERED_AFTER_RESTART'
        assert process_one_backtest(restarted.state.db_factory, tmp_path)
        completed = client.get('/api/v1/backtests/' + run_id).json()['data']
        assert completed['state'] == 'succeeded' and completed['error'] is None


def test_b1_frozen_run_parity_history_and_as_of(tmp_path: Path) -> None:
    bars = []
    for index in range(940):
        day = (date(2020, 1, 1) + timedelta(days=index)).isoformat()
        close = round(10 + index * 0.01 + (index % 17) * 0.01, 2)
        bars.append({'event_date': day, 'open': str(close), 'high': str(round(close + 0.1, 2)),
                     'low': str(round(close - 0.1, 2)), 'close': str(close), 'volume': 1000})
    old_bars = [{'date': bar['event_date'],
                 **{key: float(bar[key]) for key in ('open', 'high', 'low', 'close')},
                 'volume': bar['volume']} for bar in bars]
    assert new_check_b1('600000', old_bars, NewB1Params()) == old_check_b1(
        '600000', old_bars, OldB1Params())

    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        headers = {'X-CSRF-Token': csrf, 'Idempotency-Key': 'b1-dataset'}
        imported = client.post('/api/v1/market/datasets', json={
            'symbol': 'SH600000', 'adjustment': 'none', 'bars': bars}, headers=headers)
        assert imported.status_code == 200, imported.text
        dataset_id = imported.json()['data']['id']
        payload = {'dataset_ids': [dataset_id], 'as_of_date': bars[-1]['event_date']}
        headers['Idempotency-Key'] = 'b1-run'
        response = client.post('/api/v1/research/b1-runs', json=payload, headers=headers)
        assert response.status_code == 200, response.text
        run = response.json()['data']
        assert run['result']['total_scanned'] == 1
        assert run['result']['excluded'] == []
        assert run['result']['quality_flags'] == ['HISTORICAL_AVAILABLE_AT_UNKNOWN']
        repeat = client.post('/api/v1/research/b1-runs', json=payload, headers=headers)
        assert repeat.json()['data']['id'] == run['id']
        assert client.get('/api/v1/research/b1-runs').json()['data'][0]['id'] == run['id']
        assert client.get('/api/v1/research/b1-runs/' + run['id']).json()['data']['result'] == run['result']
        headers['Idempotency-Key'] = 'b1-early'
        early = client.post('/api/v1/research/b1-runs', json={
            **payload, 'as_of_date': bars[800]['event_date']}, headers=headers)
        assert early.status_code == 200, early.text
        assert early.json()['data']['result']['excluded'][0]['reason'] == 'INSUFFICIENT_BARS_AS_OF_DATE'
        future_bars = [{**bar, **({'available_at': '2025-01-01T00:00:00+00:00'}
                                  if index >= 890 else {})} for index, bar in enumerate(bars)]
        headers['Idempotency-Key'] = 'b1-future-dataset'
        future_dataset = client.post('/api/v1/market/datasets', json={
            'symbol': 'SH600001', 'adjustment': 'none', 'bars': future_bars}, headers=headers)
        assert future_dataset.status_code == 200, future_dataset.text
        headers['Idempotency-Key'] = 'b1-future-run'
        future = client.post('/api/v1/research/b1-runs', json={
            'dataset_ids': [future_dataset.json()['data']['id']],
            'as_of_date': bars[-1]['event_date']}, headers=headers)
        assert future.status_code == 200, future.text
        assert future.json()['data']['result']['excluded'][0]['reason'] == 'INSUFFICIENT_BARS_AS_OF_DATE'


def test_b1_tdx_universe_job_publishes_frozen_scan(tmp_path: Path, monkeypatch) -> None:
    tdx = tmp_path / 'tdx'
    folder = tdx / 'vipdoc' / 'sh' / 'lday'
    folder.mkdir(parents=True)
    path = folder / 'sh600000.day'
    records = []
    for index in range(920):
        day = date(2020, 1, 1) + timedelta(days=index)
        close = 1000 + index
        records.append(struct.pack('<IIIIIfII', int(day.strftime('%Y%m%d')),
                                   close, close + 10, close - 10, close,
                                   float(close * 1000), 100_000, 0))
    source = b''.join(records)
    path.write_bytes(source)
    monkeypatch.setenv('TRADE_TDX_ROOT', str(tdx))
    data_dir = tmp_path / 'data'
    app = create_app(data_dir, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        response = client.post('/api/v1/research/tdx-b1-jobs', json={
            'markets': ['sh'], 'as_of_date': '2022-07-08', 'max_bars': 1000},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'b1-market'})
        assert response.status_code == 200, response.text
        job = response.json()['data']
        assert job['kind'] == 'b1' and job['total'] == 1
        for _ in range(2):
            assert process_one_universe_symbol(app.state.db_factory, data_dir, tdx)
        complete = client.get('/api/v1/research/tdx-universe-jobs/' + job['id']).json()['data']
        assert complete['state'] == 'succeeded' and complete['processed'] == 1
        run = client.get('/api/v1/research/b1-runs/' + complete['run_id']).json()['data']
        assert run['result']['total_scanned'] == 1
        assert run['request']['source'] == 'tdx_universe'
        assert len(run['request']['datasets']) == 1
        assert run['result']['quality_flags'] == [
            'HISTORICAL_AVAILABLE_AT_UNKNOWN', 'TDX_INPUTS_FROZEN_SEQUENTIALLY']
    assert path.read_bytes() == source


def test_b1_positive_hit_promotes_to_sim_draft(tmp_path: Path) -> None:
    days = []
    day = date(2020, 1, 1)
    while len(days) < 940:
        if day.weekday() < 5:
            days.append(day.isoformat())
        day += timedelta(days=1)
    tail = [16.17, 15.51, 15.94, 15.64, 16.28, 15.74, 15.53, 15.52, 15.91]
    bars = []
    for index, day in enumerate(days):
        close = round(10 + index * 0.005 + max(0, index - 700) ** 2 * 0.00002, 2)
        if index >= 931:
            close = tail[index - 931]
        bars.append({'event_date': day, 'open': str(close),
                     'high': str(round(close + 0.12, 2)),
                     'low': str(round(close - 0.12, 2)),
                     'close': str(close), 'volume': 300 if index == 939 else 1000})
    legacy_bars = [{'date': bar['event_date'],
                    **{key: float(bar[key]) for key in ('open', 'high', 'low', 'close')},
                    'volume': bar['volume']} for bar in bars]
    expected = old_check_b1('600000', legacy_bars, OldB1Params())
    assert expected is not None
    assert new_check_b1('600000', legacy_bars, NewB1Params()) == expected
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        def write(path, body, key):
            response = client.post(path, json=body, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': key})
            assert response.status_code == 200, response.text
            return response.json()['data']
        dataset = write('/api/v1/market/datasets', {
            'symbol': '600000', 'adjustment': 'none', 'bars': bars}, 'b1-positive-dataset')
        run = write('/api/v1/research/b1-runs', {
            'dataset_ids': [dataset['id']], 'as_of_date': days[-1]}, 'b1-positive-run')
        assert run['result']['hit_count'] == 1
        assert run['result']['hits'][0]['kdj_j'] == expected['kdj_j']
        bad = client.post(f"/api/v1/research/b1-runs/{run['id']}/signals/600001", json={},
                          headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'b1-bad-hit'})
        assert bad.status_code == 404
        signal = write(f"/api/v1/research/b1-runs/{run['id']}/signals/600000", {}, 'b1-promote')
        assert signal['strategy_id'] == 'b1_mtf_v1' and signal['result']['signal'] is True
        assert signal['result']['executable_date'] is None
        account = write('/api/v1/sim-accounts', {
            'name': 'B1 草稿账户', 'initial_capital': '10000', 'start_date': days[-1]}, 'b1-sim')
        draft = write(f"/api/v1/sim-accounts/{account['id']}/drafts", {
            'source_run_id': signal['id'], 'quantity': 100, 'limit_price': '15.91'}, 'b1-draft')
        assert draft['symbol'] == '600000' and draft['signal_date'] == days[-1]


def test_trend_leaders_frozen_timeline_matches_daily_legacy_ranking(tmp_path: Path) -> None:
    bars_by_symbol = {}
    for symbol, slope in (('600000', 0.03), ('000001', 0.06)):
        bars = []
        for index in range(40):
            day = (date(2025, 1, 1) + timedelta(days=index)).isoformat()
            close = round(10 + index * slope, 2)
            bars.append({'event_date': day, 'open': str(close),
                         'high': str(round(close + 0.1, 2)),
                         'low': str(round(close - 0.1, 2)),
                         'close': str(close), 'volume': 100000,
                         'amount': '100000000'})
        bars_by_symbol[symbol] = bars
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        ids = []
        for symbol, bars in bars_by_symbol.items():
            response = client.post('/api/v1/market/datasets', json={
                'symbol': symbol, 'adjustment': 'none', 'bars': bars}, headers={
                'X-CSRF-Token': csrf, 'Idempotency-Key': 'trend-' + symbol})
            assert response.status_code == 200, response.text
            ids.append(response.json()['data']['id'])
        body = {'dataset_ids': ids, 'date_from': '2025-01-10',
                'date_to': '2025-02-09', 'window_days': 5,
                'daily_top_n': 1, 'board_filters': ['main'], 'min_amount_avg': 5e7}
        response = client.post('/api/v1/research/trend-leader-runs', json=body, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'trend-run'})
        assert response.status_code == 200, response.text
        run = response.json()['data']
        assert run['total_scanned'] == 2 and run['leader_count'] == 1
        assert run['result']['leaders'][0]['symbol'] == '000001'
        assert run['result']['leaders'][0]['leader_days'] == len(run['result']['dates'])
        old_bars = {symbol: [{'date': bar['event_date'], 'close': float(bar['close']),
                              'amount': float(bar['amount'])} for bar in bars]
                    for symbol, bars in bars_by_symbol.items()}
        for day in run['result']['dates']:
            assert run['result']['daily_leaders'][day] == legacy_trend_top(
                old_bars, day, 5, 1, 5e7)
        assert client.get('/api/v1/research/trend-leader-runs/' + run['id']).json()['data']['result'] == run['result']
        assert client.get('/api/v1/research/trend-leader-runs').json()['data'][0]['id'] == run['id']
        repeat = client.post('/api/v1/research/trend-leader-runs', json=body, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'trend-run'})
        assert repeat.status_code == 200 and repeat.json()['data']['id'] == run['id']


def test_trend_tdx_universe_job_publishes_ranked_frozen_result(tmp_path: Path, monkeypatch) -> None:
    tdx = tmp_path / 'tdx'
    for market, code, slope in (('sh', '600000', 2), ('sz', '000001', 4)):
        folder = tdx / 'vipdoc' / market / 'lday'
        folder.mkdir(parents=True)
        records = []
        for index in range(280):
            day = date(2024, 1, 1) + timedelta(days=index)
            close = 1000 + index * slope
            records.append(struct.pack('<IIIIIfII', int(day.strftime('%Y%m%d')),
                                       close, close + 10, close - 10, close,
                                       float(close * 100000), 100000, 0))
        (folder / f'{market}{code}.day').write_bytes(b''.join(records))
    monkeypatch.setenv('TRADE_TDX_ROOT', str(tdx))
    data_dir = tmp_path / 'data'
    app = create_app(data_dir, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        response = client.post('/api/v1/research/tdx-trend-jobs', json={
            'markets': ['sh', 'sz'], 'date_from': '2024-09-01',
            'date_to': '2024-10-06', 'window_days': 5, 'daily_top_n': 1,
            'board_filters': ['main'], 'min_amount_avg': 0, 'max_bars': 280},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'trend-market'})
        assert response.status_code == 200, response.text
        job = response.json()['data']
        assert job['kind'] == 'trend' and job['total'] == 2
        for _ in range(3):
            assert process_one_universe_symbol(app.state.db_factory, data_dir, tdx)
        finished = client.get('/api/v1/research/tdx-universe-jobs/' + job['id']).json()['data']
        assert finished['state'] == 'succeeded' and finished['processed'] == 2
        run = client.get('/api/v1/research/trend-leader-runs/' + finished['run_id']).json()['data']
        assert run['result']['data_scope'] == 'tdx_full_market'
        assert run['result']['total_scanned'] == 2
        assert run['result']['leaders'][0]['symbol'] == '000001'
        assert run['result']['leaders'][0]['leader_days'] == len(run['result']['dates'])
        assert 'TDX_INPUTS_FROZEN_SEQUENTIALLY' in run['quality_flags']
        assert len(run['request']['datasets']) == 2
        assert all(row['dataset_id'] for row in run['request']['datasets'])


def test_limit_up_ladder_selected_dataset_matches_legacy_height(tmp_path: Path) -> None:
    closes = [10.0] * 12 + [11.0, 12.1]
    bars = []
    for index, close in enumerate(closes):
        bars.append({'event_date': (date(2025, 1, 1) + timedelta(days=index)).isoformat(),
                     'open': str(close), 'high': str(round(close + 0.1, 2)),
                     'low': str(round(close - 0.1, 2)), 'close': str(close),
                     'volume': 100000, 'amount': '100000000'})
    legacy_bars = [{'date': bar['event_date'], 'close': float(bar['close'])} for bar in bars]
    assert legacy_limit_height(legacy_bars, 13, 'sh600000') == 2
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        headers = {'X-CSRF-Token': csrf, 'Idempotency-Key': 'ladder-data'}
        response = client.post('/api/v1/market/datasets', json={
            'symbol': '600000', 'adjustment': 'none', 'bars': bars}, headers=headers)
        assert response.status_code == 200, response.text
        dataset_id = response.json()['data']['id']
        headers['Idempotency-Key'] = 'ladder-run'
        body = {'dataset_ids': [dataset_id], 'date_from': '2025-01-09',
                'date_to': '2025-01-14', 'recent_days': 2,
                'historical_min_boards': 2, 'board_filters': ['main']}
        response = client.post('/api/v1/research/limit-up-ladder-runs', json=body, headers=headers)
        assert response.status_code == 200, response.text
        run = response.json()['data']
        assert run['stock_count'] == 1
        assert [point['board_height'] for point in run['result']['timeline']] == [1, 2]
        assert run['result']['summaries'][0]['latest_board_height'] == 2
        assert client.get('/api/v1/research/limit-up-ladder-runs/' + run['id']).json()['data']['result'] == run['result']
        assert client.get('/api/v1/research/limit-up-ladder-runs').json()['data'][0]['id'] == run['id']


def test_limit_up_ladder_tdx_job_publishes_frozen_timeline(tmp_path: Path, monkeypatch) -> None:
    tdx = tmp_path / 'tdx'
    folder = tdx / 'vipdoc' / 'sh' / 'lday'
    folder.mkdir(parents=True)
    closes = [1000] * 278 + [1100, 1210]
    records = []
    for index, close in enumerate(closes):
        day = date(2024, 1, 1) + timedelta(days=index)
        records.append(struct.pack('<IIIIIfII', int(day.strftime('%Y%m%d')),
                                   close, close + 10, close - 10, close,
                                   float(close * 100000), 100000, 0))
    (folder / 'sh600000.day').write_bytes(b''.join(records))
    monkeypatch.setenv('TRADE_TDX_ROOT', str(tdx))
    data_dir = tmp_path / 'data'
    app = create_app(data_dir, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        response = client.post('/api/v1/research/tdx-ladder-jobs', json={
            'markets': ['sh'], 'date_from': '2024-09-01', 'date_to': '2024-10-06',
            'recent_days': 3, 'historical_min_boards': 2,
            'board_filters': ['main'], 'max_bars': 280},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'ladder-market'})
        assert response.status_code == 200, response.text
        job = response.json()['data']
        assert job['kind'] == 'ladder' and job['total'] == 1
        for _ in range(2):
            assert process_one_universe_symbol(app.state.db_factory, data_dir, tdx)
        finished = client.get('/api/v1/research/tdx-universe-jobs/' + job['id']).json()['data']
        assert finished['state'] == 'succeeded'
        run = client.get('/api/v1/research/limit-up-ladder-runs/' + finished['run_id']).json()['data']
        assert run['result']['data_scope'] == 'tdx_full_market'
        assert [point['board_height'] for point in run['result']['timeline']] == [1, 2]
        assert len(run['request']['datasets']) == 1


def test_sector_flow_proxy_matches_legacy_fixed_index_series(monkeypatch) -> None:
    assert new_sector_codes == legacy_sector_codes
    datasets = {}
    old_series = {}
    for code, sector, slope in (('881001', '煤炭', 0.02), ('881006', '石油石化', 0.04)):
        bars = []
        old_bars = []
        for index in range(15):
            day = (date(2025, 1, 1) + timedelta(days=index)).isoformat()
            close = round(10 + index * slope, 2)
            amount = 100000000 + index * (2000000 if code == '881001' else 3000000)
            bars.append({'event_date': day, 'close': str(close), 'amount': str(amount)})
            old_bars.append({'date': day, 'close': close, 'amount': float(amount)})
        datasets[code] = {'bars': bars}
        old_series[sector] = old_bars
    monkeypatch.setattr(legacy_sector_flow, 'load_sector_series_by_name', lambda _path: old_series)
    old_rows, _, old_leaders, old_dates, _ = legacy_sector_flow.scan_sector_capital_flow(
        tdx_data_path='unused', date_from='2025-01-05', date_to='2025-01-15',
        daily_top_n=1, flow_window=5)
    result = compute_sector_flow(datasets, {'date_from': '2025-01-05',
                                             'date_to': '2025-01-15',
                                             'daily_top_n': 1, 'flow_window': 5})
    assert result['dates'] == old_dates
    assert [(row['date'], row['sector'], row['flow_score'], row['rank_flow'])
            for row in result['flow_table']] == [
                (row.date, row.sector, row.flow_score, row.rank_flow) for row in old_rows]
    assert [(row['sector'], row['leader_days']) for row in result['leaders']] == [
        (row.sector, row.leader_days) for row in old_leaders]


def test_sector_flow_tdx_job_freezes_source_and_publishes_history(tmp_path: Path, monkeypatch) -> None:
    tdx = tmp_path / 'tdx'
    folder = tdx / 'vipdoc' / 'sh' / 'lday'
    folder.mkdir(parents=True)
    for code, slope in (('881001', 2), ('881006', 4)):
        records = []
        for index in range(260):
            day = date(2024, 1, 1) + timedelta(days=index)
            close = 1000 + index * slope
            records.append(struct.pack('<IIIIIfII', int(day.strftime('%Y%m%d')),
                                       close, close + 10, close - 10, close,
                                       float((1000 + index * slope) * 100000), 100000, 0))
        (folder / f'sh{code}.day').write_bytes(b''.join(records))
    monkeypatch.setenv('TRADE_TDX_ROOT', str(tdx))
    data_dir = tmp_path / 'data'
    app = create_app(data_dir, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        response = client.post('/api/v1/research/tdx-sector-flow-jobs', json={
            'date_from': '2024-08-01', 'date_to': '2024-09-16',
            'daily_top_n': 1, 'flow_window': 5, 'max_bars': 260},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'sector-market'})
        assert response.status_code == 200, response.text
        job = response.json()['data']
        assert job['kind'] == 'sector' and job['total'] == 2
        for _ in range(3):
            assert process_one_universe_symbol(app.state.db_factory, data_dir, tdx)
        finished = client.get('/api/v1/research/tdx-universe-jobs/' + job['id']).json()['data']
        assert finished['state'] == 'succeeded'
        run = client.get('/api/v1/research/sector-flow-runs/' + finished['run_id']).json()['data']
        assert run['result']['loaded_index_count'] == 2
        assert run['result']['loaded_sector_count'] == 2
        assert run['result']['metric_label'] == 'tdx_index_amount_relative_to_rolling_mean_not_net_capital_flow'
        assert 'SECTOR_INDEX_PARTIAL' in run['quality_flags']
        assert len(run['request']['datasets']) == 2
        assert client.get('/api/v1/research/sector-flow-runs').json()['data'][0]['id'] == run['id']


def test_abnormal_domain_preserves_legacy_episode_and_cooling(tmp_path: Path) -> None:
    days = [(date(2025, 1, 2) + timedelta(days=index)).isoformat() for index in range(80)]
    close = 10.0
    bars = []
    for index, day in enumerate(days):
        previous = close
        if index:
            close *= 1.10 if index < 16 else 0.96 if 16 <= index < 20 else 1.10 if 20 <= index < 32 else 1.0
        bars.append({'date': day, 'open': previous, 'high': max(previous, close) * 1.01,
                     'low': min(previous, close) * 0.99, 'close': close,
                     'volume': 100000, 'amount': close * 100000})
    index_close = {day: 1000.0 for day in days}
    for snapshot_only in (False, True):
        for cooling in (0, 3, 5):
            args = {'symbol': 'sz300001', 'name': '测试', 'bars': bars,
                    'index_close_by_date': index_close, 'date_from': days[0],
                    'date_to': days[-1], 'include_warnings': True,
                    'snapshot_only': snapshot_only, 'trading_dates': days,
                    'cooling_days': cooling}
            assert [asdict(row) for row in new_analyze_abnormal(**args)] == [
                asdict(row) for row in legacy_analyze_abnormal(**args)]


def test_abnormal_tdx_job_uses_frozen_benchmark_and_publishes_events(tmp_path: Path, monkeypatch) -> None:
    tdx = tmp_path / 'tdx'
    folder = tdx / 'vipdoc' / 'sh' / 'lday'
    folder.mkdir(parents=True)
    stock, benchmark = [], []
    close = 1000
    for index in range(120):
        day = date(2025, 1, 1) + timedelta(days=index)
        if 0 < index <= 15:
            close = round(close * 1.12)
        item_date = int(day.strftime('%Y%m%d'))
        stock.append(struct.pack('<IIIIIfII', item_date, close, close + 10,
                                 close - 10, close, float(close * 100000), 100000, 0))
        benchmark.append(struct.pack('<IIIIIfII', item_date, 100000, 100010,
                                     99990, 100000, 100000000.0, 100000, 0))
    (folder / 'sh600000.day').write_bytes(b''.join(stock))
    (folder / 'sh000002.day').write_bytes(b''.join(benchmark))
    monkeypatch.setenv('TRADE_TDX_ROOT', str(tdx))
    data_dir = tmp_path / 'data'
    app = create_app(data_dir, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        response = client.post('/api/v1/research/tdx-abnormal-jobs', json={
            'markets': ['sh'], 'date_from': '2025-01-01', 'date_to': '2025-04-30',
            'scan_mode': 'full', 'cooling_days': 3, 'max_bars': 251},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'abnormal-market'})
        assert response.status_code == 200, response.text
        job = response.json()['data']
        assert job['kind'] == 'abnormal' and job['total'] == 1
        for _ in range(2):
            assert process_one_universe_symbol(app.state.db_factory, data_dir, tdx)
        finished = client.get('/api/v1/research/tdx-universe-jobs/' + job['id']).json()['data']
        assert finished['state'] == 'succeeded', finished
        run = client.get('/api/v1/research/abnormal-runs/' + finished['run_id']).json()['data']
        assert run['result']['algorithm_version'] == 'v4.1-fast-scan'
        assert run['result']['events']
        assert all(event['benchmark_symbol'] == 'sh000002' for event in run['result']['events'])
        assert all(event['dataset_id'] for event in run['result']['events'])
        assert run['request']['benchmark_sources']['sh_main']['dataset_id']
        assert client.get('/api/v1/research/abnormal-runs').json()['data'][0]['id'] == run['id']


def test_abnormal_domain_suspension_and_benchmark_fallback(tmp_path: Path, monkeypatch) -> None:
    days = [(date(2025, 1, 1) + timedelta(days=index)).isoformat() for index in range(100)]
    close = 10.0
    bars = []
    for index, day in enumerate(days):
        if index in (40, 41, 42, 43):
            continue
        previous = close
        if index:
            close *= 1.11 if index < 18 or 44 <= index < 56 else 1.0
        bars.append({'date': day, 'open': previous, 'high': close * 1.01,
                     'low': previous * 0.99, 'close': close,
                     'volume': 100000, 'amount': close * 100000})
    index_close = {day: 1000.0 for day in days}
    args = {'symbol': 'sh600000', 'name': '停牌样本', 'bars': bars,
            'index_close_by_date': index_close, 'date_from': days[0],
            'date_to': days[-1], 'include_warnings': True,
            'snapshot_only': False, 'trading_dates': days, 'cooling_days': 3}
    assert [asdict(row) for row in new_analyze_abnormal(**args)] == [
        asdict(row) for row in legacy_analyze_abnormal(**args)]

    tdx = tmp_path / 'tdx'
    folder = tdx / 'vipdoc' / 'sh' / 'lday'
    folder.mkdir(parents=True)
    records = []
    benchmark = []
    price = 1000
    for index, day in enumerate(days):
        if 0 < index <= 15:
            price = round(price * 1.12)
        stamp = int(day.replace('-', ''))
        records.append(struct.pack('<IIIIIfII', stamp, price, price + 10,
                                   price - 10, price, float(price * 100000), 100000, 0))
        benchmark.append(struct.pack('<IIIIIfII', stamp, 100000, 100010,
                                     99990, 100000, 100000000.0, 100000, 0))
    (folder / 'sh600000.day').write_bytes(b''.join(records))
    (folder / 'sh000001.day').write_bytes(b''.join(benchmark))
    monkeypatch.setenv('TRADE_TDX_ROOT', str(tdx))
    data_dir = tmp_path / 'data'
    app = create_app(data_dir, auto_rebuild=False)
    with started_client(app) as client:
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        response = client.post('/api/v1/research/tdx-abnormal-jobs', json={
            'markets': ['sh'], 'date_from': days[0], 'date_to': days[-1],
            'scan_mode': 'full', 'max_bars': 251},
            headers={'X-CSRF-Token': csrf, 'Idempotency-Key': 'abnormal-fallback'})
        assert response.status_code == 200, response.text
        job = response.json()['data']
        for _ in range(2):
            assert process_one_universe_symbol(app.state.db_factory, data_dir, tdx)
        finished = client.get('/api/v1/research/tdx-universe-jobs/' + job['id']).json()['data']
        assert finished['state'] == 'succeeded'
        run = client.get('/api/v1/research/abnormal-runs/' + finished['run_id']).json()['data']
        assert run['request']['benchmark_sources']['sh_main']['fallback'] is True
        assert 'BENCHMARK_FALLBACK_USED' in run['quality_flags']
        assert run['result']['events']
        assert all(event['benchmark_symbol'] == 'sh000001' for event in run['result']['events'])


def test_five_factor_valuation_parity_scenarios_and_pe_provenance(tmp_path: Path, monkeypatch) -> None:
    factors = dict(earnings_yi=12.5, growth_coef=1.3, base_pe=22.5,
                   index_coef=1.1, sentiment_coef=1.8)
    assert asdict(new_valuation_calc(NewValuationInputs(**factors))) == asdict(
        old_valuation_calc(OldValuationInputs(**factors)))

    class FakeResponse:
        def __init__(self, data): self.data = data
        def raise_for_status(self): pass
        def json(self): return {'data': self.data}
    class FakeClient:
        def __init__(self, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def get(self, url, *, params, headers):
            if params['secid'] == '1.000001':
                return FakeResponse({'f43': 350000})
            return FakeResponse({'f43': 1000, 'f58': '样本公司', 'f116': 10000000000,
                                 'f127': '软件', 'f164': 0, 'f162': 2500, 'f163': 1000})
    app = create_app(tmp_path, auto_rebuild=False)
    with started_client(app) as client:
        monkeypatch.setattr(valuation_service.httpx, 'Client', FakeClient)
        csrf = client.get('/api/v1/session').json()['data']['csrf_token']
        quote = client.get('/api/v1/research/valuation-quote?symbol=600000').json()['data']
        assert quote['price'] == 10 and quote['market_cap_yi'] == 100
        assert quote['pe_ttm'] is None and quote['pe_dynamic'] == 25 and quote['pe_static'] == 10
        assert quote['implied_earnings_yi'] is None
        assert 'TTM_PE_NOT_AVAILABLE_EARNINGS_NOT_INFERRED' in quote['quality_flags']
        assert quote['index_close'] == 3500 and quote['suggested_pe_tier'] == 'tech'
        body = {'symbol': '600000', 'scenarios': [
            {'label': '基准', 'earnings_yi': '12.5', 'growth_rate_pct': 30,
             'base_pe': 22.5, 'index_points': 3300, 'sentiment_coef': 1.8,
             'actual_cap_yi': '400'},
            {'label': '亏损', 'earnings_yi': '-1', 'growth_rate_pct': 10,
             'base_pe': 20, 'index_points': 3000, 'sentiment_coef': 1.0}]}
        response = client.post('/api/v1/research/valuation-runs', json=body, headers={
            'X-CSRF-Token': csrf, 'Idempotency-Key': 'valuation-run'})
        assert response.status_code == 200, response.text
        run = response.json()['data']
        assert run['result']['scenarios'][0]['theoretical_cap_yi'] == round(
            old_valuation_calc(OldValuationInputs(**factors)).theoretical_cap_yi, 4)
        assert run['result']['scenarios'][1]['theoretical_cap_yi'] is None
        assert run['result']['scenarios'][1]['status'] == 'not_applicable_nonpositive_earnings_or_growth'
        assert client.get('/api/v1/research/valuation-runs/' + run['id']).json()['data']['result'] == run['result']
        assert client.get('/api/v1/research/valuation-runs').json()['data'][0]['id'] == run['id']





