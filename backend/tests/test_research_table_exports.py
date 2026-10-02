import csv
import io
import json
from copy import deepcopy

from openpyxl import load_workbook
import pytest

from trade_app.api import research_table_exports as export
from trade_app.platform.db import open_database
from trade_app.platform.types import TradeError
from trade_app.platform.compute_process import ComputeProcessError
from trade_app.research import legacy_report_domain, legacy_report_service, plateau_service, signal_workspace_service
from trade_app.research import portfolio_service
from trade_app.research.backtest_models import BacktestRun
from trade_app.research.plateau_models import PlateauPoint
from test_legacy_reports import original_package, modify
from test_plateau import create_experiment
from test_signal_workspace import scenario, prepare
from test_wyckoff_research_api import client_for


def csv_rows(raw): return list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
def sheet_rows(workbook, name):
    values = list(workbook[name].values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def test_csv_html_and_workbook_formula_and_markup_are_data_and_long_text_is_lossless():
    rows = [{'name': ' =HYPERLINK("bad")', 'note': '<script>bad()</script>', 'amount': '-12.50'}]
    raw = export.csv_bytes(['name', 'note', 'amount'], rows)
    assert raw.startswith(b'\xef\xbb\xbf')
    assert csv_rows(raw)[0]['name'].startswith("'")
    assert csv_rows(raw)[0]['amount'] == '-12.50'
    html = export.table_html('<script>title</script>', ['name', 'note', 'amount'], rows, {'range': '<bad>'})
    assert b'<script>' not in html and b'&lt;script&gt;' in html
    long = 'x' * 32000 + '=at_chunk_boundary' + 'x' * 34000
    workbook = load_workbook(io.BytesIO(export.workbook_bytes([('Points', ['name', 'note', 'long'], [{**rows[0], 'long': long}])])), data_only=False)
    assert workbook['Points']['A2'].data_type == 's'
    assert ''.join(row['text'] for row in sheet_rows(workbook, 'LongText')) == long
    assert workbook['LongText']['D3'].data_type == 's'


def test_legacy_exact_trade_columns_and_all_plateau_points_preserve_failed_and_source_order(tmp_path, original_package):
    trade = {'symbol': 'sh600000', 'name': '=BAD()', 'signal_date': '2025-01-02', 'entry_date': '2025-01-03',
        'exit_date': '2025-01-05', 'quantity': 100, 'entry_price': 10, 'exit_price': 11, 'pnl_amount': 90,
        'pnl_ratio': .09, 'entry_signal': '<script>bad()</script>', 'entry_phase': 'D', 'exit_reason': '平仓'}
    raw = modify(original_package[0], 'run_result.json', lambda value: {**value, 'trades': [trade]})
    prepared = legacy_report_domain.prepare_import(raw, 'old.ftbt')
    engine, factory = open_database(tmp_path)
    try:
        with factory.begin() as session:
            report = legacy_report_service.save_report(session, prepared, prepared['preview']['preview_sha256'], 'legacy_readonly', True)
        with factory() as session:
            columns, rows, metadata = export.legacy_trades(session, report['id'])
            assert columns == export.LEGACY_TRADE_COLUMNS and len(columns) == 21
            assert rows[0]['盈亏比百分比'] == '9.00'
            assert rows[0]['质量分'] is None  # Never invent the old UI's zero fallback.
            assert metadata['date_from'] == '2025-01-01' and metadata['date_to'] == '2025-02-01'
            assert csv_rows(export.csv_bytes(columns, rows))[0]['股票名称'] == "'=BAD()"
            assert b'<script>' not in export.table_html('trade', columns, rows, metadata)
            assert export.legacy_trades(session, report['id'], 'point_test_01')[1] == []
            with pytest.raises(TradeError): export.legacy_trades(session, report['id'], 'missing')
            workbook = load_workbook(io.BytesIO(export.legacy_plateau_xlsx(session, report['id'])))
            assert workbook.sheetnames[:6] == ['PlateauSummary', 'BasePayload', 'PlateauPoints', 'PlateauRegions', 'PlateauCorr', 'PlateauNotes']
            points = sheet_rows(workbook, 'PlateauPoints')
            assert len(points) == 2 and [row['rank'] for row in points] == [1, 2]
            assert [row['status'] for row in points] == ['succeeded', 'failed']
            assert points[1]['error'] == '<script>untrusted()</script>'
            assert points[0]['window_days'] == original_package[3]['points'][0]['params']['window_days']
            assert json.loads(points[0]['source_point_json']) == original_package[3]['points'][0]
    finally: engine.dispose()


def test_saved_filtered_cross_validation_csv_matches_union_dates_scores_and_omits_excluded(scenario):
    factory, _path, _dataset, _scan = scenario
    prepared = prepare(scenario, min_overlap=2)
    with factory.begin() as session: saved = signal_workspace_service.save_report(session, prepared)
    with factory() as session:
        raw = export.cross_validation_csv(session, saved['id'])
        rows = csv_rows(raw)
        assert len(rows) == len(saved['result']['per_symbol']) == 1
        item = saved['result']['per_symbol'][0]
        assert rows[0]['代码'] == item['symbol']
        assert int(rows[0]['共振强度']) == item['overlap_count'] * len(item['signal_dates'])
        assert float(rows[0]['区间涨幅%']) == item['range_performance']['return_pct']
        assert rows[0]['报告ID'] == saved['id'] and rows[0]['扫描结束日'] == '2025-03-29'
        assert rows[0]['实际价格止日'] == '2025-03-29'
        assert json.loads(rows[0]['冻结过滤条件'])['min_overlap'] == 2
        for strategy in item['strategies']:
            key = next(key for key in rows[0] if key.endswith(f"[{strategy['strategy_id']}]_出现日期"))
            assert rows[0][key] == '; '.join(strategy['dates'])
        assert export.cross_validation_csv(session, saved['id']) == raw  # No current clock or rerun.
    excluded = prepare(scenario, markets=['bj'])
    with factory.begin() as session: excluded = signal_workspace_service.save_report(session, excluded)
    with factory() as session: assert csv_rows(export.cross_validation_csv(session, excluded['id'])) == []


def test_real_native_history_and_mixed_point_exports_are_readonly_hash_checked_api(tmp_path, monkeypatch):
    with client_for(tmp_path) as client:
        from trade_app.api.research_table_export_routes import router
        if not any(getattr(route, 'path', '') == '/api/v1/backtests/{run_id}/trades.{kind}' for route in client.app.routes):
            client.app.router.routes[0:0] = router.routes
        run, _body, _preview = create_experiment(client, tmp_path, count=3)
        factory = client.app.state.db_factory
        assert client.get('/api/v1/research/plateaus/' + run['id'] + '/export.xlsx').status_code == 409

        assert plateau_service.process_one_plateau(factory, tmp_path)
        original_runner = plateau_service.run_json_process
        def fail(*args, **kwargs): raise ComputeProcessError('TEST_FAILURE', '=malicious()')
        monkeypatch.setattr(plateau_service, 'run_json_process', fail)
        assert plateau_service.process_one_plateau(factory, tmp_path)
        monkeypatch.setattr(plateau_service, 'run_json_process', original_runner)
        assert plateau_service.process_one_plateau(factory, tmp_path)
        source = run['source_run_id']
        for kind in ('csv', 'html'):
            response = client.get('/api/v1/backtests/' + source + '/trades.' + kind)
            assert response.status_code == 200, response.text
            assert response.headers['content-disposition'].startswith('attachment;')
            assert response.headers['x-content-type-options'] == 'nosniff'
            assert 'sandbox' in response.headers['content-security-policy']
        response = client.get('/api/v1/research/plateaus/' + run['id'] + '/export.xlsx')
        assert response.status_code == 200, response.text
        workbook = load_workbook(io.BytesIO(response.content))
        points = sheet_rows(workbook, 'PlateauPoints')
        assert len(points) == 3 and [point['status'] for point in points] == ['succeeded', 'failed', 'succeeded']
        assert points[1]['total_return'] is None and points[1]['plateau_score'] is None
        with factory() as session:
            assert plateau_service.get_plateau(session, run['id'])['counts'] == {'succeeded': 2, 'failed': 1}
        with factory.begin() as session:
            row = session.get(BacktestRun, source)
            original_result = json.loads(row.result_json)
            row.result_json = json.dumps({**original_result, 'trade_count': 99})
        assert client.get('/api/v1/backtests/' + source + '/trades.csv').status_code == 409
        with factory.begin() as session:
            point = session.get(PlateauPoint, run['points'][0]['id'])
            point.result_json = '{}'
        assert client.get('/api/v1/research/plateaus/' + run['id'] + '/export.xlsx').status_code == 409


def test_real_portfolio_history_csv_preserves_all_fills_dates_and_chain_integrity(tmp_path):
    from test_portfolio import create_run
    from trade_app.research.portfolio_models import PortfolioChunk
    with client_for(tmp_path) as client:
        run, _body = create_run(client, tmp_path)
        factory = client.app.state.db_factory
        while portfolio_service.process_one_portfolio(factory, tmp_path): pass
        with factory() as session:
            columns, rows, meta = export.native_trades(session, run['id'], portfolio=True)
            saved = portfolio_service.get_portfolio(session, run['id'], full=True)
            assert len(rows) == len(saved['result']['trades'])
            assert [row['date'] for row in rows] == [row['date'] for row in saved['result']['trades']]
            assert meta['date_from'] == saved['result']['equity'][0]['date']
            assert len(csv_rows(export.csv_bytes(columns, rows))) == len(rows)
        with factory.begin() as session:
            from sqlalchemy import select
            chunk = session.scalar(select(PortfolioChunk).where(PortfolioChunk.run_id == run['id']))
            chunk.result_json = '{}'
        with factory() as session:
            with pytest.raises(TradeError): export.native_trades(session, run['id'], portfolio=True)
