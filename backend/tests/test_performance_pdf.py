from copy import deepcopy
from datetime import date, timedelta
import re
import csv
import io

import pytest
from sqlalchemy import select
from openpyxl import load_workbook

from trade_app.analytics.service import process_one
from trade_app.platform.types import TradeError
from trade_app.reviews.performance_pdf import render_statistics
from trade_app.trading.sim_models import SimWallet
from test_wyckoff_research_api import client_for, data, write


def periods(count=45):
    return {'kind': 'daily', 'projection_status': 'fresh', 'projection_version': 'projection-001',
        'calculation_version': 'nav-v2', 'method': '已确认快照计算收益，入金与出金按份额法处理。缺失价格不会补零。',
        'items': [{'key': (date(2025, 1, 1) + timedelta(days=i)).isoformat(),
            'start_date': '2025-01-01', 'end_date': '2025-02-14', 'return_pct': '1.2300',
            'return_quality': 'complete', 'baseline_date': '2024-12-31', 'last_confirmed_date': '2025-02-14',
            'confirmed_points': 2, 'min_drawdown_pct': '-3.2100', 'node_achievements': [],
            'rounds': {'closed_rounds': 1, 'closed_pnl': '123456.78', 'win_rate_pct': '100',
                      'payoff_ratio': None, 'profit_factor': None, 'round_ids': ['r' + str(i).zfill(63)]}}
            for i in range(count)]}


def pdf_valid(value):
    assert value.startswith(b'%PDF-') and value.rstrip().endswith(b'%%EOF')
    assert b'/FontFile2' in value  # Embedded repository font; no OS font dependency.
    return len(re.findall(rb'/Type\s*/Page\b', value))


def test_statistics_pdf_long_tables_chinese_and_empty_are_complete(tmp_path):
    raw = render_statistics(periods(), name='中文长账户名 <测试> & 财务统计', account_id='a' * 32,
                            account_kind='real', input_revision=89, generated_at='2026-09-26T00:00:00Z')
    assert pdf_valid(raw) >= 5
    (tmp_path / 'real-statistics.pdf').write_bytes(raw)
    empty = periods(0)
    empty['projection_status'] = 'stale'
    assert pdf_valid(render_statistics(empty, name='无数据账户', account_id='b' * 32,
                                      account_kind='real', input_revision=1)) == 1


def test_pdf_endpoint_uses_selected_period_read_only_and_validates_account(tmp_path, monkeypatch):
    import trade_app.reviews.performance_pdf as module
    with client_for(tmp_path / 'app') as client:
        account = data(write(client, '/accounts', {'name': '实盘中文统计'}))
        base = '/accounts/' + account['id']
        data(write(client, base + '/cash-flows', {'flow_date': '2025-01-01', 'kind': 'initial', 'amount': '10000'}))
        for day, amount in [('2025-01-01', '10000'), ('2025-01-02', '11000')]:
            data(write(client, base + '/snapshots/' + day, {'snap_date': day, 'total_assets': amount, 'expected_revision': 0}, 'PUT'))
        assert process_one(client.app.state.db_factory)
        original = module.render_statistics
        seen = []
        def capture(value, **kwargs):
            seen.append((deepcopy(value), kwargs))
            return original(value, **kwargs)
        monkeypatch.setattr(module, 'render_statistics', capture)
        before = data(client.get('/api/v1/accounts'))
        response = client.get('/api/v1' + base + '/exports/performance.pdf?kind=daily&limit=1')
        assert response.status_code == 200 and response.headers['cache-control'] == 'no-store'
        pdf_valid(response.content)
        assert seen[0][0]['kind'] == 'daily' and len(seen[0][0]['items']) == 1
        assert seen[0][0]['items'][0]['return_pct'] == '10.0000'
        csv_response = client.get('/api/v1' + base + '/exports/performance/summary.csv?kind=daily&limit=1')
        csv_rows = list(csv.DictReader(io.StringIO(csv_response.content.decode('utf-8-sig'))))
        assert len(csv_rows) == 1 and csv_rows[0]['return_pct'] == '10.0000'
        assert csv_rows[0]['key'] == '2025-01-02'
        xlsx = client.get('/api/v1' + base + '/exports/performance.xlsx?kind=daily&limit=1')
        workbook = load_workbook(io.BytesIO(xlsx.content), read_only=True)
        assert workbook.sheetnames == ['metadata', 'summary', 'rounds', 'nodes']
        assert workbook['summary']['D2'].value == '10.0000'
        assert client.get('/api/v1' + base + '/exports/performance/fills.csv').status_code == 400
        assert data(client.get('/api/v1/accounts')) == before
        for query in ('kind=yearly', 'date_basis=other', 'limit=367'):
            assert client.get('/api/v1' + base + '/exports/performance.pdf?' + query).status_code in (400, 422)
        assert client.get('/api/v1/accounts/' + 'f' * 32 + '/exports/performance.pdf').status_code == 404


def test_sim_pdf_uses_fifo_buy_month_and_wallet_revision(tmp_path, monkeypatch):
    import trade_app.reviews.performance_pdf as module
    with client_for(tmp_path / 'app') as client:
        account = data(write(client, '/sim-accounts', {'name': '模拟账户精确批次统计',
            'initial_capital': '100000', 'start_date': '2025-01-30'}))
        base = '/sim-accounts/' + account['id']
        def fill(day, side, qty, price):
            order = data(write(client, base + '/orders', {'symbol': '600000', 'side': side,
                'quantity': qty, 'limit_price': price, 'signal_date': day, 'submit_date': day}))
            return data(write(client, base + '/orders/' + order['id'] + '/fill', {
                'expected_revision': 1, 'fill_date': day, 'fill_price': price}))['fill']
        buy = fill('2025-01-30', 'buy', 200, '10')
        data(write(client, base + '/settle', {'to_date': '2025-02-03'}))
        sell = fill('2025-02-03', 'sell', 100, '12')
        original = module.render_statistics
        seen = []
        def capture(value, **kwargs):
            seen.append(deepcopy(value))
            return original(value, **kwargs)
        monkeypatch.setattr(module, 'render_statistics', capture)
        response = client.get('/api/v1/accounts/' + account['id'] + '/exports/performance.pdf?date_basis=buy')
        assert response.status_code == 200
        pdf_valid(response.content)
        value = seen[0]
        assert value['date_basis'] == 'buy' and value['monthly'][0]['month'] == '2025-01'
        assert value['realized_pnl'] == sell['realized_pnl']
        assert value['closed_fills'][0]['buy_allocations'][0]['buy_fill_id'] == buy['id']
        root = '/api/v1/accounts/' + account['id'] + '/exports/performance'
        workbook = load_workbook(io.BytesIO(client.get(root + '.xlsx?date_basis=buy').content), read_only=True)
        assert workbook.sheetnames == ['metadata', 'summary', 'fills', 'allocations', 'curve', 'unallocated']
        assert workbook['summary']['A2'].value == '2025-01'
        assert workbook['allocations']['A2'].value == sell['id']
        assert workbook['allocations']['B2'].value == buy['id']
        rows = list(csv.DictReader(io.StringIO(client.get(root + '/allocations.csv').content.decode('utf-8-sig'))))
        assert rows[0]['cost_basis'] == sell['buy_allocations'][0]['cost_basis']
        assert rows[0]['realized_pnl'] == sell['realized_pnl']
        with client.app.state.db_factory() as session:
            assert value['wallet_revision'] == session.scalar(select(SimWallet.revision).where(SimWallet.account_id == account['id']))
        # Long table uses the same real calculation result, repeated for layout QA only.
        value['closed_fills'] = [dict(deepcopy(value['closed_fills'][0]), fill_id=f'layout-{i:032}') for i in range(50)]
        raw = original(value, name='模拟长表版式检查 <样本>', account_id=account['id'],
                       account_kind='sim', input_revision=0, generated_at='2026-09-26T00:00:00Z')
        assert pdf_valid(raw) >= 5
        (tmp_path / 'sim-statistics.pdf').write_bytes(raw)


def test_oversized_pdf_explicitly_rejects_instead_of_truncating():
    with pytest.raises(TradeError, match='5000'):
        render_statistics(periods(5001), name='超限', account_id='a' * 32,
                          account_kind='real', input_revision=1)


def test_statistics_metadata_neutralizes_spreadsheet_formulas(tmp_path):
    with client_for(tmp_path) as client:
        account = data(write(client, '/accounts', {'name': '=SUM(1,2)'}))
        root = '/api/v1/accounts/' + account['id'] + '/exports/performance'
        workbook = load_workbook(io.BytesIO(client.get(root + '.xlsx').content), read_only=True)
        metadata = dict(workbook['metadata'].values)
        assert metadata['name'] == "'=SUM(1,2)"
        rows = list(csv.DictReader(io.StringIO(client.get(root + '/metadata.csv').content.decode('utf-8-sig'))))
        assert next(row['value'] for row in rows if row['field'] == 'name') == "'=SUM(1,2)"
