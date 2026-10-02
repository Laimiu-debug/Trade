from copy import deepcopy
import csv
import io
import json
import re
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from trade_app.api.screener_export import build_export, csv_export, xlsx_export, pdf_export
from trade_app.api.screener_export_routes import router
from trade_app.platform.types import TradeError
from trade_app.research.screener_models import ScreenerRun
from trade_app.research.b1_models import B1Run
from test_wyckoff_research_api import client_for, write


def record(kind='funnel', count=3):
    rows = [{'symbol': f'sh{600000+i}', 'name': '中文证券 <名称> & 检查' if i else '=HYPERLINK("bad")',
             'dataset_id': f'{i+1:064x}', 'as_of_date': '2025-01-02', 'score': 72,
             'ret40': 0.12345678901234566, 'turnover20': None, 'amount20': 123456789.125,
             'amplitude20': .05, 'retrace20': -.0123, 'quality_flags': ['HISTORICAL_AVAILABLE_AT_UNKNOWN'],
             'close': 12.3456, 'change_pct': -1.23, 'amplitude_pct': 4.56, 'volume_ratio': .1234,
             'kdj_j': 15.2, 'weekly_macd': .0012, 'monthly_macd': .2345} for i in range(count)]
    result = {'as_of_date': '2025-01-02', 'quality_flags': ['HISTORICAL_AVAILABLE_AT_UNKNOWN'],
        'pools': {'input': rows, 'step1': rows[1:], 'step2': [], 'step3': [], 'step4': rows[-1:]},
        'hits': rows, 'rejections': {rows[0]['dataset_id']: {'stage': 'step1', 'reasons': ['TURNOVER_MISSING']}} if rows else {},
        'config': {'mode': 'strict', 'step1': {'amount_threshold': 50000000}}, 'b1_params': {'vol_ratio': .8},
        'return_window_days': 40}
    request = {'datasets': [{'dataset_id': row['dataset_id'], 'float_shares': 123456789.125,
                            'float_shares_as_of_date': '2024-12-31'} for row in rows],
               'config': result['config'], 'b1_params': result['b1_params'], 'return_window_days': 40,
               'as_of_date': result['as_of_date']}
    return SimpleNamespace(id=('a' if kind == 'funnel' else 'b')*64, request_json=json.dumps(request, ensure_ascii=False),
        result_json=json.dumps(result, ensure_ascii=False), code_sha256='c'*64, created_at='2025-01-03T00:00:00Z')


def export_value(kind='funnel', count=3, **selection):
    row = record(kind, count)
    return build_export(SimpleNamespace(get=lambda *args: row), kind, row.id, selection)


def test_frozen_stage_selection_order_precision_and_sources():
    value = export_value(stage='input', dataset_ids=[f'{3:064x}', f'{1:064x}'], columns=['amount20', 'ret40', 'turnover20'])
    assert [row['symbol'] for row in value['rows']] == ['sh600000', 'sh600002']
    assert list(value['columns']) == ['symbol', 'name', 'amount20', 'ret40', 'turnover20']
    assert value['metadata']['parameters']['return_window_days'] == 40
    assert value['metadata']['selected_count'] == 2
    assert value['sources'][0]['float_shares'] == 123456789.125
    workbook = load_workbook(io.BytesIO(xlsx_export(value)))
    result = list(workbook['SelectedRows'].values)
    assert result[1] == ('1', 'sh600000', '\'=HYPERLINK("bad")', '123456789.125', '0.12345678901234566', None)
    assert not any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row)
    csv_rows = list(csv.reader(io.StringIO(csv_export(value).decode('utf-8-sig'))))
    assert csv_rows[1][-5:] == ['sh600000', '\'=HYPERLINK("bad")', '123456789.125', '0.12345678901234566', '']
    assert csv_rows[1][0] == value['metadata']['run_id']
    assert value['metadata']['selection_sha256'] == export_value(stage='input', dataset_ids=[f'{1:064x}', f'{3:064x}'], columns=['amount20','ret40','turnover20'])['metadata']['selection_sha256']


@pytest.mark.parametrize('selection', [
    {'stage': 'hits'}, {'stage': 'step1', 'dataset_ids': [f'{1:064x}']}, {'stage': 'input', 'dataset_ids': []},
    {'stage': 'input', 'dataset_ids': [f'{1:064x}']*2}, {'columns': ['unknown']}, {'columns': ['symbol','symbol']}, {'columns': []},
])
def test_bad_stage_subset_and_column_requests_rejected(selection):
    with pytest.raises(TradeError): export_value(**selection)


def capture_paragraphs(monkeypatch):
    import trade_app.api.screener_export as module
    original = module.Paragraph; paragraphs = []
    def wrapped(text, *args, **kwargs):
        paragraphs.append(text)
        return original(text, *args, **kwargs)
    monkeypatch.setattr(module, 'Paragraph', wrapped)
    return paragraphs


def test_b1_percent_units_and_pdf_full_rows_embedded_chinese_font(tmp_path, monkeypatch):
    paragraphs = capture_paragraphs(monkeypatch)
    value = export_value('b1', 50, columns=['close','change_pct','amplitude_pct','volume_ratio','kdj_j'])
    content = pdf_export(value)
    assert content.startswith(b'%PDF-') and b'/FontFile2' in content
    text = '\n'.join(paragraphs)
    assert len(re.findall(rb'/Type\s*/Page\b', content)) >= 3 and 'sh600049' in text and 'sh600000' in text
    assert '中文证券' in text and '百分数' in text and '-1.23' in text
    assert value['rows'][0]['change_pct'] == -1.23
    (tmp_path / 'screening-b1.pdf').write_bytes(content)


def test_empty_stage_exports_metadata_without_fabricating_security_rows(monkeypatch):
    paragraphs = capture_paragraphs(monkeypatch)
    value = export_value(stage='step2')
    workbook = load_workbook(io.BytesIO(xlsx_export(value)))
    assert workbook['SelectedRows'].max_row == 1
    rows = list(csv.reader(io.StringIO(csv_export(value).decode('utf-8-sig'))))
    start = rows[0].index('source_json')
    assert len(rows) == 2 and rows[1][start:] == ['']*(len(rows[0])-start)
    assert pdf_export(value).startswith(b'%PDF-')
    text = '\n'.join(paragraphs)
    assert '阶段为空' in text and 'sh600000' not in text


def test_export_api_reads_persisted_evidence_only_and_same_selection_all_formats(tmp_path):
    with client_for(tmp_path) as client:
        if not any(route.path.endswith('/screener-runs/{run_id}/exports/{format}') for route in client.app.routes):
            client.app.include_router(router)
        with client.app.state.db_factory.begin() as session:
            for kind, model in [('funnel',ScreenerRun),('b1',B1Run)]:
                session.add(model(**vars(record(kind))))
        for kind, prefix in [('funnel','screener'),('b1','b1')]:
            run_id = ('a' if kind=='funnel' else 'b')*64
            body = {'stage': 'step1' if kind=='funnel' else 'hits', 'dataset_ids':[f'{3:064x}'], 'columns':['symbol','name']}
            hashes = set()
            for extension, signature in [('csv',b'\xef\xbb\xbf'),('xlsx',b'PK'),('pdf',b'%PDF-')]:
                response = write(client,f'/research/{prefix}-runs/{run_id}/exports/{extension}',body)
                assert response.status_code == 200, response.text
                assert response.content.startswith(signature)
                assert response.headers['X-Export-Row-Count'] == '1'
                hashes.add(response.headers['X-Export-Selection-SHA256'])
                assert response.headers['Cache-Control'] == 'no-store'
            assert len(hashes)==1
            with client.app.state.db_factory() as session:
                assert session.get(ScreenerRun if kind=='funnel' else B1Run,run_id).result_json == record(kind).result_json
        bad=write(client,f'/research/screener-runs/{"a"*64}/exports/csv', {'stage':'step1','dataset_ids':[f'{1:064x}']})
        assert bad.status_code == 409
        assert write(client,f'/research/screener-runs/{"a"*64}/exports/csv', {'html':'untrusted'}).status_code == 422
        assert write(client,f'/research/screener-runs/{"f"*64}/exports/csv', {}).status_code == 404
