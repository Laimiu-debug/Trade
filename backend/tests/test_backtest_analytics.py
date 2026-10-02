import copy
import io
import json
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from trade_app.research.backtest_analytics import build_backtest_analysis, completed_trades
from trade_app.research.backtest_service import process_one_backtest
from test_wyckoff_research_api import client_for, data, dataset, write


def sample(pnls=(10, -20, 0, 30, -10, 5, -4, 14, -2, 12)):
    trades, equity, bars = [], [], []
    capital = 1000
    day = date(2025, 1, 1)
    for value in pnls:
        buy, sell = day.isoformat(), (day + timedelta(days=1)).isoformat()
        # Cost = 101 including entry fee; sale proceeds = 101 + pnl.
        trades += [{'date': buy, 'side': 'buy', 'price': '1', 'quantity': 100, 'fees': '1', 'realized_pnl': None},
                   {'date': sell, 'side': 'sell', 'price': str((102 + value) / 100), 'quantity': 100,
                    'fees': '1', 'realized_pnl': str(value), 'holding_bars': 1}]
        equity += [{'date': buy, 'total_assets': str(capital - 1), 'cash': str(capital - 101), 'quantity': 100},
                   {'date': sell, 'total_assets': str(capital + value), 'cash': str(capital + value), 'quantity': 0}]
        capital += value
        day += timedelta(days=2)
    for row in equity:
        bars.append({'event_date': row['date'], 'open': '1', 'high': '2', 'low': '.5', 'close': '1',
                     'volume': 1000, 'available_at': row['date'] + 'T08:00:00+00:00'})
    return {'initial_capital': '1000', 'equity': equity, 'trades': trades, 'signal_count': len(pnls)}, bars


def test_pairing_includes_both_fees_and_respects_actual_position_weight():
    result, bars = sample()
    pairs = completed_trades(result)
    assert pairs[0]['net_pnl'] == '10.00'
    assert pairs[0]['position_return'] == pytest.approx(10 / 101)
    assert pairs[0]['account_return'] == .01
    output = build_backtest_analysis(result, bars)
    risk = output['risk']
    assert risk['profit_factor'] == pytest.approx(71 / 36)
    assert risk['expectancy'] == pytest.approx(35 / 101 / 10)
    assert risk['fill_rate'] == 1
    assert risk['max_concurrent_positions'] == 1


def test_monte_carlo_frozen_seed_and_missing_values_are_explicit():
    result, bars = sample()
    original = copy.deepcopy(result)
    first = build_backtest_analysis(result, bars, seed=32, iterations=100)
    assert first == build_backtest_analysis(result, bars, seed=32, iterations=100)
    assert first['monte_carlo'] != build_backtest_analysis(result, bars, seed=33, iterations=100)['monte_carlo']
    assert result == original
    assert first['monte_carlo']['ruin_probability'] == 0
    assert first['walk_forward']['status'] == 'not_generated'
    json.dumps(first, allow_nan=False)
    short, short_bars = sample((0, 0))
    short['equity'] = [{**row, 'total_assets': '1000'} for row in short['equity']]
    output = build_backtest_analysis(short, short_bars)
    assert output['risk']['sharpe'] is None
    assert output['risk']['profit_factor'] is None
    assert output['monte_carlo']['status'] == 'insufficient_sample'
    assert output['monte_carlo']['iterations'] == 0
    assert output['stability']['monthly_return_std'] is None


def test_regime_never_reads_future_or_late_available_bars():
    result, bars = sample(tuple([10] * 13))
    for index, bar in enumerate(bars):
        bar['close'] = str(1 + index / 10)
    baseline = build_backtest_analysis(result, bars)['regimes']
    changed = copy.deepcopy(bars)
    changed[-1]['close'] = '10000'
    changed.append({**changed[-1], 'event_date': '2099-01-01', 'close': '.01', 'available_at': '2099-01-01T08:00:00+00:00'})
    assert build_backtest_analysis(result, changed)['regimes'] == baseline
    delayed = [{**bar, 'available_at': '2099-01-01T08:00:00+00:00'} for bar in bars]
    buckets = build_backtest_analysis(result, delayed)['regimes']['buckets']
    assert next(row for row in buckets if row['regime'] == 'unknown')['trade_count'] == 13


@pytest.mark.parametrize('seed,iterations', [(True, 100), (-1, 100), (1, 99), (1, 2001), (1, 100.0)])
def test_monte_carlo_budget(seed, iterations):
    result, bars = sample()
    with pytest.raises(ValueError):
        build_backtest_analysis(result, bars, seed=seed, iterations=iterations)


def test_unrecovered_drawdown_is_not_claimed_as_recovered():
    result, bars = sample((-50, -50, -50))
    recovery = build_backtest_analysis(result, bars)['risk']['recovery']
    assert recovery['status'] == 'unrecovered'
    assert recovery['calendar_days'] is None and recovery['recovery_date'] is None


def test_worker_freezes_opt_in_analysis_export_and_report_round_trip(tmp_path):
    with client_for(tmp_path) as client:
        ds = dataset(client)
        body = {'dataset_id': ds['id'], 'strategy_id': 'trend_king_v1', 'advanced_analysis': True,
                'analysis_seed': 45, 'analysis_iterations': 100}
        queued = data(write(client, '/backtests', body))
        assert process_one_backtest(client.app.state.db_factory, tmp_path)
        run = data(client.get('/api/v1/backtests/' + queued['id']))
        assert run['state'] == 'succeeded', run['error']
        analysis = run['result']['advanced_analysis']
        assert analysis['status'] == 'generated' and analysis['monte_carlo']['seed'] == 45
        workbook = load_workbook(io.BytesIO(client.get('/api/v1/backtests/' + run['id'] + '/export.xlsx').content))
        assert {'Analysis', 'Monthly', 'Regimes', 'RoundTrips'} <= set(workbook.sheetnames)
        report = data(write(client, '/research/reports', {'backtest_run_id': run['id'], 'title': '高级分析冻结'}))
        stored = data(client.get('/api/v1/research/reports/' + report['id']))
        assert stored['payload']['run']['result']['advanced_analysis'] == analysis
        assert '高级分析与计算口径' in client.get('/api/v1/research/reports/' + report['id'] + '/report.html').text
        disabled = data(write(client, '/backtests', {**body, 'advanced_analysis': False}))
        assert disabled['id'] != run['id']
        assert process_one_backtest(client.app.state.db_factory, tmp_path)
        assert data(client.get('/api/v1/backtests/' + disabled['id']))['result']['advanced_analysis']['status'] == 'not_generated'
        assert write(client, '/backtests', {**body, 'unimplemented_stop': True}).status_code == 422
